"""统计真实业务文件中，同一项目编号下各候选共享字段是否真的稳定。

只输出计数，不打印任何业务值。用于决定 Task 1 共享字段校验该覆盖哪些列。
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import load_workbook  # noqa: E402

CANDIDATES = {
    "department": 3,
    "branch_company": 4,
    "account_manager": 5,
    "team_level3_name": 9,
    "customer_unit_name": 10,
    "end_user_name": 11,
    "regional_platform": 12,
    "project_name": 14,
}


def main() -> int:
    source = Path(sys.argv[1])
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(min_row=3, values_only=True)

    seen: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    counts: dict[str, int] = defaultdict(int)
    for row in rows:
        project_code = row[1] if len(row) > 1 else None
        if not project_code or not str(project_code).strip():
            continue
        project_code = str(project_code).strip()
        counts[project_code] += 1
        for field, column in CANDIDATES.items():
            value = row[column - 1] if len(row) >= column else None
            text = str(value).strip() if value not in (None, "") else ""
            if text:
                seen[project_code][field].add(text)
    workbook.close()

    multi_line_projects = {code for code, lines in counts.items() if lines > 1}
    print(f"数据行所属项目数：{len(counts)}，其中多明细项目：{len(multi_line_projects)}")
    print(f"{'字段':<22}{'多值项目数':>10}{'占比':>9}")
    for field in CANDIDATES:
        unstable = sum(
            1 for code in multi_line_projects if len(seen[code][field]) > 1
        )
        ratio = unstable / len(multi_line_projects) if multi_line_projects else 0
        print(f"{field:<22}{unstable:>10}{ratio:>9.1%}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
