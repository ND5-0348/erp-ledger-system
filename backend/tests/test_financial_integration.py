from __future__ import annotations

import json
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from conftest import TEST_PASSWORD
from app.auth import ALL_PERMISSIONS, CurrentUser
from app.config import ROOT_DIR, settings
from app.db import db
from app.ledger_excel import (
    EDITABLE_ORDER_KEYS,
    PURCHASE_EDITABLE_COLUMN_NUMBERS,
    PURCHASE_EDITOR_KEYS_BY_COLUMN,
    SALES_EDITABLE_COLUMN_NUMBERS,
    SALES_EDITOR_KEYS_BY_COLUMN,
    SAMPLE_ORDER_NO,
    SAMPLE_PROJECT_CODE,
    TEMPLATE_HEADERS,
)
from app.routers import purchases

# 建库、清库、client、headers 等夹具统一在 backend/tests/conftest.py，
# 该文件负责拒绝业务库名与业务备份目录。


def _payload(suffix: str) -> dict[str, str]:
    return {
        "amount_type": "gross",
        "project_code": f"QA-{suffix}",
        "project_name": "Financial Integration Test",
        "department": "QA",
        "branch_company": "QA Branch",
        "account_manager": "QA Manager",
        "order_no": f"SO-{suffix}",
        "order_date": "2026-07-14",
        "business_type": "QA",
        "statistical_category": "QA",
        "team_name": "QA Team",
        "customer_unit_name": "QA Customer",
        "user_name": "QA User",
        "regional_platform": "QA Platform",
        "goods_name": "QA Equipment",
        "specification_model": "QA-SPEC",
        "unit_name": "unit",
        "quantity": "10.000000",
        "net_unit_price": "100.000000",
        "unit_price": "113.000000",
        "net_revenue": "1000.00",
        "order_value": "1130.00",
        "supplier_name": "QA Supplier",
        "purchase_unit_price_no_tax": "70.000000",
        "purchase_unit_price": "79.100000",
        "cost_no_tax": "700.00",
        "purchase_amount": "791.00",
        "delivery_date": "2026-07-15",
        "delivery_quantity": "5.000000",
        "delivery_revenue_no_tax": "500.00",
        "delivery_value": "565.00",
        "delivery_cost_no_tax": "350.00",
        "delivery_cost": "395.50",
        "pending_delivery_quantity": "5.000000",
        "pending_delivery_amount_no_tax": "500.00",
        "pending_delivery_amount": "565.00",
    }


def _create_order(client: TestClient, headers: dict[str, str], suffix: str, **updates: str) -> tuple[int, dict[str, str]]:
    payload = _payload(suffix)
    payload.update(updates)
    response = client.post("/api/orders", json=payload, headers=headers)
    assert response.status_code == 200, response.text
    return int(response.json()["order_line_id"]), payload


def _d(value: object) -> Decimal:
    return Decimal(str(value))


def _finance(order_line_id: int) -> dict:
    with db() as conn:
        row = conn.execute(
            text("SELECT * FROM v_order_line_finance WHERE order_line_id = :id"),
            {"id": order_line_id},
        ).mappings().one()
    return dict(row)


def _ledger(project_code: str) -> dict:
    with db() as conn:
        row = conn.execute(
            text("SELECT * FROM v_project_ledger_summary WHERE project_code = :code"),
            {"code": project_code},
        ).mappings().one()
    return dict(row)


def _count(table: str, condition: str = "1=1", **params: object) -> int:
    with db() as conn:
        return int(conn.execute(text(f"SELECT COUNT(*) FROM `{table}` WHERE {condition}"), params).scalar() or 0)


def _action_count(action: str) -> int:
    return _count("operation_log", "action_name = :action", action=action)


def _dashboard(client: TestClient, headers: dict[str, str]) -> dict:
    response = client.get("/api/dashboard/summary", headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def _payment(amount: str, payment_date: str = "2026-07-20") -> dict[str, str]:
    return {
        "due_payment_date": "2026-07-20",
        "payment_date": payment_date,
        "payment_voucher_no": f"PAY-{amount}",
        "payment_amount": amount,
    }


def _receipt(amount: str, ratio: str = "0.000000") -> dict[str, str]:
    return {
        "receipt_date": "2026-07-30",
        "payment_notice_no": f"RCPT-{amount}",
        "receipt_amount": amount,
        "receipt_ratio": ratio,
    }


def test_project_name_and_supplier_remain_bound_to_each_order_line(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    suffix = "PROJECT-NAME-GRAIN"
    first_id, _ = _create_order(
        client,
        headers,
        suffix,
        project_name="项目名称甲",
        goods_name="设备甲",
        specification_model="SPEC-A",
        supplier_name="采购商甲",
    )
    second_payload = _payload(suffix)
    second_payload.update(
        {
            "project_name": "项目名称乙",
            "goods_name": "设备乙",
            "specification_model": "SPEC-B",
            "supplier_name": "采购商乙",
        }
    )
    second_response = client.post("/api/orders", json=second_payload, headers=headers)
    assert second_response.status_code == 200, second_response.text
    second_id = int(second_response.json()["order_line_id"])

    response = client.get(f"/api/orders?project_id=QA-{suffix}", headers=headers)
    assert response.status_code == 200, response.text
    rows = {int(row["order_line_id"]): row for row in response.json()["items"]}
    assert rows[first_id]["project_name"] == "项目名称甲"
    assert rows[first_id]["supplier_name"] == "采购商甲"
    assert rows[second_id]["project_name"] == "项目名称乙"
    assert rows[second_id]["supplier_name"] == "采购商乙"

    updated = {**second_payload, "project_name": "项目名称乙-已修改"}
    update_response = client.put(f"/api/orders/{second_id}", json=updated, headers=headers)
    assert update_response.status_code == 200, update_response.text
    assert _finance(first_id)["project_name"] == "项目名称甲"
    assert _finance(second_id)["project_name"] == "项目名称乙-已修改"

    with db() as conn:
        order_summary = conn.execute(
            text(
                """
                SELECT project_name, line_count
                FROM v_order_ledger_summary
                WHERE project_code = :project_code AND order_no = :order_no
                """
            ),
            {"project_code": f"QA-{suffix}", "order_no": f"SO-{suffix}"},
        ).mappings().one()
    assert int(order_summary["line_count"]) == 2
    assert set(str(order_summary["project_name"]).split("；")) == {"项目名称甲", "项目名称乙-已修改"}


def _basic_payload(payload: dict[str, str]) -> dict[str, str]:
    return {key: value for key, value in payload.items() if key in EDITABLE_ORDER_KEYS}


def _purchase_editor_payload(editor_response: dict, row: dict) -> dict:
    payload = {"order_line_id": row["order_line_id"]}
    for index, column in enumerate(editor_response["columns"]):
        if column["editable"] and column["key"]:
            payload[column["key"]] = row["values"][index]
    return payload


def _sales_editor_payload(editor_response: dict, row: dict) -> dict:
    payload = {"order_line_id": row["order_line_id"]}
    for index, column in enumerate(editor_response["columns"]):
        if column["editable"] and column["key"]:
            payload[column["key"]] = row["values"][index]
    return payload


def test_orders_list_includes_the_latest_data_modification_time(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    _create_order(client, headers, "LATEST-MODIFIED")

    response = client.get("/api/orders", params={"project_id": "QA-LATEST-MODIFIED"}, headers=headers)

    assert response.status_code == 200, response.text
    item = response.json()["items"][0]
    assert item["last_modified_at"]
    assert _d(item["gross_profit"]) == Decimal("339.00")
    assert _d(item["accounts_receivable"]) == Decimal("1130.00")
    assert _d(item["accounts_payable"]) == Decimal("791.00")


def test_batch_basic_editor_schema_create_update_and_readonly_boundary(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    schema_response = client.get("/api/orders/batch-editor/schema", headers=headers)
    assert schema_response.status_code == 200, schema_response.text
    schema = schema_response.json()
    assert schema["editable_through"] == "W"
    assert len(schema["columns"]) == len(TEMPLATE_HEADERS) == 91
    assert [column["key"] for column in schema["columns"][:23]] == EDITABLE_ORDER_KEYS
    assert all(column["editable"] for column in schema["columns"][:23])
    assert not any(column["editable"] for column in schema["columns"][23:])
    assert {
        column["excel_column"] for column in schema["columns"] if column["required"]
    } == {"B", "J", "M", "O", "W"}

    first = _payload("BATCH-BASIC")
    second = {
        **first,
        "goods_name": "QA Equipment 2",
        "specification_model": "QA-SPEC-2",
        "order_value": "2260.00",
    }
    create_response = client.post(
        "/api/orders/batch-basic",
        json={"items": [_basic_payload(first), _basic_payload(second)]},
        headers=headers,
    )
    assert create_response.status_code == 200, create_response.text
    order_line_ids = create_response.json()["order_line_ids"]
    assert len(order_line_ids) == 2
    assert _action_count("batch_create_basic_orders") == 1

    rows_response = client.post(
        "/api/orders/batch-editor/rows",
        json={"order_line_ids": order_line_ids},
        headers=headers,
    )
    assert rows_response.status_code == 200, rows_response.text
    editor_rows = rows_response.json()["rows"]
    assert len(editor_rows) == 2
    assert all(len(row["values"]) == 91 for row in editor_rows)
    first_editor = next(row for row in editor_rows if row["order_line_id"] == order_line_ids[0])
    assert first_editor["values"][1] == first["project_code"]
    assert first_editor["values"][22] == first["order_value"]
    assert first_editor["values"][23] is None

    update_items = []
    for order_line_id, source, manager in (
        (order_line_ids[0], first, "批量客户经理"),
        (order_line_ids[1], second, "批量客户经理"),
    ):
        item = _basic_payload(source)
        item.update(
            {
                "order_line_id": order_line_id,
                "account_manager": manager,
                "customer_unit_name": "批量修改客户单位",
            }
        )
        update_items.append(item)
    update_response = client.put(
        "/api/orders/batch-basic",
        json={"items": update_items},
        headers=headers,
    )
    assert update_response.status_code == 200, update_response.text
    assert update_response.json()["updated"] == 2
    assert _action_count("batch_update_basic_order") == 1
    with db() as conn:
        logs = conn.execute(
            text(
                """
                SELECT detail
                FROM operation_log
                WHERE action_name = 'batch_update_basic_order'
                ORDER BY id
                """
            )
        ).scalars().all()
        manager = conn.execute(
            text("SELECT account_manager FROM project WHERE project_code = :code"),
            {"code": first["project_code"]},
        ).scalar_one()
        goods = conn.execute(
            text("SELECT goods_name FROM order_line WHERE id IN (:first, :second) ORDER BY id"),
            {"first": order_line_ids[0], "second": order_line_ids[1]},
        ).scalars().all()
    assert len(logs) == 1
    audit_detail = json.loads(str(logs[0]))
    audit_entries = audit_detail["batch_entries"]
    assert audit_detail["summary"] == "在线表格批量修改 2 条基本信息，共 4 个单元格"
    assert audit_detail["before"] is None
    assert audit_detail["after"] is None
    assert len(audit_entries) == 2
    assert all(set(entry["before"]) == {"order_line_id", *EDITABLE_ORDER_KEYS} for entry in audit_entries)
    assert all(entry["before"]["account_manager"] == first["account_manager"] for entry in audit_entries)
    assert all(entry["after"]["account_manager"] == "批量客户经理" for entry in audit_entries)
    assert all(entry["before"]["customer_unit_name"] == first["customer_unit_name"] for entry in audit_entries)
    assert all(entry["after"]["customer_unit_name"] == "批量修改客户单位" for entry in audit_entries)
    assert manager == "批量客户经理"
    assert goods == [first["goods_name"], second["goods_name"]]

    unchanged_response = client.put(
        "/api/orders/batch-basic",
        json={"items": update_items},
        headers=headers,
    )
    assert unchanged_response.status_code == 200, unchanged_response.text
    assert unchanged_response.json() == {"updated": 0, "order_line_ids": []}
    assert _action_count("batch_update_basic_order") == 1

    forbidden_readonly = {
        **update_items[0],
        "supplier_name": "不允许通过基本信息批量修改",
    }
    forbidden_response = client.put(
        "/api/orders/batch-basic",
        json={"items": [forbidden_readonly]},
        headers=headers,
    )
    assert forbidden_response.status_code == 422, forbidden_response.text


def test_purchase_batch_editor_round_trip_and_readonly_boundary(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    order_line_id, _ = _create_order(client, headers, "PURCHASE-BATCH")
    rows_response = client.post(
        "/api/purchases/batch-editor/rows",
        json={"order_line_ids": [order_line_id]},
        headers=headers,
    )
    assert rows_response.status_code == 200, rows_response.text
    editor = rows_response.json()
    assert editor["editable_from"] == "X"
    assert editor["editable_through"] == "BO"
    assert editor["fixed_columns"] == ["B", "C", "D", "E", "F", "G", "M", "N", "O"]
    assert {
        column["key"] for column in editor["columns"] if column["editable"]
    } == {
        PURCHASE_EDITOR_KEYS_BY_COLUMN[column_no]
        for column_no in PURCHASE_EDITABLE_COLUMN_NUMBERS
    }
    assert not {
        "pending_booked_amount",
        "total_paid",
        "accounts_payable",
        "gross_profit_no_tax",
        "tax_difference",
        "gross_profit",
        "gross_profit_margin_no_tax",
    } & {
        column["key"] for column in editor["columns"] if column["editable"]
    }

    item = _purchase_editor_payload(editor, editor["rows"][0])
    item.update(
        {
            "supplier_name": "批量采购厂商",
            "purchase_amount": "900.00",
            "delivery_quantity": "6.000000",
            "purchase_contract_no": "PC-BATCH-001",
            "received_invoice_date": "2026-07-21",
            "purchase_invoice_no": "PINV-BATCH-001",
            "purchase_invoice_amount": "300.00",
            "warehouse_date": "2026-07-22",
            "warehouse_voucher_no": "WH-BATCH-001",
            "warehouse_amount": "250.00",
            "booked_date": "2026-07-23",
            "booked_voucher_code": "BOOK-BATCH-001",
            "booked_amount": "200.00",
            "payment1_due_date": "2026-08-01",
            "payment1_date": "2026-07-24",
            "payment1_voucher_no": "PAY1-BATCH-001",
            "payment1_amount": "100.00",
            "payment2_date": "2026-07-25",
            "payment2_voucher_no": "PAY2-BATCH-001",
            "payment2_amount": "50.00",
        }
    )
    update_response = client.put(
        "/api/purchases/batch",
        json={"items": [item]},
        headers=headers,
    )
    assert update_response.status_code == 200, update_response.text
    assert update_response.json()["updated"] == 1
    assert _action_count("batch_update_purchases") == 1
    with db() as conn:
        audit_detail = json.loads(
            conn.execute(
                text(
                    """
                    SELECT detail
                    FROM operation_log
                    WHERE action_name = 'batch_update_purchases'
                    ORDER BY id DESC
                    LIMIT 1
                    """
                )
            ).scalar_one()
        )
    assert audit_detail["summary"].startswith("在线表格批量修改 1 条采购信息，共 ")
    assert len(audit_detail["batch_entries"]) == 1
    assert audit_detail["batch_entries"][0]["before"]["supplier_name"] == "QA Supplier"
    assert audit_detail["batch_entries"][0]["after"]["supplier_name"] == "批量采购厂商"

    refreshed_response = client.post(
        "/api/purchases/batch-editor/rows",
        json={"order_line_ids": [order_line_id]},
        headers=headers,
    )
    assert refreshed_response.status_code == 200, refreshed_response.text
    refreshed = refreshed_response.json()
    values_by_key = {
        column["key"]: refreshed["rows"][0]["values"][index]
        for index, column in enumerate(refreshed["columns"])
        if column["key"]
    }
    assert values_by_key["supplier_name"] == "批量采购厂商"
    assert values_by_key["purchase_contract_no"] == "PC-BATCH-001"
    assert values_by_key["purchase_invoice_no"] == "PINV-BATCH-001"
    assert values_by_key["warehouse_voucher_no"] == "WH-BATCH-001"
    assert values_by_key["booked_voucher_code"] == "BOOK-BATCH-001"
    assert values_by_key["payment1_voucher_no"] == "PAY1-BATCH-001"
    assert values_by_key["payment2_voucher_no"] == "PAY2-BATCH-001"
    assert _d(values_by_key["total_paid"]) == Decimal("150.00")
    assert _d(values_by_key["accounts_payable"]) == Decimal("750.00")

    unchanged_response = client.put(
        "/api/purchases/batch",
        json={"items": [item]},
        headers=headers,
    )
    assert unchanged_response.status_code == 200, unchanged_response.text
    assert unchanged_response.json() == {"updated": 0, "order_line_ids": []}
    assert _action_count("batch_update_purchases") == 1

    forbidden = {**item, "gross_profit": "1.00"}
    forbidden_response = client.put(
        "/api/purchases/batch",
        json={"items": [forbidden]},
        headers=headers,
    )
    assert forbidden_response.status_code == 422, forbidden_response.text


def test_purchase_batch_update_is_atomic(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    first_id, _ = _create_order(client, headers, "PURCHASE-BATCH-ATOMIC-1")
    second_id, _ = _create_order(client, headers, "PURCHASE-BATCH-ATOMIC-2")
    rows_response = client.post(
        "/api/purchases/batch-editor/rows",
        json={"order_line_ids": [first_id, second_id]},
        headers=headers,
    )
    assert rows_response.status_code == 200, rows_response.text
    editor = rows_response.json()
    rows_by_id = {row["order_line_id"]: row for row in editor["rows"]}
    first = _purchase_editor_payload(editor, rows_by_id[first_id])
    second = _purchase_editor_payload(editor, rows_by_id[second_id])
    first["supplier_name"] = "该修改必须回滚"
    second["purchase_amount"] = "100.00"
    second["payment1_amount"] = "101.00"

    response = client.put(
        "/api/purchases/batch",
        json={"items": [first, second]},
        headers=headers,
    )
    assert response.status_code == 422, response.text
    assert _finance(first_id)["supplier_name"] == "QA Supplier"
    assert _action_count("batch_update_purchases") == 0


def test_sales_batch_editor_round_trip_and_readonly_boundary(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    order_line_id, _ = _create_order(client, headers, "SALES-BATCH")
    rows_response = client.post(
        "/api/sales/batch-editor/rows",
        json={"order_line_ids": [order_line_id]},
        headers=headers,
    )
    assert rows_response.status_code == 200, rows_response.text
    editor = rows_response.json()
    assert editor["editable_from"] == "BP"
    assert editor["editable_through"] == "CM"
    assert editor["fixed_columns"] == ["B", "C", "D", "E", "F", "G", "M", "N", "O"]
    assert {
        column["key"] for column in editor["columns"] if column["editable"]
    } == {
        SALES_EDITOR_KEYS_BY_COLUMN[column_no]
        for column_no in SALES_EDITABLE_COLUMN_NUMBERS
    }
    assert not {"total_received", "accounts_receivable"} & {
        column["key"] for column in editor["columns"] if column["editable"]
    }

    item = _sales_editor_payload(editor, editor["rows"][0])
    item.update(
        {
            "contract_signed_date": "2026-07-21",
            "sales_contract_no": "SC-BATCH-001",
            "sales_contract_value": "1130.00",
            "sales_performance_period": "合同签订后30日内",
            "sales_unsigned_contract_amount": "0.00",
            "invoice_doc_no": "SID-BATCH-001",
            "invoice_date": "2026-07-22",
            "sales_invoice_no": "SINV-BATCH-001",
            "sales_invoice_amount": "600.00",
            "pending_invoice_amount": "530.00",
            "delivered_not_invoiced_amount": "0.00",
            "receipt1_date": "2026-07-23",
            "receipt1_notice_no": "RCPT1-BATCH-001",
            "receipt1_amount": "400.00",
            "receipt1_ratio": "0.353982",
            "receipt2_date": "2026-07-24",
            "receipt2_notice_no": "RCPT2-BATCH-001",
            "receipt2_amount": "300.00",
            "receipt2_ratio": "0.265487",
            "close_status": "进行中",
            "labor_cost": "12.34",
            "other_cost": "5.67",
        }
    )
    update_response = client.put(
        "/api/sales/batch",
        json={"items": [item]},
        headers=headers,
    )
    assert update_response.status_code == 200, update_response.text
    assert update_response.json()["updated"] == 1
    assert _action_count("batch_update_sales") == 1
    with db() as conn:
        audit_detail = json.loads(
            conn.execute(
                text(
                    """
                    SELECT detail
                    FROM operation_log
                    WHERE action_name = 'batch_update_sales'
                    ORDER BY id DESC
                    LIMIT 1
                    """
                )
            ).scalar_one()
        )
    assert audit_detail["summary"].startswith("在线表格批量修改 1 条销售信息，共 ")
    assert len(audit_detail["batch_entries"]) == 1
    assert audit_detail["batch_entries"][0]["before"]["sales_contract_no"] is None
    assert audit_detail["batch_entries"][0]["after"]["sales_contract_no"] == "SC-BATCH-001"

    refreshed_response = client.post(
        "/api/sales/batch-editor/rows",
        json={"order_line_ids": [order_line_id]},
        headers=headers,
    )
    assert refreshed_response.status_code == 200, refreshed_response.text
    refreshed = refreshed_response.json()
    values_by_key = {
        column["key"]: refreshed["rows"][0]["values"][index]
        for index, column in enumerate(refreshed["columns"])
        if column["key"]
    }
    assert values_by_key["sales_contract_no"] == "SC-BATCH-001"
    assert values_by_key["sales_invoice_no"] == "SINV-BATCH-001"
    assert values_by_key["receipt1_notice_no"] == "RCPT1-BATCH-001"
    assert values_by_key["receipt2_notice_no"] == "RCPT2-BATCH-001"
    assert _d(values_by_key["total_received"]) == Decimal("700.00")
    assert _d(values_by_key["accounts_receivable"]) == Decimal("430.00")
    assert values_by_key["close_status"] == "进行中"
    assert _d(values_by_key["labor_cost"]) == Decimal("12.34")
    assert _d(values_by_key["other_cost"]) == Decimal("5.67")

    unchanged_response = client.put(
        "/api/sales/batch",
        json={"items": [item]},
        headers=headers,
    )
    assert unchanged_response.status_code == 200, unchanged_response.text
    assert unchanged_response.json() == {"updated": 0, "order_line_ids": []}
    assert _action_count("batch_update_sales") == 1

    forbidden = {**item, "total_received": "1.00"}
    forbidden_response = client.put(
        "/api/sales/batch",
        json={"items": [forbidden]},
        headers=headers,
    )
    assert forbidden_response.status_code == 422, forbidden_response.text


def test_sales_batch_rejects_conflicting_shared_close_status_atomically(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    suffix = "SALES-BATCH-CLOSE-CONFLICT"
    first_id, _ = _create_order(client, headers, suffix, goods_name="设备甲")
    second_payload = _payload(suffix)
    second_payload["goods_name"] = "设备乙"
    second_response = client.post("/api/orders", json=second_payload, headers=headers)
    assert second_response.status_code == 200, second_response.text
    second_id = int(second_response.json()["order_line_id"])

    rows_response = client.post(
        "/api/sales/batch-editor/rows",
        json={"order_line_ids": [first_id, second_id]},
        headers=headers,
    )
    assert rows_response.status_code == 200, rows_response.text
    editor = rows_response.json()
    rows_by_id = {row["order_line_id"]: row for row in editor["rows"]}
    first = _sales_editor_payload(editor, rows_by_id[first_id])
    second = _sales_editor_payload(editor, rows_by_id[second_id])
    first.update({"sales_contract_no": "必须回滚", "close_status": "进行中"})
    second["close_status"] = "关闭"

    response = client.put(
        "/api/sales/batch",
        json={"items": [first, second]},
        headers=headers,
    )
    assert response.status_code == 422, response.text
    assert _finance(first_id)["sales_contract_no"] is None
    assert _finance(first_id)["close_status"] is None
    assert _action_count("batch_update_sales") == 0


def test_batch_basic_create_rejects_conflicting_shared_fields_atomically(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    first = _payload("BATCH-CREATE-CONFLICT")
    second = {
        **first,
        "account_manager": "另一个客户经理",
        "goods_name": "QA Equipment 2",
        "specification_model": "QA-SPEC-2",
    }
    conflict = client.post(
        "/api/orders/batch-basic",
        json={"items": [_basic_payload(first), _basic_payload(second)]},
        headers=headers,
    )
    assert conflict.status_code == 422, conflict.text
    with db() as conn:
        stored_lines = conn.execute(
            text(
                """
                SELECT COUNT(*)
                FROM order_line ol
                JOIN sales_order so ON so.id = ol.sales_order_id
                JOIN project p ON p.id = so.project_id
                WHERE p.project_code = :project_code
                """
            ),
            {"project_code": first["project_code"]},
        ).scalar_one()
    assert stored_lines == 0
    assert _action_count("batch_create_basic_orders") == 0


def test_batch_basic_update_is_atomic_and_rejects_conflicting_shared_order_fields(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    first_id, first = _create_order(client, headers, "BATCH-ATOMIC")
    second = {
        **first,
        "goods_name": "QA Equipment 2",
        "specification_model": "QA-SPEC-2",
    }
    second_response = client.post("/api/orders", json=second, headers=headers)
    assert second_response.status_code == 200, second_response.text
    second_id = int(second_response.json()["order_line_id"])

    first_update = {
        **_basic_payload(first),
        "order_line_id": first_id,
        "order_date": "2026-08-01",
    }
    second_update = {
        **_basic_payload(second),
        "order_line_id": second_id,
        "order_date": "2026-08-02",
    }
    conflict = client.put(
        "/api/orders/batch-basic",
        json={"items": [first_update, second_update]},
        headers=headers,
    )
    assert conflict.status_code == 422, conflict.text
    with db() as conn:
        stored_date = conn.execute(
            text("SELECT order_date FROM sales_order WHERE order_no = :order_no"),
            {"order_no": first["order_no"]},
        ).scalar_one()
    assert str(stored_date) == first["order_date"]
    assert _action_count("batch_update_basic_order") == 0


def test_purchase_summary_can_be_edited_from_purchase_detail(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    order_line_id, payload = _create_order(client, headers, "PURCHASE-SUMMARY")
    response = client.put(
        f"/api/purchases/{order_line_id}/summary",
        headers=headers,
        json={
            "supplier_name": "修改后的采购商",
            "purchase_tax_rate": "13.000000",
            "purchase_unit_price_no_tax": "72.000000",
            "purchase_unit_price": "81.360000",
            "cost_no_tax": "720.00",
            "purchase_amount": "813.60",
            "labor_cost": "12.34",
            "other_cost": "5.67",
        },
    )
    assert response.status_code == 200, response.text
    summary = response.json()["summary"]
    assert summary["supplier_name"] == "修改后的采购商"
    assert _d(summary["purchase_tax_rate"]) == Decimal("13.000000")
    assert _d(summary["purchase_unit_price_no_tax"]) == Decimal("72.000000")
    assert _d(summary["purchase_amount"]) == Decimal("813.60")
    assert _d(summary["labor_cost"]) == Decimal("12.34")
    assert _d(summary["other_cost"]) == Decimal("5.67")
    assert _d(summary["purchase_tax_amount"]) == Decimal("93.60")

    invalid_tax_rate = client.put(
        f"/api/purchases/{order_line_id}/summary",
        headers=headers,
        json={"purchase_tax_rate": "100.000001"},
    )
    assert invalid_tax_rate.status_code == 422, invalid_tax_rate.text

    with db() as conn:
        log = conn.execute(
            text(
                """
                SELECT detail
                FROM operation_log
                WHERE action_name = 'update_purchase_summary'
                ORDER BY id DESC
                LIMIT 1
                """
            )
        ).scalar_one()
    detail = json.loads(str(log))
    assert detail["after"]["project_code"] == payload["project_code"]
    assert detail["after"]["order_no"] == payload["order_no"]
    assert detail["after"]["goods_name"] == payload["goods_name"]
    assert detail["after"]["supplier_name"] == "修改后的采购商"


def _excel_import_file() -> bytes:
    template_path = ROOT_DIR / "backend" / "templates" / "市场部业务台账模板.xlsx"
    workbook = load_workbook(template_path)
    worksheet = workbook["Sheet1"]
    values = [
        "全额", "XL-IMPORT-001", "QA", "QA Branch", "QA Manager", date(2026, 7, 21), "商品销售", "常规",
        "QA Team", "QA Customer", "QA User", "QA Platform", "SO-XL-IMPORT-001", "Excel Import Project",
        "Excel Equipment", "XL-SPEC", "台", 2, Decimal("0.13"), Decimal("100.00"), Decimal("113.00"), Decimal("200.00"),
        Decimal("226.00"), "Excel Supplier", Decimal("0.13"), Decimal("70.00"), Decimal("79.10"), Decimal("140.00"),
        Decimal("158.20"), date(2026, 7, 22), 1, Decimal("100.00"), Decimal("113.00"), Decimal("70.00"),
        Decimal("79.10"), 1, Decimal("100.00"), Decimal("113.00"), "PC-XL-001", "验收后付款", "30天",
        Decimal("158.20"), Decimal("0.00"), date(2026, 7, 23), "PINV-XL-001", Decimal("158.20"),
        date(2026, 7, 24), "WH-XL-001", Decimal("158.20"), date(2026, 7, 25), "BOOK-XL-001",
        Decimal("158.20"), Decimal("0.00"), date(2026, 7, 25), date(2026, 7, 26), "PAY-XL-001",
        Decimal("100.00"), date(2026, 7, 27), "PAY-XL-002", Decimal("58.20"), Decimal("158.20"),
        Decimal("0.00"), Decimal("60.00"), Decimal("7.80"), Decimal("0.00"), Decimal("67.80"),
        Decimal("0.30"), date(2026, 7, 21), "SC-XL-001", Decimal("226.00"), "30天", Decimal("0.00"),
        "DOC-XL-001", date(2026, 7, 28), "SINV-XL-001", Decimal("226.00"), Decimal("0.00"),
        Decimal("0.00"), date(2026, 7, 29), "REC-XL-001", Decimal("100.00"), Decimal("44.2478"),
        date(2026, 7, 30), "REC-XL-002", Decimal("126.00"), Decimal("55.7522"), Decimal("226.00"),
        Decimal("0.00"), "进行中", Decimal("12.34"), Decimal("5.67"),
    ]
    assert len(values) == len(TEMPLATE_HEADERS) == 91
    for column, value in enumerate(values, start=1):
        worksheet.cell(3, column, value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize('spec', [None, '型号A'])
def test_import_same_goods_in_two_named_projects(client: TestClient, headers: dict[str, str], spec) -> None:
    workbook = load_workbook(BytesIO(_excel_import_file()))
    sheet = workbook.worksheets[0]
    sheet.cell(3, 14, '项目甲')
    sheet.cell(3, 16).value = spec
    values = [sheet.cell(3, column).value for column in range(1, 92)]
    values[13] = '项目乙'
    sheet.append(values)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = client.post('/api/orders/import-excel?filename=multi-project.xlsx',
                           content=output.getvalue(), headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()['success_rows'] == 2
    assert _count('sales_order') == 1
    assert _count('order_line') == 2
    repeated = client.post('/api/orders/import-excel?filename=multi-project.xlsx',
                           content=output.getvalue(), headers=headers)
    assert repeated.status_code == 422, repeated.text
    assert _count('order_line') == 2


def test_excel_template_import_export_round_trip(client: TestClient, headers: dict[str, str]) -> None:
    template_response = client.get("/api/orders/template", headers=headers)
    assert template_response.status_code == 200, template_response.text
    assert _action_count("download_order_template") == 1
    template = load_workbook(BytesIO(template_response.content), read_only=True, data_only=True)
    assert template["Sheet1"].cell(1, 1).value == "订单情况（王淼）"
    assert template["Sheet1"].cell(1, 39).value == "采购合同（周航）"
    assert [template["Sheet1"].cell(2, column).value for column in range(1, 92)] == TEMPLATE_HEADERS
    assert template["Sheet1"].cell(3, 2).value == SAMPLE_PROJECT_CODE
    assert template["Sheet1"].cell(3, 13).value == SAMPLE_ORDER_NO
    assert Decimal(str(template["Sheet1"].cell(3, 19).value)) == Decimal("0.13")
    assert template["Sheet1"].cell(3, 24).value == "示例采购厂商"
    assert Decimal(str(template["Sheet1"].cell(3, 25).value)) == Decimal("0.13")
    assert template["Sheet1"].cell(3, 90).value is None
    assert template["Sheet1"].cell(3, 91).value is None
    template.close()

    untouched_template = client.post(
        "/api/orders/import-excel?filename=市场部业务台账模板.xlsx",
        content=template_response.content,
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert untouched_template.status_code == 422, untouched_template.text
    assert "没有可导入的业务数据" in untouched_template.json()["detail"]
    assert _count("order_line") == 0

    partial_workbook = load_workbook(BytesIO(template_response.content))
    partial_workbook["Sheet1"].cell(3, 2, "XL-PARTIAL-EXAMPLE")
    partial_output = BytesIO()
    partial_workbook.save(partial_output)
    partial_workbook.close()
    partial_template = client.post(
        "/api/orders/import-excel?filename=市场部业务台账模板.xlsx",
        content=partial_output.getvalue(),
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert partial_template.status_code == 422, partial_template.text
    assert "仍包含示例占位内容" in partial_template.json()["detail"]
    assert _count("order_line") == 0

    content = _excel_import_file()
    import_response = client.post(
        "/api/orders/import-excel?filename=市场部业务台账模板.xlsx",
        content=content,
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert import_response.status_code == 200, import_response.text
    assert import_response.json()["success_rows"] == 1
    assert _count("order_line") == 1
    assert _count("finance_payment_entry") == 1
    assert _count("purchase_payment") == 2
    assert _count("sales_receipt") == 2
    with db() as conn:
        booked = conn.execute(text("SELECT booked_amount FROM finance_payment_entry")).scalar_one()
        purchase_info = conn.execute(
            text("SELECT supplier_name, labor_cost, other_cost FROM purchase_info")
        ).mappings().one()
    assert _d(booked) == Decimal("158.20")
    assert purchase_info["supplier_name"] == "Excel Supplier"
    assert _d(purchase_info["labor_cost"]) == Decimal("12.34")
    assert _d(purchase_info["other_cost"]) == Decimal("5.67")

    duplicate_response = client.post(
        "/api/orders/import-excel?filename=市场部业务台账模板.xlsx",
        content=content,
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert duplicate_response.status_code == 422, duplicate_response.text
    assert "整批回滚" in duplicate_response.json()["detail"]
    assert _count("order_line") == 1

    create_scoped_user = client.post(
        "/api/auth/users",
        json={
            "username": "excel-scope-user",
            "password": "Excel-Scope-Test-20260721!",
            "display_name": "Excel Scope User",
            "role_code": "order_entry",
            "permissions": ["order_entry"],
            "department_scope": ["OTHER"],
            "department_can_view": True,
            "department_can_entry": True,
        },
        headers=headers,
    )
    assert create_scoped_user.status_code == 200, create_scoped_user.text
    scoped_login = client.post(
        "/api/auth/login",
        json={"username": "excel-scope-user", "password": "Excel-Scope-Test-20260721!"},
    )
    assert scoped_login.status_code == 200, scoped_login.text
    scoped_headers = {"Authorization": f"Bearer {scoped_login.json()['access_token']}"}

    scoped_export_response = client.get("/api/orders/export", headers=scoped_headers)
    assert scoped_export_response.status_code == 200, scoped_export_response.text
    scoped_export = load_workbook(BytesIO(scoped_export_response.content), read_only=True, data_only=True)
    assert scoped_export["Sheet1"].max_row == 2
    scoped_export.close()

    forbidden_import = client.post(
        "/api/orders/import-excel?filename=市场部业务台账模板.xlsx",
        content=content,
        headers={**scoped_headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert forbidden_import.status_code == 403, forbidden_import.text
    assert "无权向部门" in forbidden_import.json()["detail"]
    assert _count("order_line") == 1

    export_response = client.get("/api/orders/export", headers=headers)
    assert export_response.status_code == 200, export_response.text
    assert _action_count("export_orders") == 2
    exported = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
    worksheet = exported["Sheet1"]
    assert [worksheet.cell(2, column).value for column in range(1, 92)] == TEMPLATE_HEADERS
    assert worksheet.cell(3, 2).value == "XL-IMPORT-001"
    assert worksheet.cell(3, 24).value == "Excel Supplier"
    assert Decimal(str(worksheet.cell(3, 19).value)) == Decimal("0.13")
    assert Decimal(str(worksheet.cell(3, 25).value)) == Decimal("0.13")
    assert Decimal(str(worksheet.cell(3, 52).value)) == Decimal("158.2")
    assert Decimal(str(worksheet.cell(3, 87).value)) == Decimal("226")
    assert Decimal(str(worksheet.cell(3, 90).value)) == Decimal("12.34")
    assert Decimal(str(worksheet.cell(3, 91).value)) == Decimal("5.67")
    exported.close()

    filtered_export_response = client.get(
        "/api/orders/export",
        params={
            "project_id": "XL-IMPORT",
            "department": "QA",
            "manager": "QA Manager",
            "client_unit": "QA Customer",
            "order_id": "SO-XL-IMPORT",
            "order_status": "closed",
            "supplier_name": "Excel Supplier",
            "start_date": "2026-07-21",
            "end_date": "2026-07-21",
            "invoice_start_date": "2026-07-28",
            "invoice_end_date": "2026-07-28",
        },
        headers=headers,
    )
    assert filtered_export_response.status_code == 200, filtered_export_response.text
    filtered_export = load_workbook(
        BytesIO(filtered_export_response.content),
        read_only=True,
        data_only=True,
    )
    assert filtered_export["Sheet1"].max_row == 3
    assert filtered_export["Sheet1"].cell(3, 2).value == "XL-IMPORT-001"
    filtered_export.close()

    empty_filtered_export_response = client.get(
        "/api/orders/export",
        params={"supplier_name": "不存在的采购厂商"},
        headers=headers,
    )
    assert empty_filtered_export_response.status_code == 200, empty_filtered_export_response.text
    empty_filtered_export = load_workbook(
        BytesIO(empty_filtered_export_response.content),
        read_only=True,
        data_only=True,
    )
    assert empty_filtered_export["Sheet1"].max_row == 2
    empty_filtered_export.close()

    invalid_date_range_response = client.get(
        "/api/orders/export",
        params={"start_date": "2026-07-22", "end_date": "2026-07-21"},
        headers=headers,
    )
    assert invalid_date_range_response.status_code == 400
    assert "开始日期不能晚于结束日期" in invalid_date_range_response.json()["detail"]


def test_ledger_export_filters_apply_to_exported_order_lines(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    shared_project = {
        "project_code": "QA-EXPORT-FILTER-PROJECT",
        "project_name": "Export Filter Project",
        "department": "EXPORT-QA",
        "branch_company": "Export Branch",
        "account_manager": "Export Manager",
        "customer_unit_name": "Export Customer",
    }
    first_line_id, _ = _create_order(
        client,
        headers,
        "EXPORT-FILTER-A",
        **shared_project,
        order_no="SO-EXPORT-FILTER-A",
        order_date="2026-08-01",
        supplier_name="Export Supplier A",
        goods_name="Export Goods A",
    )
    second_line_id, _ = _create_order(
        client,
        headers,
        "EXPORT-FILTER-B",
        **shared_project,
        order_no="SO-EXPORT-FILTER-B",
        order_date="2026-08-02",
        supplier_name="Export Supplier B",
        goods_name="Export Goods B",
    )

    for line_id, suffix, invoice_date in (
        (first_line_id, "A", "2026-08-10"),
        (second_line_id, "B", "2026-08-11"),
    ):
        response = client.post(
            f"/api/sales/{line_id}/invoices",
            json={
                "invoice_doc_no": f"DOC-EXPORT-{suffix}",
                "invoice_date": invoice_date,
                "invoice_no": f"INV-EXPORT-{suffix}",
                "invoice_amount": "1130.00",
                "pending_invoice_amount": "0.00",
                "delivered_not_invoiced_amount": "0.00",
            },
            headers=headers,
        )
        assert response.status_code == 200, response.text

    def exported_order_nos(params: dict[str, str]) -> list[str]:
        response = client.get("/api/orders/export", params=params, headers=headers)
        assert response.status_code == 200, response.text
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        worksheet = workbook["Sheet1"]
        values = [
            str(worksheet.cell(row, 13).value)
            for row in range(3, worksheet.max_row + 1)
            if worksheet.cell(row, 13).value is not None
        ]
        workbook.close()
        return values

    both_orders = ["SO-EXPORT-FILTER-B", "SO-EXPORT-FILTER-A"]
    project_level_cases = (
        {"project_id": "EXPORT-FILTER-PROJECT"},
        {"department": "EXPORT-QA"},
        {"manager": "Export Manager"},
        {"client_unit": "Export Customer"},
        {"order_status": "open"},
    )
    for params in project_level_cases:
        assert exported_order_nos(params) == both_orders

    row_level_cases = (
        ({"order_id": "FILTER-A"}, ["SO-EXPORT-FILTER-A"]),
        ({"start_date": "2026-08-01", "end_date": "2026-08-01"}, ["SO-EXPORT-FILTER-A"]),
        ({"supplier_name": "Supplier A"}, ["SO-EXPORT-FILTER-A"]),
        (
            {"invoice_start_date": "2026-08-10", "invoice_end_date": "2026-08-10"},
            ["SO-EXPORT-FILTER-A"],
        ),
        (
            {
                "project_id": "EXPORT-FILTER-PROJECT",
                "order_id": "FILTER-A",
                "supplier_name": "Supplier A",
                "invoice_start_date": "2026-08-10",
                "invoice_end_date": "2026-08-10",
            },
            ["SO-EXPORT-FILTER-A"],
        ),
        ({"order_id": "FILTER-A", "supplier_name": "Supplier B"}, []),
        ({"client_unit": "Missing Customer"}, []),
    )
    for params, expected in row_level_cases:
        assert exported_order_nos(params) == expected

    page_response = client.get(
        "/api/ledgers",
        params={"project_id": "EXPORT-FILTER-PROJECT", "order_id": "FILTER-A"},
        headers=headers,
    )
    assert page_response.status_code == 200, page_response.text
    assert page_response.json()["total"] == 1
    assert [item["project_code"] for item in page_response.json()["items"]] == [
        "QA-EXPORT-FILTER-PROJECT"
    ]


def test_release_multi_project_excel_round_trip_and_contamination_guard(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    input_path_value = os.getenv("RELEASE_QA_INPUT_XLSX")
    invalid_path_value = os.getenv("RELEASE_QA_INVALID_XLSX")
    output_dir_value = os.getenv("RELEASE_QA_OUTPUT_DIR")
    if not input_path_value or not invalid_path_value or not output_dir_value:
        pytest.skip("Release QA workbook paths are not configured.")

    input_path = Path(input_path_value)
    invalid_path = Path(invalid_path_value)
    output_dir = Path(output_dir_value)
    assert input_path.is_file()
    assert invalid_path.is_file()
    output_dir.mkdir(parents=True, exist_ok=True)

    business_tables = (
        "import_batch",
        "ledger_raw_row",
        "project",
        "sales_order",
        "order_line",
        "purchase_info",
        "delivery_record",
        "purchase_contract",
        "purchase_invoice",
        "warehouse_entry",
        "finance_invoice_check",
        "finance_payment_entry",
        "purchase_payment",
        "sales_contract",
        "sales_invoice",
        "sales_receipt",
    )

    def normalized(value: object) -> object:
        return json.loads(json.dumps(value, ensure_ascii=False, default=str))

    def business_snapshot() -> dict[str, list[dict]]:
        snapshot: dict[str, list[dict]] = {}
        with db() as conn:
            for table_name in business_tables:
                rows = conn.execute(text(f"SELECT * FROM `{table_name}` ORDER BY id")).mappings().all()
                snapshot[table_name] = normalized([dict(row) for row in rows])
        return snapshot

    def finance_snapshot(order_line_id: int) -> dict:
        with db() as conn:
            row = conn.execute(
                text("SELECT * FROM v_order_line_finance WHERE order_line_id = :id"),
                {"id": order_line_id},
            ).mappings().one()
        return normalized(dict(row))

    def editor_response(endpoint: str, order_line_ids: list[int]) -> dict:
        response = client.post(endpoint, json={"order_line_ids": order_line_ids}, headers=headers)
        assert response.status_code == 200, response.text
        return response.json()

    def editable_payload(editor: dict, row: dict) -> dict:
        payload = {"order_line_id": int(row["order_line_id"])}
        for index, column in enumerate(editor["columns"]):
            if column["editable"] and column["key"]:
                payload[column["key"]] = row["values"][index]
        return payload

    assert all(_count(table_name) == 0 for table_name in business_tables)
    workbook_bytes = input_path.read_bytes()
    import_response = client.post(
        "/api/orders/import-excel?filename=qa_multi_project_import.xlsx",
        content=workbook_bytes,
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert import_response.status_code == 200, import_response.text
    assert import_response.json()["success_rows"] == 4
    expected_counts = {
        "import_batch": 1,
        "ledger_raw_row": 4,
        "project": 2,
        "sales_order": 3,
        "order_line": 4,
        "purchase_info": 4,
        "delivery_record": 4,
        "purchase_contract": 4,
        "purchase_invoice": 4,
        "warehouse_entry": 4,
        "finance_invoice_check": 0,
        "finance_payment_entry": 4,
        "purchase_payment": 8,
        "sales_contract": 4,
        "sales_invoice": 4,
        "sales_receipt": 8,
    }
    actual_counts = {table_name: _count(table_name) for table_name in business_tables}
    assert actual_counts == expected_counts

    with db() as conn:
        imported_rows = conn.execute(
            text(
                """
                SELECT order_line_id, project_code, order_no, goods_name,
                       specification_model, supplier_name
                FROM v_order_line_finance
                ORDER BY order_line_id
                """
            )
        ).mappings().all()
    row_ids = {
        (str(row["project_code"]), str(row["order_no"]), str(row["goods_name"])): int(row["order_line_id"])
        for row in imported_rows
    }
    assert set(row_ids) == {
        ("QA-REL-001", "QA-SO-001", "边界路由器"),
        ("QA-REL-001", "QA-SO-001", "光模块"),
        ("QA-REL-001", "QA-SO-002", "维护服务"),
        ("QA-REL-002", "QA-SO-003", "零税率测试设备"),
    }
    first_id = row_ids[("QA-REL-001", "QA-SO-001", "边界路由器")]
    sibling_id = row_ids[("QA-REL-001", "QA-SO-001", "光模块")]
    third_id = row_ids[("QA-REL-001", "QA-SO-002", "维护服务")]
    sentinel_id = row_ids[("QA-REL-002", "QA-SO-003", "零税率测试设备")]
    sibling_before = finance_snapshot(sibling_id)
    sentinel_before = finance_snapshot(sentinel_id)

    basic_editor = editor_response("/api/orders/batch-editor/rows", [first_id, sibling_id])
    basic_rows = {int(row["order_line_id"]): row for row in basic_editor["rows"]}
    first_basic = editable_payload(basic_editor, basic_rows[first_id])
    sibling_basic = editable_payload(basic_editor, basic_rows[sibling_id])
    first_basic["goods_name"] = "边界路由器-已验收修改"
    basic_update = client.put(
        "/api/orders/batch-basic",
        json={
            "items": [
                first_basic,
                sibling_basic,
            ]
        },
        headers=headers,
    )
    assert basic_update.status_code == 200, basic_update.text
    assert basic_update.json() == {"updated": 1, "order_line_ids": [first_id]}
    basic_log_count = _action_count("batch_update_basic_order")
    basic_noop = client.put(
        "/api/orders/batch-basic",
        json={"items": [first_basic]},
        headers=headers,
    )
    assert basic_noop.status_code == 200, basic_noop.text
    assert basic_noop.json() == {"updated": 0, "order_line_ids": []}
    assert _action_count("batch_update_basic_order") == basic_log_count

    purchase_editor = editor_response("/api/purchases/batch-editor/rows", [first_id, sibling_id])
    purchase_rows = {int(row["order_line_id"]): row for row in purchase_editor["rows"]}
    first_purchase = editable_payload(purchase_editor, purchase_rows[first_id])
    sibling_purchase = editable_payload(purchase_editor, purchase_rows[sibling_id])
    first_purchase["supplier_name"] = "QA采购厂商甲-已验收修改"
    purchase_update = client.put(
        "/api/purchases/batch",
        json={
            "items": [
                first_purchase,
                sibling_purchase,
            ]
        },
        headers=headers,
    )
    assert purchase_update.status_code == 200, purchase_update.text
    assert purchase_update.json() == {"updated": 1, "order_line_ids": [first_id]}

    sales_editor = editor_response("/api/sales/batch-editor/rows", [third_id, sentinel_id])
    sales_rows = {int(row["order_line_id"]): row for row in sales_editor["rows"]}
    third_sales = editable_payload(sales_editor, sales_rows[third_id])
    sentinel_sales = editable_payload(sales_editor, sales_rows[sentinel_id])
    third_sales["sales_contract_no"] = "QA-SC-002-已验收修改"
    sales_update = client.put(
        "/api/sales/batch",
        json={
            "items": [
                third_sales,
                sentinel_sales,
            ]
        },
        headers=headers,
    )
    assert sales_update.status_code == 200, sales_update.text
    assert sales_update.json() == {"updated": 1, "order_line_ids": [third_id]}

    first_after = finance_snapshot(first_id)
    third_after = finance_snapshot(third_id)
    assert first_after["goods_name"] == "边界路由器-已验收修改"
    assert first_after["supplier_name"] == "QA采购厂商甲-已验收修改"
    assert third_after["sales_contract_no"] == "QA-SC-002-已验收修改"
    assert finance_snapshot(sibling_id) == sibling_before
    assert finance_snapshot(sentinel_id) == sentinel_before

    with db() as conn:
        audit_rows = conn.execute(
            text(
                """
                SELECT action_name, detail
                FROM operation_log
                WHERE action_name IN (
                  'batch_update_basic_order',
                  'batch_update_purchases',
                  'batch_update_sales'
                )
                ORDER BY id
                """
            )
        ).mappings().all()
    audit_changed_keys: dict[str, list[str]] = {}
    for row in audit_rows:
        detail = json.loads(str(row["detail"]))
        assert len(detail["batch_entries"]) == 1
        entry = detail["batch_entries"][0]
        changed_keys = sorted(
            key
            for key in set(entry["before"]) | set(entry["after"])
            if entry["before"].get(key) != entry["after"].get(key)
        )
        audit_changed_keys[str(row["action_name"])] = changed_keys
    assert audit_changed_keys["batch_update_basic_order"] == ["goods_name"]
    assert audit_changed_keys["batch_update_purchases"] == ["supplier_name"]
    assert audit_changed_keys["batch_update_sales"] == ["sales_contract_no"]

    before_invalid = business_snapshot()
    invalid_response = client.post(
        "/api/orders/import-excel?filename=qa_invalid_tax_rollback.xlsx",
        content=invalid_path.read_bytes(),
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert invalid_response.status_code == 422, invalid_response.text
    assert business_snapshot() == before_invalid

    duplicate_response = client.post(
        "/api/orders/import-excel?filename=qa_multi_project_import.xlsx",
        content=workbook_bytes,
        headers={**headers, "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    assert duplicate_response.status_code == 422, duplicate_response.text
    assert business_snapshot() == before_invalid

    before_export = business_snapshot()
    filtered_export = client.get(
        "/api/orders/export",
        params={"project_id": "QA-REL-001"},
        headers=headers,
    )
    assert filtered_export.status_code == 200, filtered_export.text
    (output_dir / "qa_filtered_export.xlsx").write_bytes(filtered_export.content)
    full_export = client.get("/api/orders/export", headers=headers)
    assert full_export.status_code == 200, full_export.text
    (output_dir / "qa_full_export.xlsx").write_bytes(full_export.content)
    assert business_snapshot() == before_export

    before_backup_restore = business_snapshot()
    backup_response = client.post("/api/backups", headers=headers)
    assert backup_response.status_code == 200, backup_response.text
    backup_id = int(backup_response.json()["id"])
    restore_probe_editor = editor_response("/api/orders/batch-editor/rows", [first_id])
    restore_probe = editable_payload(restore_probe_editor, restore_probe_editor["rows"][0])
    restore_probe["goods_name"] = "恢复演练临时值"
    restore_probe_response = client.put(
        "/api/orders/batch-basic",
        json={"items": [restore_probe]},
        headers=headers,
    )
    assert restore_probe_response.status_code == 200, restore_probe_response.text
    assert finance_snapshot(first_id)["goods_name"] == "恢复演练临时值"
    restore_response = client.post(f"/api/backups/{backup_id}/restore", headers=headers)
    assert restore_response.status_code == 200, restore_response.text
    assert restore_response.json()["restored"] is True
    assert business_snapshot() == before_backup_restore

    result = {
        "database": settings.mysql_database,
        "imported_rows": 4,
        "expected_counts": expected_counts,
        "actual_counts": actual_counts,
        "modified_order_line_ids": {
            "basic_and_purchase": first_id,
            "sales": third_id,
        },
        "unchanged_order_line_ids": [sibling_id, sentinel_id],
        "audit_changed_keys": audit_changed_keys,
        "invalid_import_status": invalid_response.status_code,
        "duplicate_import_status": duplicate_response.status_code,
        "filtered_export_file": "qa_filtered_export.xlsx",
        "full_export_file": "qa_full_export.xlsx",
        "business_tables_unchanged_by_export": True,
        "backup_restore_id": backup_id,
        "business_tables_restored_exactly": True,
    }
    (output_dir / "qa_api_database_results.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _admin_user() -> CurrentUser:
    with db() as conn:
        user_id = conn.execute(text("SELECT id FROM erp_user WHERE username = 'admin'")).scalar_one()
    return CurrentUser(
        id=int(user_id),
        username="admin",
        display_name="admin",
        role_code="admin",
        permissions=sorted(ALL_PERMISSIONS),
        department_scope=[],
        department_can_view=True,
        department_can_entry=True,
    )


def test_user_deactivation_and_permanent_deletion_lifecycle(
    client: TestClient,
    headers: dict[str, str],
) -> None:
    username = "lifecycle-user"
    password = "Lifecycle-Test-20260730!"
    created = client.post(
        "/api/auth/users",
        json={
            "username": username,
            "password": password,
            "display_name": "Lifecycle User",
            "role_code": "viewer",
        },
        headers=headers,
    )
    assert created.status_code == 200, created.text
    created_user = next(item for item in created.json()["items"] if item["username"] == username)
    user_id = int(created_user["id"])
    reset_password = "Lifecycle-Reset-20260730!"

    invalid_reset = client.post(
        f"/api/auth/users/{user_id}/reset-password",
        json={"password": "short"},
        headers=headers,
    )
    assert invalid_reset.status_code == 422

    reset = client.post(
        f"/api/auth/users/{user_id}/reset-password",
        json={"password": reset_password},
        headers=headers,
    )
    assert reset.status_code == 200, reset.text
    assert client.post(
        "/api/auth/login",
        json={"username": username, "password": password},
    ).status_code == 401
    assert client.post(
        "/api/auth/login",
        json={"username": username, "password": reset_password},
    ).status_code == 200
    with db() as conn:
        reset_log_detail = conn.execute(
            text(
                """
                SELECT detail
                FROM operation_log
                WHERE action_name = 'reset_user_password'
                ORDER BY id DESC
                LIMIT 1
                """
            )
        ).scalar_one()
    assert reset_password not in str(reset_log_detail)

    permanent_while_active = client.delete(
        f"/api/auth/users/{user_id}/permanent",
        headers=headers,
    )
    assert permanent_while_active.status_code == 400
    assert "先停用账号" in permanent_while_active.json()["detail"]

    login = client.post(
        "/api/auth/login",
        json={"username": username, "password": reset_password},
    )
    assert login.status_code == 200, login.text
    user_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert client.get("/api/orders/export", headers=user_headers).status_code == 200

    with db() as conn:
        conn.execute(
            text(
                """
                INSERT INTO backup_record
                  (file_name, backup_type, status, created_by)
                VALUES
                  ('lifecycle-test.json.gz', 'test', 'success', :user_id)
                """
            ),
            {"user_id": user_id},
        )

    deactivated = client.delete(f"/api/auth/users/{user_id}", headers=headers)
    assert deactivated.status_code == 200, deactivated.text
    assert all(item["username"] != username for item in deactivated.json()["items"])

    inactive = client.get(
        "/api/auth/users",
        params={"status": "inactive"},
        headers=headers,
    )
    assert inactive.status_code == 200, inactive.text
    assert any(item["username"] == username for item in inactive.json()["items"])
    assert client.post(
        "/api/auth/login",
        json={"username": username, "password": reset_password},
    ).status_code == 401

    reset_inactive = client.post(
        f"/api/auth/users/{user_id}/reset-password",
        json={"password": "Should-Not-Apply-20260730!"},
        headers=headers,
    )
    assert reset_inactive.status_code == 400
    assert "先恢复账号" in reset_inactive.json()["detail"]

    restored = client.post(
        f"/api/auth/users/{user_id}/restore",
        headers=headers,
    )
    assert restored.status_code == 200, restored.text
    assert any(item["username"] == username for item in restored.json()["items"])
    assert client.post(
        "/api/auth/login",
        json={"username": username, "password": reset_password},
    ).status_code == 200
    inactive_after_restore = client.get(
        "/api/auth/users",
        params={"status": "inactive"},
        headers=headers,
    )
    assert inactive_after_restore.status_code == 200, inactive_after_restore.text
    assert all(item["username"] != username for item in inactive_after_restore.json()["items"])

    already_active = client.post(
        f"/api/auth/users/{user_id}/restore",
        headers=headers,
    )
    assert already_active.status_code == 400
    assert "已启用" in already_active.json()["detail"]

    deactivated_again = client.delete(f"/api/auth/users/{user_id}", headers=headers)
    assert deactivated_again.status_code == 200, deactivated_again.text

    permanently_deleted = client.delete(
        f"/api/auth/users/{user_id}/permanent",
        headers=headers,
    )
    assert permanently_deleted.status_code == 200, permanently_deleted.text
    assert all(item["username"] != username for item in permanently_deleted.json()["items"])
    assert _count("erp_user", "id = :id", id=user_id) == 0
    with db() as conn:
        detached_logs = conn.execute(
            text(
                """
                SELECT COUNT(*)
                FROM operation_log
                WHERE user_id IS NULL AND user_name LIKE :user_name
                """
            ),
            {"user_name": f"%{username}%"},
        ).scalar()
        detached_backup_creator = conn.execute(
            text(
                """
                SELECT created_by
                FROM backup_record
                WHERE file_name = 'lifecycle-test.json.gz'
                """
            )
        ).scalar()
    assert int(detached_logs or 0) >= 1
    assert detached_backup_creator is None

    recreated = client.post(
        "/api/auth/users",
        json={
            "username": username,
            "password": password,
            "display_name": "Lifecycle User Recreated",
            "role_code": "viewer",
        },
        headers=headers,
    )
    assert recreated.status_code == 200, recreated.text


def test_tax_calculation_and_finance_entry_mapping(client: TestClient, headers: dict[str, str]) -> None:
    payload = _payload("TAX-MAP")
    payload.update(
        {
            "sales_tax_rate": "13.000000",
            "net_unit_price": "100.000000",
            "unit_price": "0.000000",
            "net_revenue": "0.00",
            "order_value": "0.00",
            "purchase_tax_rate": "13.000000",
            "purchase_unit_price_no_tax": "70.000000",
            "purchase_unit_price": "0.000000",
            "cost_no_tax": "0.00",
            "purchase_amount": "0.00",
            "labor_cost": "12.34",
            "other_cost": "5.67",
        }
    )
    created = client.post("/api/orders", json=payload, headers=headers)
    assert created.status_code == 200, created.text
    order_line_id = int(created.json()["order_line_id"])
    finance = _finance(order_line_id)
    assert _d(finance["sales_unit_price"]) == Decimal("113.000000")
    assert _d(finance["revenue_no_tax"]) == Decimal("1000.00")
    assert _d(finance["order_value"]) == Decimal("1130.00")
    assert _d(finance["sales_tax_amount"]) == Decimal("130.00")
    assert _d(finance["purchase_unit_price"]) == Decimal("79.100000")
    assert _d(finance["cost_no_tax"]) == Decimal("700.00")
    assert _d(finance["purchase_amount"]) == Decimal("791.00")
    assert _d(finance["purchase_tax_amount"]) == Decimal("91.00")
    assert _d(finance["labor_cost"]) == Decimal("12.34")
    assert _d(finance["other_cost"]) == Decimal("5.67")

    warehouse = client.post(
        f"/api/purchases/{order_line_id}/warehouse-entries",
        json={"warehouse_date": "2026-07-16", "voucher_no": "WH-TAX", "warehouse_amount": "791.00", "warehouse_amount_no_tax": "700.00"},
        headers=headers,
    )
    assert warehouse.status_code == 200, warehouse.text
    checked = client.post(
        f"/api/purchases/{order_line_id}/finance-invoice-checks",
        json={"received_invoice_date": "2026-07-17", "received_invoice_amount": "791.00", "voucher_code": "FI-TAX"},
        headers=headers,
    )
    assert checked.status_code == 200, checked.text
    paid = client.post(
        f"/api/purchases/{order_line_id}/finance-payments",
        json={"payment_date": "2026-07-18", "voucher_code": "FP-TAX", "booked_amount": "300.00"},
        headers=headers,
    )
    assert paid.status_code == 200, paid.text
    detail = paid.json()
    assert len(detail["warehouse_entries"]) == 1
    assert len(detail["finance_invoice_checks"]) == 1
    assert len(detail["finance_payments"]) == 1
    assert _d(detail["summary"]["total_finance_checked"]) == Decimal("791.00")
    assert _d(detail["summary"]["total_finance_paid"]) == Decimal("300.00")
    assert _d(detail["summary"]["financial_accounts_payable"]) == Decimal("491.00")


def test_order_user_change_audit_contains_chinese_context(client: TestClient, headers: dict[str, str]) -> None:
    order_line_id, payload = _create_order(client, headers, "AUDIT-USER")
    logs_before_query = _count("operation_log")
    query = client.get("/api/orders", params={"project_id": payload["project_code"]}, headers=headers)
    assert query.status_code == 200, query.text
    detail_query = client.get(f"/api/purchases/{order_line_id}", headers=headers)
    assert detail_query.status_code == 200, detail_query.text
    assert _count("operation_log") == logs_before_query

    updated = {**payload, "user_name": "修改后的用户"}
    response = client.put(f"/api/orders/{order_line_id}", json=updated, headers=headers)
    assert response.status_code == 200, response.text

    with db() as conn:
        row = conn.execute(
            text(
                """
                SELECT user_name, detail, status
                FROM operation_log
                WHERE action_name = 'update_order'
                ORDER BY id DESC LIMIT 1
                """
            )
        ).mappings().one()
    detail = json.loads(str(row["detail"]))
    assert row["user_name"] == "系统管理员（账号：admin）"
    assert row["status"] == "success"
    assert detail["before"]["end_user_name"] == "QA User"
    assert detail["after"]["end_user_name"] == "修改后的用户"
    assert detail["after"]["project_code"] == payload["project_code"]
    assert detail["after"]["order_no"] == payload["order_no"]
    assert detail["after"]["goods_name"] == payload["goods_name"]


def test_failed_data_change_is_logged_but_login_is_not(client: TestClient, headers: dict[str, str]) -> None:
    logs_before_login = _count("operation_log")
    login = client.post("/api/auth/login", json={"username": "admin", "password": TEST_PASSWORD})
    assert login.status_code == 200, login.text
    assert _count("operation_log") == logs_before_login

    invalid = client.post("/api/orders", json={"project_code": "缺少必填字段"}, headers=headers)
    assert invalid.status_code == 422, invalid.text
    with db() as conn:
        row = conn.execute(
            text(
                """
                SELECT user_name, module_name, detail, status
                FROM operation_log
                WHERE action_name = 'create_failed'
                ORDER BY id DESC LIMIT 1
                """
            )
        ).mappings().one()
    detail = json.loads(str(row["detail"]))
    assert row["user_name"] == "系统管理员（账号：admin）"
    assert row["module_name"] == "订单管理"
    assert row["status"] == "failed"
    assert detail["summary"] == "新增数据失败：字段校验未通过"


@pytest.mark.parametrize("case", [f"N-{index:02d}" for index in range(1, 13)])
def test_normal_financial_cases(case: str, client: TestClient, headers: dict[str, str]) -> None:
    order_line_id, payload = _create_order(client, headers, case)

    if case == "N-01":
        finance = _finance(order_line_id)
        assert _d(finance["gross_profit_no_tax"]) == Decimal("300.00")
        assert _d(finance["gross_profit"]) == Decimal("339.00")
        assert _d(finance["accounts_receivable"]) == Decimal("1130.00")
        assert _d(finance["accounts_payable"]) == Decimal("791.00")
        assert _count("project") == _count("sales_order") == _count("order_line") == 1
        assert _count("purchase_info", "order_line_id = :id", id=order_line_id) == 1
        assert _count("delivery_record", "order_line_id = :id", id=order_line_id) == 1
        assert _d(_ledger(payload["project_code"])["order_amount"]) == Decimal("1130.00")
        assert _dashboard(client, headers) == {
            "orderAmount": 1130.0, "grossProfit": 339.0, "orderCount": 1,
            "accountsReceivable": 1130.0, "accountsPayable": 791.0, "closedCount": 0,
        }
        assert _action_count("create_order") == 1

    elif case == "N-02":
        response = client.post(
            f"/api/purchases/{order_line_id}/contracts",
            json={"purchase_contract_no": "PC-N02", "signed_amount": "791.00", "unsigned_amount": "0.00"},
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert _d(response.json()["summary"]["purchase_contract_signed_amount"]) == Decimal("791.00")
        assert _d(_finance(order_line_id)["accounts_payable"]) == Decimal("791.00")
        assert _action_count("create_purchase_contract") == 1

    elif case == "N-03":
        response = client.post(
            f"/api/purchases/{order_line_id}/invoices",
            json={"received_invoice_date": "2026-07-15", "invoice_no": "PI-N03", "invoice_amount": "791.00"},
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert response.json()["invoices"][0]["phase_no"] == 1
        assert _d(response.json()["invoices"][0]["invoice_amount"]) == Decimal("791.00")
        assert _d(_finance(order_line_id)["accounts_payable"]) == Decimal("791.00")
        assert _action_count("create_purchase_invoice") == 1

    elif case == "N-04":
        response = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("300.00"), headers=headers)
        assert response.status_code == 200, response.text
        assert response.json()["payments"][0]["phase_no"] == 1
        assert response.json()["payments"][0]["due_payment_date"] == "2026-07-20"
        finance = _finance(order_line_id)
        assert _d(finance["total_paid"]) == Decimal("300.00")
        assert _d(finance["accounts_payable"]) == Decimal("491.00")
        assert _dashboard(client, headers)["accountsPayable"] == 491.0

    elif case == "N-05":
        client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("300.00"), headers=headers)
        response = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("491.00", "2026-07-25"), headers=headers)
        assert response.status_code == 200, response.text
        assert response.json()["payments"][1]["phase_no"] == 2
        assert response.json()["payments"][1]["due_payment_date"] is None
        finance = _finance(order_line_id)
        assert _d(finance["total_paid"]) == Decimal("791.00")
        assert _d(finance["accounts_payable"]) == Decimal("0.00")

    elif case == "N-06":
        response = client.post(
            f"/api/sales/{order_line_id}/contracts",
            json={"sales_contract_no": "SC-N06", "contract_value": "1130.00", "unsigned_contract_amount": "0.00"},
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert _d(response.json()["summary"]["sales_contract_value"]) == Decimal("1130.00")
        assert _d(_finance(order_line_id)["accounts_receivable"]) == Decimal("1130.00")
        assert _action_count("create_sales_contract") == 1

    elif case == "N-07":
        response = client.post(
            f"/api/sales/{order_line_id}/invoices",
            json={
                "invoice_doc_no": "SID-N07", "invoice_date": "2026-07-20", "invoice_no": "SI-N07",
                "invoice_amount": "565.00", "pending_invoice_amount": "565.00",
                "delivered_not_invoiced_amount": "0.00",
            },
            headers=headers,
        )
        assert response.status_code == 200, response.text
        assert response.json()["invoices"][0]["phase_no"] == 1
        assert _d(_finance(order_line_id)["sales_invoice_amount"]) == Decimal("565.00")
        assert _d(_finance(order_line_id)["accounts_receivable"]) == Decimal("1130.00")
        sales_rows = client.get("/api/sales", headers=headers)
        assert sales_rows.status_code == 200, sales_rows.text
        sales_row = next(item for item in sales_rows.json()["items"] if item["order_line_id"] == order_line_id)
        assert sales_row["invoice_dates"] == "2026-07-20"

    elif case == "N-08":
        response = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("400.00", "35.398230"), headers=headers)
        assert response.status_code == 200, response.text
        finance = _finance(order_line_id)
        assert _d(finance["total_received"]) == Decimal("400.00")
        assert _d(finance["accounts_receivable"]) == Decimal("730.00")
        assert _dashboard(client, headers)["closedCount"] == 0

    elif case == "N-09":
        client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("400.00"), headers=headers)
        response = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("730.00"), headers=headers)
        assert response.status_code == 200, response.text
        finance = _finance(order_line_id)
        assert _d(finance["total_received"]) == Decimal("1130.00")
        assert _d(finance["accounts_receivable"]) == Decimal("0.00")
        assert _ledger(payload["project_code"])["computed_close_status"] == "closed"
        assert _dashboard(client, headers)["closedCount"] == 1

    elif case == "N-10":
        first = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("300.00"), headers=headers)
        client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("491.00"), headers=headers)
        payment_id = first.json()["payments"][0]["id"]
        response = client.put(f"/api/purchases/payments/{payment_id}", json=_payment("250.00"), headers=headers)
        assert response.status_code == 200, response.text
        assert _count("purchase_payment", "order_line_id = :id", id=order_line_id) == 2
        finance = _finance(order_line_id)
        assert _d(finance["total_paid"]) == Decimal("741.00")
        assert _d(finance["accounts_payable"]) == Decimal("50.00")
        with db() as conn:
            detail = conn.execute(text("SELECT detail FROM operation_log WHERE action_name = 'update_purchase_payment'")).scalar_one()
        assert '"300.00"' in detail and '"250.00"' in detail

    elif case == "N-11":
        client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("400.00"), headers=headers)
        second = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("730.00"), headers=headers)
        receipt_id = second.json()["receipts"][1]["id"]
        response = client.delete(f"/api/sales/receipts/{receipt_id}", headers=headers)
        assert response.status_code == 200, response.text
        assert _count("sales_receipt", "id = :id AND deleted_at IS NOT NULL", id=receipt_id) == 1
        finance = _finance(order_line_id)
        assert _d(finance["total_received"]) == Decimal("400.00")
        assert _d(finance["accounts_receivable"]) == Decimal("730.00")
        assert _dashboard(client, headers)["closedCount"] == 0

    elif case == "N-12":
        response = client.delete(f"/api/orders/{order_line_id}", headers=headers)
        assert response.status_code == 200, response.text
        assert _count("order_line", "id = :id AND deleted_at IS NOT NULL", id=order_line_id) == 1
        assert _count("sales_order", "deleted_at IS NOT NULL") == 1
        assert _count("v_order_line_finance", "order_line_id = :id", id=order_line_id) == 0
        listed = client.get("/api/orders", params={"project_id": payload["project_code"]}, headers=headers)
        assert listed.status_code == 200
        assert listed.json()["total"] == 0
        assert _action_count("delete_order") == 1


@pytest.mark.parametrize("case", [f"B-{index:02d}" for index in range(1, 9)])
def test_boundary_financial_cases(case: str, client: TestClient, headers: dict[str, str]) -> None:
    if case == "B-01":
        maximum = "9999999999999999.99"
        order_line_id, payload = _create_order(client, headers, case, order_value=maximum, purchase_amount="0.00")
        assert _d(_finance(order_line_id)["order_value"]) == Decimal(maximum)
        response = client.get("/api/orders", params={"project_id": payload["project_code"]}, headers=headers)
        assert response.status_code == 200, response.text
        assert _d(response.json()["items"][0]["order_value"]) == Decimal(maximum)

    elif case == "B-02":
        quantity = "999999999999.999999"
        order_line_id, _ = _create_order(
            client, headers, case, quantity=quantity, net_unit_price=quantity, unit_price=quantity,
            order_value="9999999999999999.99", purchase_amount="0.00",
        )
        with db() as conn:
            stored = conn.execute(text("SELECT quantity FROM order_line WHERE id = :id"), {"id": order_line_id}).scalar_one()
        assert _d(stored) == Decimal(quantity)

    elif case == "B-03":
        order_line_id, _ = _create_order(client, headers, case, quantity="0.000000", delivery_quantity="0.000000")
        response = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("0.00", "100.000000"), headers=headers)
        assert response.status_code == 200, response.text
        finance = _finance(order_line_id)
        assert _d(finance["total_received"]) == Decimal("0.00")
        assert _d(finance["accounts_receivable"]) == Decimal("1130.00")
        assert _dashboard(client, headers)["closedCount"] == 0

    elif case == "B-04":
        _, payload = _create_order(client, headers, case, order_date="2099-12-31")
        logs_before = _count("operation_log")
        for limit in (1, 500):
            response = client.get("/api/orders", params={"project_id": payload["project_code"], "limit": limit, "offset": 0}, headers=headers)
            assert response.status_code == 200, response.text
            assert len(response.json()["items"]) <= limit
        assert _count("operation_log") == logs_before

    elif case == "B-05":
        order_line_id, _ = _create_order(client, headers, case)
        client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("100.00"), headers=headers)
        second = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("100.00"), headers=headers)
        payment_id = second.json()["payments"][1]["id"]
        assert client.delete(f"/api/purchases/payments/{payment_id}", headers=headers).status_code == 200
        third = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("100.00"), headers=headers)
        assert [item["phase_no"] for item in third.json()["payments"]] == [1, 3]
        assert _d(_finance(order_line_id)["total_paid"]) == Decimal("200.00")

    elif case == "B-06":
        order_line_id, _ = _create_order(client, headers, case)
        record = purchases.PurchasePaymentCreate(payment_amount=Decimal("1.00"), payment_date=date(2026, 7, 20))
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = [pool.submit(purchases.add_purchase_payment, order_line_id, record, _admin_user()) for _ in range(2)]
            [result.result() for result in results]
        with db() as conn:
            phases = conn.execute(
                text("SELECT phase_no FROM purchase_payment WHERE order_line_id = :id AND deleted_at IS NULL ORDER BY phase_no"),
                {"id": order_line_id},
            ).scalars().all()
        assert phases == [1, 2]
        assert _d(_finance(order_line_id)["total_paid"]) == Decimal("2.00")

    elif case == "B-07":
        order_line_id, _ = _create_order(client, headers, case)
        client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("1129.99"), headers=headers)
        client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("0.01"), headers=headers)
        client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("790.99"), headers=headers)
        client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("0.01"), headers=headers)
        finance = _finance(order_line_id)
        assert _d(finance["accounts_receivable"]) == Decimal("0.00")
        assert _d(finance["accounts_payable"]) == Decimal("0.00")
        assert _dashboard(client, headers)["closedCount"] == 1

    elif case == "B-08":
        order_line_id, _ = _create_order(client, headers, case)
        assert client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("1130.00"), headers=headers).status_code == 200
        assert client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("791.00"), headers=headers).status_code == 200
        receipt_before = _count("sales_receipt", "order_line_id = :id", id=order_line_id)
        payment_before = _count("purchase_payment", "order_line_id = :id", id=order_line_id)
        receipt_response = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("0.01"), headers=headers)
        payment_response = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("0.01"), headers=headers)
        receipt_after = _count("sales_receipt", "order_line_id = :id", id=order_line_id)
        payment_after = _count("purchase_payment", "order_line_id = :id", id=order_line_id)
        finance = _finance(order_line_id)
        assert (receipt_response.status_code, payment_response.status_code) == (422, 422), {
            "receipt_count_before_after": (receipt_before, receipt_after),
            "payment_count_before_after": (payment_before, payment_after),
            "accounts_receivable": str(finance["accounts_receivable"]),
            "accounts_payable": str(finance["accounts_payable"]),
        }
        assert (receipt_after, payment_after) == (receipt_before, payment_before)
        assert _d(finance["accounts_receivable"]) == Decimal("0.00")
        assert _d(finance["accounts_payable"]) == Decimal("0.00")


@pytest.mark.parametrize("case", [f"I-{index:02d}" for index in range(1, 13)])
def test_invalid_input_cases(case: str, client: TestClient, headers: dict[str, str]) -> None:
    if case == "I-01":
        order_payload = {"project_code": "QA-I01", "order_no": "SO-I01", "goods_name": "x", "customer_unit_name": "y", "order_value": "-0.01"}
        assert client.post("/api/orders", json=order_payload, headers=headers).status_code == 422
        order_line_id, _ = _create_order(client, headers, case)
        assert client.post(f"/api/purchases/{order_line_id}/payments", json=_payment("-1.00"), headers=headers).status_code == 422
        assert client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("-1.00"), headers=headers).status_code == 422
        assert _count("purchase_payment") == _count("sales_receipt") == 0
        assert _d(_finance(order_line_id)["accounts_payable"]) == Decimal("791.00")

    elif case == "I-02":
        order_line_id, _ = _create_order(client, headers, case)
        for amount in ("100.001", "99999999999999999.99", "NaN", "Infinity", "not-a-number"):
            response = client.post(f"/api/purchases/{order_line_id}/payments", json=_payment(amount), headers=headers)
            assert response.status_code == 422, response.text
        assert _count("purchase_payment") == 0
        assert _d(_finance(order_line_id)["accounts_payable"]) == Decimal("791.00")

    elif case == "I-03":
        payload = _payload(case)
        payload["quantity"] = "1.0000001"
        response = client.post("/api/orders", json=payload, headers=headers)
        assert response.status_code == 422, response.text
        assert _count("project") == _count("sales_order") == _count("order_line") == 0

    elif case == "I-04":
        for bad_date in ("not-a-date", "2026-02-30", "2100-01-01"):
            payload = _payload(f"{case}-{bad_date[:4]}")
            payload["order_date"] = bad_date
            response = client.post("/api/orders", json=payload, headers=headers)
            assert response.status_code == 422, response.text
        assert _count("order_line") == 0
        assert _action_count("create_order") == 0

    elif case == "I-05":
        order_line_id, _ = _create_order(client, headers, case)
        for ratio in ("-0.000001", "100.000001", "1.0000001"):
            response = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("1.00", ratio), headers=headers)
            assert response.status_code == 422, response.text
        assert _count("sales_receipt") == 0

    elif case == "I-06":
        for field in ("project_code", "order_no", "goods_name", "customer_unit_name", "order_value"):
            payload = _payload(f"{case}-{field}")
            payload.pop(field)
            response = client.post("/api/orders", json=payload, headers=headers)
            assert response.status_code == 422, response.text
        assert _count("project") == _count("sales_order") == _count("order_line") == 0

    elif case == "I-07":
        # 判重键含数量与单价：完全一致的明细拒绝，数量或单价不同的视为不同明细。
        order_line_id, payload = _create_order(client, headers, case)
        response = client.post("/api/orders", json=payload, headers=headers)
        assert response.status_code == 409, response.text
        with db() as conn:
            quantity = conn.execute(text("SELECT quantity FROM order_line WHERE id = :id"), {"id": order_line_id}).scalar_one()
        assert _d(quantity) == Decimal("10.0000")
        assert _count("order_line") == 1
        for field, value in (("quantity", "99.000000"), ("unit_price", "120.000000")):
            response = client.post("/api/orders", json={**payload, field: value}, headers=headers)
            assert response.status_code == 200, response.text
        assert _count("order_line") == 3
        assert _action_count("create_order") == 3

    elif case == "I-08":
        valid = _payload("I08-valid")
        invalid = _payload("I08-invalid")
        invalid["order_value"] = "-1.00"
        response = client.post("/api/orders/batch", json={"items": [valid, invalid]}, headers=headers)
        assert response.status_code == 422, response.text
        assert _count("project") == _count("sales_order") == _count("order_line") == 0
        assert _action_count("batch_create_orders") == 0

    elif case == "I-09":
        response = client.post("/api/sales/999999/receipts", json=_receipt("1.00"), headers=headers)
        assert response.status_code == 404, response.text
        order_line_id, _ = _create_order(client, headers, case)
        created = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("1.00"), headers=headers)
        receipt_id = created.json()["receipts"][0]["id"]
        assert client.delete(f"/api/sales/receipts/{receipt_id}", headers=headers).status_code == 200
        assert client.put(f"/api/sales/receipts/{receipt_id}", json=_receipt("2.00"), headers=headers).status_code == 404
        assert _count("sales_receipt", "id = :id AND deleted_at IS NOT NULL", id=receipt_id) == 1

    elif case == "I-10":
        order_line_id, _ = _create_order(client, headers, case)
        created = client.post(
            "/api/auth/users",
            json={"username": "viewer-i10", "password": "Viewer-Test-20260714!", "display_name": "Viewer", "role_code": "viewer"},
            headers=headers,
        )
        assert created.status_code == 200, created.text
        login = client.post("/api/auth/login", json={"username": "viewer-i10", "password": "Viewer-Test-20260714!"})
        viewer_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        response = client.post(f"/api/sales/{order_line_id}/receipts", json=_receipt("1.00"), headers=viewer_headers)
        assert response.status_code == 403, response.text
        assert _count("sales_receipt") == 0
        assert _d(_finance(order_line_id)["accounts_receivable"]) == Decimal("1130.00")

    elif case == "I-11":
        _create_order(client, headers, case)
        logs_before = _count("operation_log")
        for params in ({"limit": 0}, {"limit": 501}, {"offset": -1}):
            response = client.get("/api/orders", params=params, headers=headers)
            assert response.status_code == 422, response.text
        assert _count("operation_log") == logs_before

    elif case == "I-12":
        payload = _payload(case)
        payload["project_code"] = "P" * 65
        response = client.post("/api/orders", json=payload, headers=headers)
        assert _count("project") == _count("sales_order") == _count("order_line") == 0
        assert response.status_code == 422, response.text
