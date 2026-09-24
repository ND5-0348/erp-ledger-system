from io import BytesIO
from decimal import Decimal
from openpyxl import load_workbook
from sqlalchemy import text
from app.db import db
from test_source_import import source_file
from test_template_0916 import workbook_bytes
from test_import_permissions import _create_user, _login


def upload(client,headers,content,preview=False,token=None):
    return client.post(f'/api/orders/source-import?filename=source.xlsx&preview={str(preview).lower()}',content=content,
                       headers={**headers,**({'X-Duplicate-Confirmation':token} if token else {})})


def modified(content, project=None):
    wb=load_workbook(BytesIO(content));wb.properties.title='resaved source'
    if project:
        for no in (3,4):wb.active.cell(no,2).value=project
    return workbook_bytes(wb)


def test_preview_totals_and_resaved_duplicates_require_bound_confirmation(client,headers):
    content=source_file()
    pre=upload(client,headers,content,True)
    assert pre.status_code==200,pre.text
    assert pre.json()['layout']=='0916 新版'
    assert pre.json()['summary']['line_count']==2
    assert Decimal(pre.json()['summary']['invoice_amount'])==400
    first=upload(client,headers,content)
    assert first.status_code==200,first.text
    newer=modified(content)
    pre=upload(client,headers,newer,True)
    assert pre.status_code==200,pre.text
    duplicate=pre.json()['duplicates']
    assert duplicate['count']==2 and duplicate['rows'][0]['matches'][0]['import_batch_id']==first.json()['batch_id']
    assert upload(client,headers,newer).status_code==409
    assert upload(client,headers,newer,token='x'*64).status_code==409
    done=upload(client,headers,newer,token=duplicate['confirmation_token'])
    assert done.status_code==200,done.text
    with db() as conn: assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one()==4


def test_identical_rows_within_file_are_kept_separately(client,headers):
    wb=load_workbook(BytesIO(source_file()))
    for col in range(1,93):wb.active.cell(4,col).value=wb.active.cell(3,col).value
    result=upload(client,headers,workbook_bytes(wb))
    assert result.status_code==200,result.text
    assert result.json()['success_rows']==2
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM ledger_raw_row')).scalar_one()==2


def test_duplicate_review_amount_uses_same_rounding_as_import(client,headers):
    wb=load_workbook(BytesIO(source_file()))
    for row in (3,4):wb.active.cell(row,23).value=8.459999999999999
    content=workbook_bytes(wb)
    assert upload(client,headers,content).status_code==200
    preview=upload(client,headers,modified(content),True)
    assert preview.status_code==200,preview.text
    for row in preview.json()['duplicates']['rows']:
        assert row['order_amount']=='8.46'
        assert all(Decimal(m['order_value'])==Decimal('8.46') for m in row['matches'])


def test_undo_blocks_later_additions_to_same_project(client,headers):
    content=source_file()
    first=upload(client,headers,content).json()['batch_id']
    undo=client.get(f'/api/import-batches/{first}',headers=headers).json()['undo']
    assert undo['can_revert']
    wb=load_workbook(BytesIO(content))
    for row in (3,4):wb.active.cell(row,13).value='NEW-ORDER'
    added=upload(client,headers,workbook_bytes(wb))
    assert added.status_code==200,added.text
    assert not client.get(f'/api/import-batches/{first}',headers=headers).json()['undo']['can_revert']
    reverted=client.post(f'/api/import-batches/{first}/revert',headers=headers,json={'token':undo['token'],'confirmation':f'撤销批次 {first}'})
    assert reverted.status_code==409
    assert client.get('/api/orders',headers=headers).json()['total']==4


def test_undo_preserves_other_batch_and_accounts_and_blocks_replay(client,headers):
    content=source_file();first=upload(client,headers,content).json()['batch_id']
    other=upload(client,headers,modified(content,'OTHER-PROJECT'))
    assert other.status_code==200,other.text
    detail=client.get(f'/api/import-batches/{first}',headers=headers)
    assert detail.status_code==200,detail.text
    undo=detail.json()['undo']; assert undo['can_revert'],undo
    assert detail.json()['backup_file'] and not detail.json()['summary_is_current']
    result=client.post(f'/api/import-batches/{first}/revert',json={'token':undo['token'],'confirmation':f'撤销批次 {first}'},headers=headers)
    assert result.status_code==200,result.text
    assert client.get('/api/orders',headers=headers).json()['total']==2
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one()==4
        assert conn.execute(text("SELECT status FROM import_batch WHERE id=:id"),{'id':first}).scalar_one()=='reverted'
        assert conn.execute(text('SELECT COUNT(*) FROM erp_user')).scalar_one()>=1
    assert upload(client,headers,content).status_code==409
    assert client.get(f'/api/import-batches/{first}',headers=headers).json()['undo']['can_revert'] is False


def test_undo_rechecks_updates_and_legacy_batches_have_no_invented_baseline(client,headers):
    batch=upload(client,headers,source_file()).json()['batch_id']
    undo=client.get(f'/api/import-batches/{batch}',headers=headers).json()['undo']
    assert undo['can_revert']
    with db() as conn:conn.execute(text("UPDATE order_line SET goods_name='changed after preview' ORDER BY id LIMIT 1"))
    response=client.post(f'/api/import-batches/{batch}/revert',json={'token':undo['token'],'confirmation':f'撤销批次 {batch}'},headers=headers)
    assert response.status_code==409
    assert client.get('/api/orders',headers=headers).json()['total']==2
    with db() as conn:conn.execute(text('UPDATE import_batch SET baseline_sha256=NULL,review_json=NULL WHERE id=:id'),{'id':batch})
    detail=client.get(f'/api/import-batches/{batch}',headers=headers).json()
    assert detail['summary_is_current'] and not detail['undo']['can_revert']


def test_batch_history_and_undo_permissions(client,headers):
    batch=upload(client,headers,source_file()).json()['batch_id']
    assert _create_user(client,headers,'batch_reader',['ledger_import'],department_scope=['第二部门'],department_all=False).status_code==200
    scoped=_login(client,'batch_reader')
    assert client.get('/api/import-batches',headers=scoped).json()['total']==0
    assert client.get(f'/api/import-batches/{batch}',headers=scoped).status_code==404
    assert client.post(f'/api/import-batches/{batch}/revert',headers=scoped,json={'token':'x'*64,'confirmation':f'撤销批次 {batch}'}).status_code==403
