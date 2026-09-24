"""H3 page contract: keep original values and return revalidated display data."""
from test_legacy_import_flow import _preview, _row


def test_read_preview_shows_effective_values_after_resolution(client, headers):
    response = _preview(client, headers, [_row({
        13: "SO-OLD/SO-NEW", 74: "2026/01/01/2026/02/01", 76: "300",
        55: "2026-03-01", 57: "40",
    })])
    session_id = response.json()["session_id"]
    before = client.get(f"/api/orders/import-preview/{session_id}", headers=headers).json()
    assert any(issue["code"] == "AMOUNT_SPLIT_REQUIRED" for issue in before["rows"]["items"][0]["effective"]["issues"])
    updated = client.put(f"/api/orders/import-preview/{session_id}/resolutions", headers=headers, json={"items": [{
        "excel_row_no": 3, "resolution": {"finance": {"sales_invoice": {"phases": [
            {"source_group": 1, "date": "2026-01-01", "amount": "100", "document_no": "INV-1"},
            {"source_group": 1, "date": "2026-02-01", "amount": "200", "document_no": "INV-2"},
        ]}}},
    }]})
    assert updated.status_code == 200, updated.text
    after = client.get(f"/api/orders/import-preview/{session_id}", headers=headers).json()
    row = after["rows"]["items"][0]
    assert after["summary"]["comparable"]
    assert row["parsed"]["finance"]["sales_invoice"]["amount_raw"] == "300"
    assert not row["effective"]["issues"]
    assert row["effective"]["order_history"] == ["SO-OLD", "SO-NEW"]
    assert [phase["amount"] for phase in row["effective"]["finance"]["sales_invoice"]["phases"]] == ["100", "200"]
    assert row["effective"]["finance"]["purchase_payment"]["phases"][0]["amount"] == "40"
    assert row["resolved_at"]


def test_read_committed_preview_returns_original_result(client, headers):
    from test_legacy_import_flow import _workbook
    content = _workbook([_row()])
    created = client.post('/api/orders/import-preview?filename=read-result.xlsx', content=content, headers=headers)
    session_id = created.json()["session_id"]
    committed = client.post(f"/api/orders/import-preview/{session_id}/commit", content=content, headers=headers)
    assert committed.status_code == 200, committed.text
    read = client.get(f"/api/orders/import-preview/{session_id}", headers=headers)
    assert read.status_code == 200, read.text
    assert read.json()["session"]["status"] == "committed"
    assert read.json()["session"]["result"]["success_rows"] == committed.json()["success_rows"]
    assert read.json()["session"]["result"]["file_sha256"] == committed.json()["file_sha256"]


def test_preview_preserves_tax_formats_and_refreshes_multi_phase_balances(client, headers):
    from decimal import Decimal
    from io import BytesIO
    from openpyxl import load_workbook
    from sqlalchemy import text
    from app.db import db
    from test_legacy_import_flow import _workbook

    workbook = load_workbook(BytesIO(_workbook([_row({
        19: 0.5, 25: 0.005, 26: 10,
        74: "2026-01-01/2026-02-01", 76: "60/40",
        79: "2026-03-01", 81: "2",
    })])))
    workbook.active.cell(3, 19).number_format = "General"
    workbook.active.cell(3, 25).number_format = "0.00%"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    content = stream.getvalue()
    created = client.post('/api/orders/import-preview?filename=tax-formats.xlsx', content=content, headers=headers)
    assert created.status_code == 200, created.text
    committed = client.post(f"/api/orders/import-preview/{created.json()['session_id']}/commit", content=content, headers=headers)
    assert committed.status_code == 200, committed.text
    with db() as conn:
        line = conn.execute(text('SELECT * FROM v_order_line_finance')).mappings().one()
        assert line['sales_tax_rate'] == Decimal('0.5')
        assert line['purchase_tax_rate'] == Decimal('0.5')
        assert line['order_value'] == Decimal('200')
        assert line['purchase_amount'] == Decimal('20.10')
        invoices = conn.execute(text('SELECT pending_invoice_amount FROM sales_invoice')).scalars().all()
        assert invoices == [Decimal('100'), Decimal('100')]
        assert conn.execute(text('SELECT receipt_ratio FROM sales_receipt')).scalar_one() == Decimal('2')
