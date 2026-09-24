from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.ledger_excel import _literal_export_cell


def test_exported_business_text_is_not_an_excel_formula() -> None:
    workbook = Workbook()
    formula_like_text = '=HYPERLINK("https://example.invalid", "click")'
    cell = _literal_export_cell(workbook.active, 1, 1, formula_like_text)
    assert cell.data_type == "s"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    loaded = load_workbook(output)
    assert loaded.active.cell(1, 1).value == formula_like_text
    assert loaded.active.cell(1, 1).data_type == "s"
    loaded.close()
    workbook.close()
