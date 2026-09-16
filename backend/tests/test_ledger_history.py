"""H2：订单号别名与客户经理历史的持久化。

合成数据；历史链通过 register_* 直接构造，模拟预检确认后的写入（H3 接入前）。
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.config import ROOT_DIR
from app.db import db
from app.ledger_history import (
    conflicts_in_project,
    register_current_manager,
    register_current_number,
    find_order_by_number,
    load_manager_history,
    load_order_numbers,
    manager_candidates,
    register_manager_history,
    register_order_numbers,
)

TEMPLATE_PATH = ROOT_DIR / "backend" / "templates" / "市场部业务台账模板.xlsx"


def _row(
    project_code: str, order_no: str, manager: str = "张三", overrides: dict[int, object] | None = None
) -> list[object]:
    values: list[object] = [
        "全额", project_code, "销售部", "安徽分公司", manager,
        date(2026, 7, 21), "商品销售", "常规", "三级一组",
        "客户单位A", "最终用户A", "区域平台A", order_no, "子项目甲",
        "服务器", "规格A", "台", 2, None, None, Decimal("100.00"), None, None, "供应商A",
    ]
    for position, value in (overrides or {}).items():
        values[int(position) - 1] = value
    return values


def _import(client: TestClient, headers: dict[str, str], rows: list[list[object]], name: str = "case.xlsx"):
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook["Sheet1"]
    for excel_row in range(3, 21):
        for column in range(1, 92):
            worksheet.cell(excel_row, column).value = None
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            worksheet.cell(3 + offset, column, value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return client.post(
        f"/api/orders/import-excel?filename={name}", content=output.getvalue(), headers=headers
    )


def _ids(project_code: str, order_no: str) -> dict[str, int]:
    with db() as conn:
        row = conn.execute(
            text(
                """
                SELECT p.id AS project_id, so.id AS sales_order_id
                FROM project p JOIN sales_order so ON so.project_id = p.id
                WHERE p.project_code = :code AND so.order_no = :order_no
                """
            ),
            {"code": project_code, "order_no": order_no},
        ).mappings().first()
    return dict(row) if row else {}


# --- 订单号别名 -------------------------------------------------------------


def test_plain_import_registers_the_current_number_only(client: TestClient, headers: dict[str, str]) -> None:
    assert _import(client, headers, [_row("P-H1", "SO-001")], "a.xlsx").status_code == 200
    ids = _ids("P-H1", "SO-001")
    with db() as conn:
        assert load_order_numbers(conn, ids["sales_order_id"]) == ["SO-001"]


def test_aliases_resolve_to_one_order(client: TestClient, headers: dict[str, str]) -> None:
    """A/B/C 是同一个订单的改号链：任一别名都指向同一订单，不会新增订单。"""
    assert _import(client, headers, [_row("P-H2", "SO-C")], "a.xlsx").status_code == 200
    ids = _ids("P-H2", "SO-C")
    with db() as conn:
        register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-B", "SO-C"])
        for alias in ("SO-A", "SO-B", "SO-C"):
            assert find_order_by_number(conn, ids["project_id"], alias) == ids["sales_order_id"]
        register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-B", "SO-C"])
        assert load_order_numbers(conn, ids["sales_order_id"]) == ["SO-A", "SO-B", "SO-C"]
        orders = conn.execute(
            text("SELECT COUNT(*) FROM sales_order WHERE project_id = :id"),
            {"id": ids["project_id"]},
        ).scalar()
    assert orders == 1, "别名不应产生第二个订单"


def test_alias_conflicting_with_another_orders_current_number(
    client: TestClient, headers: dict[str, str]
) -> None:
    """历史号撞上同框架另一个订单的当前号：必须能被检出（导入侧阻断）。"""
    assert _import(client, headers, [_row("P-H3", "SO-001")], "a.xlsx").status_code == 200
    assert _import(client, headers, [_row("P-H3", "SO-002", overrides={15: "交换机"})], "b.xlsx").status_code == 200
    first = _ids("P-H3", "SO-001")
    second = _ids("P-H3", "SO-002")
    with db() as conn:
        conflicts = conflicts_in_project(
            conn,
            first["project_id"],
            ["SO-OLD", "SO-002"],
            exclude_sales_order_id=first["sales_order_id"],
        )
    assert conflicts == {"SO-002": second["sales_order_id"]}


def test_history_number_claimed_by_two_orders_is_detected(
    client: TestClient, headers: dict[str, str]
) -> None:
    assert _import(client, headers, [_row("P-H4", "SO-001")], "a.xlsx").status_code == 200
    assert _import(client, headers, [_row("P-H4", "SO-002", overrides={15: "交换机"})], "b.xlsx").status_code == 200
    first = _ids("P-H4", "SO-001")
    second = _ids("P-H4", "SO-002")
    with db() as conn:
        register_order_numbers(conn, first["sales_order_id"], ["SO-001", "SO-LEGACY"])
        conflicts = conflicts_in_project(
            conn,
            first["project_id"],
            ["SO-LEGACY"],
            exclude_sales_order_id=second["sales_order_id"],
        )
    assert conflicts == {"SO-LEGACY": first["sales_order_id"]}


def test_same_number_in_another_framework_is_allowed(client: TestClient, headers: dict[str, str]) -> None:
    """不同框架可以有同编号：冲突判定限定在框架内。"""
    assert _import(client, headers, [_row("P-H5", "SO-SHARED")], "a.xlsx").status_code == 200
    assert _import(client, headers, [_row("P-H6", "SO-SHARED")], "b.xlsx").status_code == 200
    with db() as conn:
        other = _ids("P-H6", "SO-SHARED")
        conflicts = conflicts_in_project(
            conn,
            other["project_id"],
            ["SO-SHARED"],
            exclude_sales_order_id=other["sales_order_id"],
        )
    assert conflicts == {}, "跨框架同号不应被判为冲突"


# --- 客户经理历史 -----------------------------------------------------------


def test_manager_history_keeps_repeated_tenure(client: TestClient, headers: dict[str, str]) -> None:
    """甲→乙→甲 是两次交接，必须保留三个位置。"""
    assert _import(client, headers, [_row("P-H7", "SO-001", manager="甲")], "a.xlsx").status_code == 200
    ids = _ids("P-H7", "SO-001")
    with db() as conn:
        register_manager_history(conn, ids["project_id"], ["甲", "乙", "甲"])
        history = load_manager_history(conn, ids["project_id"])
    assert [item["manager_name"] for item in history] == ["甲", "乙", "甲"]
    assert [item["history_order"] for item in history] == [1, 2, 3]


def test_manager_history_without_dates_keeps_effective_from_null(
    client: TestClient, headers: dict[str, str]
) -> None:
    """缺少变更日期时留空，不编造生效时间。"""
    assert _import(client, headers, [_row("P-H8", "SO-001", manager="甲")], "a.xlsx").status_code == 200
    ids = _ids("P-H8", "SO-001")
    with db() as conn:
        register_manager_history(conn, ids["project_id"], ["甲", "乙"])
        history = load_manager_history(conn, ids["project_id"])
    assert all(item["effective_from"] is None for item in history)


def test_manager_candidates_include_current_and_history(client: TestClient, headers: dict[str, str]) -> None:
    assert _import(client, headers, [_row("P-H9", "SO-001", manager="丙")], "a.xlsx").status_code == 200
    ids = _ids("P-H9", "SO-001")
    with db() as conn:
        register_manager_history(conn, ids["project_id"], ["甲", "乙", "丙"])
        names = manager_candidates(conn, ids["project_id"])
    assert names == ["丙", "甲", "乙"]  # 现任在前，其后是历次（去重、保序）


def test_partial_registration_never_prunes_confirmed_history(
    client: TestClient, headers: dict[str, str]
) -> None:
    """普通导入只登记现任/当前号，绝不能删掉预检已确认的历史。"""
    assert _import(client, headers, [_row("P-H10", "SO-001", manager="丙")], "a.xlsx").status_code == 200
    ids = _ids("P-H10", "SO-001")
    with db() as conn:
        register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-B", "SO-001"])
        register_manager_history(conn, ids["project_id"], ["甲", "乙", "丙"])
        # 再走一次“普通导入只登记当前值”的路径：已有历史时它必须什么都不做
        assert register_current_number(conn, ids["sales_order_id"], "SO-001") is False
        assert register_current_manager(conn, ids["project_id"], "丙") is False
        numbers = load_order_numbers(conn, ids["sales_order_id"])
        managers = [item["manager_name"] for item in load_manager_history(conn, ids["project_id"])]
    assert numbers == ["SO-A", "SO-B", "SO-001"]
    assert managers == ["甲", "乙", "丙"]


def test_register_rejects_shortening_confirmed_history(
    client: TestClient, headers: dict[str, str]
) -> None:
    """只补不删：不允许用更短的链把已确认的别名历史截短。

    （原 `test_prune_removes_only_shorter_tail` 断言的是"默认就会删掉尾部"，
    那正是 Codex 判定的缺陷——截短后搜索再也命中不到被删掉的历史别名。）
    """
    assert _import(client, headers, [_row("P-H11", "SO-001", manager="丙")], "a.xlsx").status_code == 200
    ids = _ids("P-H11", "SO-001")
    with db() as conn:
        register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-B", "SO-C"])
        with pytest.raises(ValueError):
            register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-C"])
        assert load_order_numbers(conn, ids["sales_order_id"]) == ["SO-A", "SO-B", "SO-C"]


def test_explicit_prune_replaces_the_whole_chain(
    client: TestClient, headers: dict[str, str]
) -> None:
    """只有显式 prune=True（整份文件全量替换的维护场景）才允许整体替换成更短的链。"""
    assert _import(client, headers, [_row("P-H11", "SO-001", manager="丙")], "a.xlsx").status_code == 200
    ids = _ids("P-H11", "SO-001")
    with db() as conn:
        register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-B", "SO-C"])
        register_order_numbers(conn, ids["sales_order_id"], ["SO-A", "SO-C"], prune=True)
        assert load_order_numbers(conn, ids["sales_order_id"]) == ["SO-A", "SO-C"]


# --- 事务边界 ---------------------------------------------------------------


def test_failed_import_rolls_back_history_too(client: TestClient, headers: dict[str, str]) -> None:
    """整批回滚时，历史表也不能留下半个订单号的记录。"""
    good = _row("P-H12", "SO-001")
    bad = _row("P-H12", "SO-002", overrides={15: "交换机"})
    bad[1] = None  # 缺项目编号但有业务载荷 → 整批中止
    response = _import(client, headers, [good, bad], "mixed.xlsx")
    assert response.status_code == 422, response.text
    with db() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM sales_order")).scalar() == 0
        assert conn.execute(text("SELECT COUNT(*) FROM sales_order_number_history")).scalar() == 0
        assert conn.execute(text("SELECT COUNT(*) FROM project_manager_history")).scalar() == 0


def test_import_registers_current_manager(client: TestClient, headers: dict[str, str]) -> None:
    assert _import(client, headers, [_row("P-H13", "SO-001", manager="王五")], "a.xlsx").status_code == 200
    ids = _ids("P-H13", "SO-001")
    with db() as conn:
        history = load_manager_history(conn, ids["project_id"])
    assert [item["manager_name"] for item in history] == ["王五"]
    assert history[0]["source"] == "legacy_import"
