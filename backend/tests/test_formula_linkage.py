from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import text

from app.db import db
from test_financial_integration import _create_order, _finance
# Database/client fixtures come from the shared, isolated conftest.py (PR #12).


def request(client, headers, method, url, data=None):
    response = client.request(method, url, json=data, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def batch_item(client, headers, scope, line):
    editor = request(client, headers, 'POST', f'/api/{scope}/batch-editor/rows', {'order_line_ids': [line]})
    return {'order_line_id': line, **{
        col['key']: editor['rows'][0]['values'][index]
        for index, col in enumerate(editor['columns']) if col['editable'] and col['key']
    }}


def test_formula_linkage_after_edits_and_deletions(client, headers):
    line, payload = _create_order(client, headers, 'FORMULAS', sales_tax_rate='13', purchase_tax_rate='13')
    purchase = request(client, headers, 'POST', f'/api/purchases/{line}/contracts', {'signed_amount': 300, 'unsigned_amount': 999})
    assert Decimal(str(purchase['contracts'][0]['unsigned_amount'])) == Decimal('491')
    contract_id = purchase['contracts'][0]['id']
    first = request(client, headers, 'POST', f'/api/sales/{line}/invoices', {'invoice_amount': 200, 'pending_invoice_amount': 999})
    first_id = first['invoices'][0]['id']
    second = request(client, headers, 'POST', f'/api/sales/{line}/invoices', {'invoice_amount': 200})
    second_id = second['invoices'][1]['id']
    received = request(client, headers, 'POST', f'/api/sales/{line}/receipts', {'receipt_amount': 2, 'receipt_ratio': 99})
    assert Decimal(str(received['receipts'][0]['receipt_ratio'])) == Decimal('0.5')
    assert all(Decimal(str(r['pending_invoice_amount'])) == Decimal('730') for r in received['invoices'])
    assert all(Decimal(str(r['delivered_not_invoiced_amount'])) == Decimal('165') for r in received['invoices'])
    updated = request(client, headers, 'PUT', f'/api/sales/invoices/{first_id}', {'invoice_amount': 600})
    assert Decimal(str(updated['receipts'][0]['receipt_ratio'])) == Decimal('0.25')
    assert Decimal(str(updated['invoices'][0]['delivered_not_invoiced_amount'])) == Decimal('-235')
    updated = request(client, headers, 'DELETE', f'/api/sales/invoices/{second_id}')
    assert abs(float(updated['receipts'][0]['receipt_ratio']) - 1/3) < 0.000001
    updated = request(client, headers, 'DELETE', f'/api/sales/invoices/{first_id}')
    assert updated['receipts'][0]['receipt_ratio'] is None

    payload.update(quantity='20', unit_price='226')
    request(client, headers, 'PUT', f'/api/orders/{line}', payload)
    state = _finance(line)
    assert state['order_value'] == 4520
    assert state['delivery_value'] == 1130
    assert state['purchase_amount'] == 1582
    purchase = request(client, headers, 'GET', f'/api/purchases/{line}')
    assert Decimal(str(purchase['contracts'][0]['unsigned_amount'])) == Decimal('1282')
    request(client, headers, 'DELETE', f'/api/purchases/contracts/{contract_id}')


def test_all_batch_paths_recalculate_without_submitted_formula_fields(client, headers):
    line, _ = _create_order(client, headers, 'FORMULA-BATCH', sales_tax_rate='13', purchase_tax_rate='13')
    basic = batch_item(client, headers, 'orders', line)
    assert 'order_value' not in basic
    basic.update(quantity=20, net_unit_price=200)
    request(client, headers, 'PUT', '/api/orders/batch-basic', {'items': [basic]})
    state = _finance(line)
    assert state['purchase_amount'] == 1582
    assert state['pending_delivery_amount'] == 3390
    purchase = batch_item(client, headers, 'purchases', line)
    assert 'purchase_unsigned_amount' not in purchase
    purchase.update(delivery_quantity=2, purchase_unit_price=90.4)
    request(client, headers, 'PUT', '/api/purchases/batch', {'items': [purchase]})
    state = _finance(line)
    assert state['purchase_unit_price_no_tax'] == 80
    assert state['purchase_amount'] == 1808
    assert state['delivery_cost'] == Decimal('180.8')
    sales = batch_item(client, headers, 'sales', line)
    assert 'receipt1_ratio' not in sales
    sales.update(sales_invoice_amount=400, receipt1_amount=2)
    request(client, headers, 'PUT', '/api/sales/batch', {'items': [sales]})
    detail = request(client, headers, 'GET', f'/api/sales/{line}')
    assert Decimal(str(detail['receipts'][0]['receipt_ratio'])) == Decimal('0.5')
    assert len(detail['receipts']) == 1
    assert Decimal(str(detail['invoices'][0]['pending_invoice_amount'])) == Decimal('4120')


def test_excel_low_percentage_round_trip_ignores_stale_formulas(client, headers):
    line, payload = _create_order(client, headers, 'FORMULA-EXCEL', sales_tax_rate='0.5', purchase_tax_rate='1')
    request(client, headers, 'POST', f'/api/sales/{line}/invoices', {'invoice_amount': 400})
    request(client, headers, 'POST', f'/api/sales/{line}/receipts', {'receipt_amount': 2})
    response = client.get('/api/orders/export', headers=headers)
    assert response.status_code == 200, response.text
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    row = next(row for row in range(3, ws.max_row + 1) if ws.cell(row, 2).value == payload['project_code'])
    assert ws.cell(row, 19).value == 0.005
    assert ws.cell(row, 25).value == 0.01
    assert ws.cell(row, 82).value == 0.005
    ws.cell(row, 2, 'QA-FORMULA-REIMPORT')
    ws.cell(row, 13, 'SO-FORMULA-REIMPORT')
    # A normal numeric input is already percentage points, even below one.
    ws.cell(row, 25, 0.5).number_format = 'General'
    for col in (22, 23, 28, 29, 32, 33, 34, 35, 36, 37, 38, 77, 78, 82):
        ws.cell(row, col, 999)
    output = BytesIO()
    wb.save(output)
    imported = client.post('/api/orders/import-excel?filename=formulas.xlsx', content=output.getvalue(), headers=headers)
    assert imported.status_code == 200, imported.text
    with db() as conn:
        state = dict(conn.execute(text("SELECT * FROM v_order_line_finance WHERE project_code='QA-FORMULA-REIMPORT'")).mappings().one())
    assert state['sales_tax_rate'] == Decimal('0.5')
    assert state['purchase_tax_rate'] == Decimal('0.5')
    assert state['order_value'] == 1005
    assert state['purchase_amount'] == Decimal('703.5')
    assert state['pending_delivery_amount'] == Decimal('502.5')
    detail = request(client, headers, 'GET', f"/api/sales/{state['order_line_id']}")
    assert Decimal(str(detail['receipts'][0]['receipt_ratio'])) == Decimal('0.5')
    assert Decimal(str(detail['invoices'][0]['pending_invoice_amount'])) == Decimal('605')


def test_exported_formula_balances_do_not_create_empty_import_records(client, headers):
    line, payload = _create_order(client, headers, 'FORMULA-EMPTY')
    response = client.get('/api/orders/export', headers=headers)
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    row = next(i for i in range(3, ws.max_row + 1) if ws.cell(i, 2).value == payload['project_code'])
    ws.cell(row, 2, 'QA-FORMULA-EMPTY-COPY')
    ws.cell(row, 13, 'SO-FORMULA-EMPTY-COPY')
    output = BytesIO()
    wb.save(output)
    imported = client.post('/api/orders/import-excel?filename=empty.xlsx', content=output.getvalue(), headers=headers)
    assert imported.status_code == 200, imported.text
    with db() as conn:
        copied = conn.execute(text("SELECT order_line_id FROM v_order_line_finance WHERE project_code='QA-FORMULA-EMPTY-COPY'")).scalar_one()
        for table in ('purchase_contract', 'sales_invoice', 'sales_receipt'):
            assert conn.execute(text(f"SELECT COUNT(*) FROM {table} WHERE order_line_id=:id"), {'id': copied}).scalar() == 0


def test_small_invoice_large_receipt_preserves_calculated_ratio(client, headers):
    line, _ = _create_order(client, headers, 'FORMULA-LARGE-RATIO')
    request(client, headers, 'POST', f'/api/sales/{line}/invoices', {'invoice_amount': '0.01'})
    detail = request(client, headers, 'POST', f'/api/sales/{line}/receipts', {'receipt_amount': '1130'})
    assert Decimal(detail['receipts'][0]['receipt_ratio']) == Decimal('11300000')
    with db() as conn:
        assert conn.execute(text('SELECT receipt_ratio FROM sales_receipt WHERE order_line_id=:id'), {'id': line}).scalar_one() == Decimal('11300000')
