"""A-修正3：四级结构（框架项目 → 订单 → 子项目 → 明细）的必需回归用例。

全部使用合成数据，在隔离数据库上验证，不以“真实文件导入成功”作为完成标准。
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from conftest import TEST_PASSWORD
from app.config import ROOT_DIR
from app.db import db

TEMPLATE_PATH = ROOT_DIR / "backend" / "templates" / "市场部业务台账模板.xlsx"
USER_PASSWORD = "Four-Level-Test-20260915!"
PROJECT_CODE = "P-LEVEL-001"


def _row(overrides: dict[int, object] | None = None) -> list[object]:
    """91 列模板的一行；默认全部落在同一个订单、同一个子项目里。"""
    values: list[object] = [
        "全额", PROJECT_CODE, "销售部", "安徽分公司", "张三",
        date(2026, 7, 21), "商品销售", "常规", "三级一组",
        "客户单位A", "最终用户A", "区域平台A", "SO-LEVEL-001", "子项目甲",
        "服务器", "规格A", "台", 2, None, None, Decimal("100.00"), None, None, "供应商A",
    ]
    for position, value in (overrides or {}).items():
        values[position - 1] = value
    return values


def _workbook(rows: list[list[object]]) -> bytes:
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook["Sheet1"]
    for excel_row in range(3, 21):
        for column in range(1, 92):
            worksheet.cell(excel_row, column).value = None
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            worksheet.cell(3 + offset, column, value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _import(client: TestClient, headers: dict[str, str], rows: list[list[object]], name: str = "case.xlsx"):
    return client.post(
        f"/api/orders/import-excel?filename={name}", content=_workbook(rows), headers=headers
    )


def _count(table: str) -> int:
    with db() as conn:
        return int(conn.execute(text(f"SELECT COUNT(*) FROM `{table}`")).scalar() or 0)


def _sub_projects() -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            text(
                """
                SELECT sp.name, sp.customer_unit_name, sp.end_user_name, sp.regional_platform
                FROM sub_project sp
                JOIN sales_order so ON so.id = sp.sales_order_id
                JOIN project p ON p.id = so.project_id
                WHERE p.project_code = :code AND sp.deleted_at IS NULL
                ORDER BY so.order_no, sp.name
                """
            ),
            {"code": PROJECT_CODE},
        ).mappings().all()
    return [dict(row) for row in rows]


# --- 框架项目层 -------------------------------------------------------------


def test_framework_manager_conflict_across_orders_is_rejected(client: TestClient, headers: dict[str, str]) -> None:
    """同框架不同订单、客户经理不同：拒绝并说明框架字段冲突，原数据不变。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row({13: "SO-LEVEL-002", 5: "李四"})], "manager.xlsx")
    assert response.status_code == 422, response.text
    assert "客户经理" in response.json()["detail"]
    assert _count("project") == 1
    assert _count("order_line") == 1


def test_framework_code_reuse_without_department_rights_is_rejected(
    client: TestClient, headers: dict[str, str]
) -> None:
    """无权部门的账号复用已有框架编号：403，框架归属和明细均不变。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    created = client.post(
        "/api/auth/users",
        headers=headers,
        json={
            "username": "level-dept-user",
            "password": USER_PASSWORD,
            "display_name": "level-dept-user",
            "role_code": "viewer",
            "permissions": ["ledger_import"],
            "department_scope": ["市场部"],
            "department_can_view": True,
            "department_can_entry": True,
            "department_all": False,
        },
    )
    assert created.status_code == 200, created.text
    try:
        login = client.post(
            "/api/auth/login", json={"username": "level-dept-user", "password": USER_PASSWORD}
        )
        scoped = {"Authorization": f"Bearer {login.json()['access_token']}"}
        response = _import(client, scoped, [_row({3: "市场部", 13: "SO-LEVEL-002"})], "dept.xlsx")
        assert response.status_code == 403, response.text
        assert _count("project") == 1
        assert _count("order_line") == 1
    finally:
        with db() as conn:
            user_id = conn.execute(
                text("SELECT id FROM erp_user WHERE username = 'level-dept-user'")
            ).scalar()
            if user_id:
                conn.execute(text("DELETE FROM operation_log WHERE user_id = :id"), {"id": int(user_id)})
                conn.execute(text("DELETE FROM erp_user WHERE id = :id"), {"id": int(user_id)})


# --- 子项目层 ---------------------------------------------------------------


def test_sub_projects_keep_their_own_customer_fields(client: TestClient, headers: dict[str, str]) -> None:
    """同框架不同子项目：框架字段相同，客户单位/最终用户/平台各自保存，不串值。"""
    first = _import(client, headers, [_row()], "a.xlsx")
    assert first.status_code == 200, first.text
    second = _import(
        client,
        headers,
        [_row({13: "SO-LEVEL-002", 14: "子项目乙", 10: "客户单位B", 11: "最终用户B", 12: "区域平台B"})],
        "b.xlsx",
    )
    assert second.status_code == 200, second.text

    stored = {(item["name"]): item for item in _sub_projects()}
    assert stored["子项目甲"]["customer_unit_name"] == "客户单位A"
    assert stored["子项目甲"]["regional_platform"] == "区域平台A"
    assert stored["子项目乙"]["customer_unit_name"] == "客户单位B"
    assert stored["子项目乙"]["regional_platform"] == "区域平台B"


@pytest.mark.parametrize("position,label", [(10, "客户单位"), (11, "最终用户"), (12, "区域平台")])
def test_sub_project_field_conflict_inside_one_sub_project(
    client: TestClient, headers: dict[str, str], position: int, label: str
) -> None:
    """同一订单同一子项目内，这三项不同必须分别拒绝，不能用最后一行覆盖。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row({15: "交换机", position: "改过的值"})], "conflict.xlsx")
    assert response.status_code == 422, response.text
    assert label in response.json()["detail"]
    assert _count("order_line") == 1


def test_same_sub_project_name_in_another_order_is_a_different_sub_project(
    client: TestClient, headers: dict[str, str]
) -> None:
    """同框架不同订单可以有同名子项目、客户不同：按订单隔离，不按名称误合并。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row({13: "SO-LEVEL-002", 10: "客户单位B"})], "second.xlsx")
    assert response.status_code == 200, response.text

    names = [item["name"] for item in _sub_projects()]
    assert names.count("子项目甲") == 2, names
    customers = sorted(item["customer_unit_name"] for item in _sub_projects())
    assert customers == ["客户单位A", "客户单位B"]


# --- 明细层 -----------------------------------------------------------------


def test_same_line_values_in_another_sub_project_are_allowed(
    client: TestClient, headers: dict[str, str]
) -> None:
    """同订单不同子项目，五项明细属性完全相同：允许。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row({14: "子项目乙"})], "second.xlsx")
    assert response.status_code == 200, response.text
    assert _count("order_line") == 2


def test_identical_line_inside_one_sub_project_is_rejected(client: TestClient, headers: dict[str, str]) -> None:
    """同子项目五项全部相同：拒绝。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row()], "duplicate.xlsx")
    assert response.status_code == 422, response.text
    assert _count("order_line") == 1


@pytest.mark.parametrize(
    "overrides",
    [
        {18: 5},                 # 数量不同
        {21: Decimal("88.00")},  # 销售单价不同
        {16: "规格B"},            # 规格不同
        {15: "交换机"},           # 物资名称不同
        {24: "供应商B"},          # 采购厂商不同
    ],
)
def test_each_distinguishing_field_allows_a_new_line(
    client: TestClient, headers: dict[str, str], overrides: dict[int, object]
) -> None:
    """五项里任一项不同都允许成为新明细，采购厂商不丢失。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row(overrides)], "second.xlsx")
    assert response.status_code == 200, response.text
    assert _count("order_line") == 2


def test_two_sub_projects_may_share_goods_with_blank_specification(
    client: TestClient, headers: dict[str, str]
) -> None:
    """两个子项目使用相同物资且规格均为空：成功，各自保留所属关系。"""
    assert _import(client, headers, [_row({16: None})], "first.xlsx").status_code == 200
    response = _import(client, headers, [_row({14: "子项目乙", 16: None})], "second.xlsx")
    assert response.status_code == 200, response.text
    with db() as conn:
        rows = conn.execute(
            text(
                """
                SELECT ol.sub_project_id, sp.name
                FROM order_line ol JOIN sub_project sp ON sp.id = ol.sub_project_id
                ORDER BY ol.id
                """
            )
        ).mappings().all()
    assert [dict(row)["name"] for row in rows] == ["子项目甲", "子项目乙"]


def test_purchase_entry_rejects_a_line_that_becomes_identical(
    client: TestClient, headers: dict[str, str]
) -> None:
    """采购入口改厂商使明细与另一条完全相同：同样拒绝，不能绕过订单入口判重。"""
    assert _import(client, headers, [_row()], "first.xlsx").status_code == 200
    assert _import(client, headers, [_row({24: "供应商B"})], "second.xlsx").status_code == 200
    with db() as conn:
        target = conn.execute(
            text("SELECT id FROM order_line WHERE id = (SELECT MAX(id) FROM order_line)")
        ).scalar()
    response = client.put(
        f"/api/purchases/{int(target)}/summary",
        json={"supplier_name": "供应商A"},
        headers=headers,
    )
    assert response.status_code in (409, 422), response.text
    assert _count("order_line") == 2


# --- 筛选与导出 -------------------------------------------------------------


def test_export_filter_does_not_borrow_another_sub_projects_customer(
    client: TestClient, headers: dict[str, str]
) -> None:
    """按某个子项目的客户单位筛选导出：不携带同框架其他子项目的值。"""
    assert _import(client, headers, [_row()], "a.xlsx").status_code == 200
    assert (
        _import(
            client,
            headers,
            [_row({13: "SO-LEVEL-002", 14: "子项目乙", 10: "客户单位B", 12: "区域平台B"})],
            "b.xlsx",
        ).status_code
        == 200
    )

    response = client.get("/api/orders/export", params={"client_unit": "客户单位B"}, headers=headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    rows = list(worksheet.iter_rows(min_row=3, values_only=True))
    workbook.close()
    assert len(rows) == 1, rows
    assert rows[0][9] == "客户单位B"
    assert rows[0][11] == "区域平台B"
