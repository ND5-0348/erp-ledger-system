from io import BytesIO
from decimal import Decimal
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from app.db import db
from test_template_0916 import current_import_file, workbook_bytes
from test_import_permissions import _create_user, _login


def source_file():
    wb = load_workbook(BytesIO(current_import_file()))
    ws = wb.active
    ws['E3'] = '甲'
    ws['M3'] = 'A/B'
    ws['BV3'] = '2026/8/1，2026/8/2'
    ws['BX3'] = 200
    ws['BW3'] = '票1/票2'
    ws['AU3'] = None
    ws['AW3'] = 50
    values = [c.value for c in ws[3]]
    ws.append(values)
    ws['E4'] = '乙'
    ws['C4'] = '第二部门'
    ws['I4'] = '第二团队'
    return workbook_bytes(wb)


def test_source_rows_preview_commit_and_duplicate_file(client, headers):
    content = source_file()
    preview = client.post('/api/orders/source-import?filename=real.xlsx',content=content,headers=headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()['success_rows'] == 2
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
    imported = client.post('/api/orders/source-import?filename=real.xlsx&preview=false',content=content,headers=headers)
    assert imported.status_code == 200, imported.text
    with db() as conn:
        lines = conn.execute(text('SELECT * FROM v_order_line_finance ORDER BY order_line_id')).mappings().all()
        assert [r['account_manager'] for r in lines] == ['甲','乙']
        assert [r['order_no'] for r in lines] == ['A/B','A/B']
        assert lines[1]['department'] == '第二部门'
        invoices = conn.execute(text('SELECT * FROM sales_invoice')).mappings().all()
        assert len(invoices) == 2
        assert all(r['invoice_date'] is None and r['invoice_amount'] == Decimal('200') and r['invoice_date_text'] == '2026/8/1，2026/8/2' for r in invoices)
        assert conn.execute(text('SELECT COUNT(*) FROM project_manager_history')).scalar_one() == 0
    repeat = client.post('/api/orders/source-import?filename=real.xlsx&preview=false',content=content,headers=headers)
    assert repeat.status_code == 409
    # List and export must expose row ownership, never the first manager on a framework.
    listed = client.get('/api/orders',headers=headers).json()['items']
    assert {r['account_manager'] for r in listed} == {'甲','乙'}
    export = client.get('/api/orders/export',headers=headers)
    book = load_workbook(BytesIO(export.content),data_only=True)
    assert {r[4] for r in book.active.iter_rows(min_row=3,values_only=True)} == {'甲','乙'}
    assert _create_user(client,headers,'source_scope',['order_edit'],department_scope=['第二部门'],department_all=False).status_code == 200
    scoped = _login(client,'source_scope')
    visible = client.get('/api/orders',headers=scoped).json()
    assert visible['total'] == 1 and visible['items'][0]['account_manager'] == '乙'
    ledger = client.get('/api/ledgers',headers=scoped)
    assert ledger.status_code == 200,ledger.text
    assert len(ledger.json()['items']) == 1
    assert Decimal(ledger.json()['items'][0]['order_amount']) == lines[1]['order_value']
    assert ledger.json()['items'][0]['account_manager'] == '乙'
    assert client.get(f"/api/history/lines/{lines[0]['order_line_id']}",headers=scoped).status_code == 403


@pytest.mark.skipif(not os.environ.get('REAL_SOURCE_WORKBOOK'), reason='local acceptance workbook not configured')
def test_real_source_acceptance(client, headers):
    content = Path(os.environ['REAL_SOURCE_WORKBOOK']).read_bytes()
    response = client.post('/api/orders/source-import?filename=acceptance.xlsx&preview=false',content=content,headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()['success_rows'] == 1048
    from decimal import ROUND_HALF_UP
    source = load_workbook(BytesIO(content),data_only=True).worksheets[0]
    with db() as conn:
        rows = conn.execute(text('SELECT v.*,ol.source_excel_row_no FROM v_order_line_finance v JOIN order_line ol ON ol.id=v.order_line_id')).mappings().all()
        assert len(rows) == 1048
        for row in rows:
            no = row['source_excel_row_no']
            for col,key in [(2,'project_code'),(3,'department'),(4,'branch_company'),(5,'account_manager'),(9,'team_level3_name'),(13,'order_no')]:
                assert row[key] == str(source.cell(no,col).value).strip(), (no,key)
            for col,key in [(22,'revenue_no_tax'),(23,'order_value'),(28,'cost_no_tax'),(29,'purchase_amount'),(33,'delivery_value'),(76,'sales_invoice_amount')]:
                value = source.cell(no,col).value
                expected = Decimal(str(value)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP) if value is not None else None
                assert (row[key] or 0) == (expected or 0), (no,key,row[key],expected)
            receipts = sum((Decimal(str(source.cell(no,col).value or 0)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP) for col in (81,85)),Decimal(0))
            assert row['total_received'] == receipts, no
            assert row['delivery_accounts_receivable'] == (row['delivery_value'] or 0)-receipts
            assert row['invoice_accounts_receivable'] == (row['sales_invoice_amount'] or 0)-receipts
    assert client.get('/api/orders',headers=headers).json()['total'] == 1048
    assert client.get('/api/ledgers?limit=500',headers=headers).json()['total'] == 85


def test_source_import_failure_rolls_back_every_row(client, headers):
    wb = load_workbook(BytesIO(source_file()))
    wb.active['F4'] = 'bad date'
    result = client.post('/api/orders/source-import?filename=bad.xlsx&preview=false',content=workbook_bytes(wb),headers=headers)
    assert result.status_code == 422
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
        assert conn.execute(text('SELECT COUNT(*) FROM import_batch')).scalar_one() == 0
