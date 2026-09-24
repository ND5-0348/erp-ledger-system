from io import BytesIO
from decimal import Decimal
from openpyxl import load_workbook
from sqlalchemy import text
from app.db import db
from test_financial_integration import _create_order


def rename(client, headers, line, old, new):
    return client.post('/api/history/rename-orders', headers=headers, json={'items':[{'order_line_id':line,'expected_order_no':old,'order_no':new,'reason':'测试正式改号'}]})


def test_alias_queries_export_keep_one_order_and_all_phases(client, headers):
    line,payload=_create_order(client,headers,'H4')
    assert rename(client,headers,line,payload['order_no'],'H4-B').status_code==200
    assert rename(client,headers,line,'H4-B','H4-C').status_code==200
    for i in range(4):
        response=client.post(f'/api/sales/{line}/invoices',headers=headers,json={'invoice_date':f'2026-09-0{i+1}','invoice_amount':'10','invoice_no':f'INV-{i}'})
        assert response.status_code==200,response.text
    for endpoint in ['orders','purchases','sales','ledgers']:
        for number in [payload['order_no'],'H4-B','H4-C']:
            response=client.get(f'/api/{endpoint}',params={'order_id':number},headers=headers)
            assert response.status_code==200,response.text
            assert response.json()['total']==1
    current=client.get('/api/orders',headers=headers).json()['items'][0]
    assert current['order_number_history']==[payload['order_no'],'H4-B','H4-C']
    response=client.get('/api/history/export',params={'order_id':payload['order_no']},headers=headers)
    assert response.status_code==200,response.text
    workbook=load_workbook(BytesIO(response.content))
    assert workbook['明细'].max_row==2
    assert workbook['订单号历史'].max_row==4
    assert workbook['销售开票'].max_row==5
    standard=client.get('/api/orders/export',params={'order_id':payload['order_no']},headers=headers)
    assert standard.status_code==200,standard.text
    assert load_workbook(BytesIO(standard.content)).active.cell(3,13).value=='H4-C'
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one()==1
        assert conn.execute(text('SELECT SUM(order_value) FROM order_line')).scalar_one()==Decimal('1130')


def test_transfer_preserves_repeated_history_and_blocks_stale_or_plain_edit(client, headers):
    line,payload=_create_order(client,headers,'MANAGER')
    for manager in ['乙',payload['account_manager']]:
        ctx=client.get(f'/api/history/lines/{line}',headers=headers).json()
        current=ctx['current']
        change={'order_line_id':line,'expected':{k:current[k] for k in ['account_manager','department','branch_company','team_level3_name']},'account_manager':manager,'department':current['department'],'branch_company':current['branch_company'],'team_level3_name':current['team_level3_name'],'reason':'测试交接'}
        assert client.post('/api/history/transfer-project',headers=headers,json=change).status_code==200
    ctx=client.get(f'/api/history/lines/{line}',headers=headers).json()
    assert [r['manager_name'] for r in ctx['managers']]==[payload['account_manager'],'乙',payload['account_manager']]
    assert all(r['effective_from'] is None for r in ctx['managers'])
    for endpoint in ['sales','purchases','ledgers']:
        assert client.get(f'/api/{endpoint}',params={'manager':'乙'},headers=headers).json()['total']==0
        response=client.get(f'/api/{endpoint}',params={'manager':'乙','include_history_manager':'true'},headers=headers)
        assert response.status_code==200,response.text
        assert response.json()['total']==1
    assert client.put(f'/api/orders/{line}',headers=headers,json={**payload,'account_manager':'丙'}).status_code==409
    assert client.put(f'/api/orders/{line}',headers=headers,json={**payload,'order_no':'BYPASS'}).status_code==409
    assert client.post('/api/history/transfer-project',headers=headers,json=change).status_code==409


def test_rename_batch_conflict_rolls_back_and_old_alias_cannot_create_order(client, headers):
    line,payload=_create_order(client,headers,'CONFLICT')
    other,_=_create_order(client,headers,'OTHER',project_code=payload['project_code'])
    response=client.post('/api/history/rename-orders',headers=headers,json={'items':[
        {'order_line_id':line,'expected_order_no':payload['order_no'],'order_no':'NEW','reason':'test'},
        {'order_line_id':other,'expected_order_no':'SO-OTHER','order_no':'NEW','reason':'test'}]})
    assert response.status_code==409,response.text
    assert client.get(f'/api/history/lines/{line}',headers=headers).json()['current']['order_no']==payload['order_no']
    assert rename(client,headers,line,payload['order_no'],'NEW').status_code==200
    assert client.post('/api/orders',headers=headers,json={**payload,'goods_name':'another'}).status_code==409


def test_history_scope_and_exact_phase_dates(client, headers):
    from test_import_permissions import _create_user, _login
    line,payload=_create_order(client,headers,'SCOPE',department='甲部')
    assert _create_user(client,headers,'h4limited',['order_edit'],department_scope=['甲部'],department_all=False).status_code==200
    limited=_login(client,'h4limited')
    ctx=client.get(f'/api/history/lines/{line}',headers=limited).json()['current']
    change={'order_line_id':line,'expected':{k:ctx[k] for k in ['account_manager','department','branch_company','team_level3_name']},'account_manager':'乙','department':'乙部','reason':'跨部门'}
    assert client.post('/api/history/transfer-project',headers=limited,json=change).status_code==403
    hidden,_=_create_order(client,headers,'HIDDEN',department='乙部')
    assert client.get(f'/api/history/lines/{hidden}',headers=limited).status_code==403
    assert rename(client,limited,hidden,'SO-HIDDEN','X').status_code==403
    assert client.get('/api/orders',headers=limited,params={'order_id':'HIDDEN'}).json()['total']==0
    for kind,path,field in [('sales','receipts','receipt'),('purchases','payments','payment')]:
        for date,amount in [('2026-01-01','10'),('2026-02-01','20')]:
            response=client.post(f'/api/{kind}/{line}/{path}',headers=headers,json={f'{field}_date':date,f'{field}_amount':amount})
            assert response.status_code==200,response.text
        response=client.get(f'/api/{kind}',headers=limited,params={f'{field}_start_date':'2026-01-01',f'{field}_end_date':'2026-01-01'})
        assert response.status_code==200,response.text
        assert response.json()['total']==1
        assert Decimal(str(response.json()['items'][0]['total_received' if kind=='sales' else 'total_paid']))==10
