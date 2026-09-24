"""92-column contract, source-driven balances, and legacy financial alignment."""
from io import BytesIO
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db import db, _split_sql
from app.config import ROOT_DIR
from app.ledger_excel import template_bytes, TEMPLATE_HEADERS, editor_columns
from app.maintenance_import import normalize_source
from test_financial_integration import _create_order, _finance, _excel_import_file


def workbook_bytes(wb):
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def current_import_file():
    old = load_workbook(BytesIO(_excel_import_file()))
    values = [c.value for c in old.active[3]]
    old.close()
    values.insert(88, Decimal('999999'))  # Deliberately forged cached balance.
    wb = load_workbook(BytesIO(template_bytes()))
    for i, value in enumerate(values, 1):
        wb.active.cell(3, i).value = value
    return workbook_bytes(wb)


def test_template_headers_styles_and_readonly_balances():
    wb = load_workbook(BytesIO(template_bytes()))
    ws = wb.active
    assert ws.max_column == 92
    assert [c.value for c in ws[2]] == TEMPLATE_HEADERS
    assert [ws.cell(2, i).value for i in range(88, 93)] == [
        '交付应收款', '开票应收款', '是否关闭', '人工成本', '其他成本']
    assert 'CI1:CL1' in ws.merged_cells and 'CM1:CN1' in ws.merged_cells
    assert ws['CM1'].value == '补充成本'
    for col in [6,13,15,22,23,28,29,33,69,70]:
        assert ws.cell(2,col).fill.fgColor.rgb == 'FFFEFF00'
    cols = editor_columns('sales')
    assert all(not cols[i-1]['editable'] for i in (87,88,89,90))
    assert [cols[i-1]['key'] for i in (88,89,91,92)] == [
        'delivery_accounts_receivable','invoice_accounts_receivable','labor_cost','other_cost']
    assert cols[90]['editable'] and cols[91]['editable']
    wb.close()


@pytest.mark.parametrize('version', [91,92])
def test_old_and_new_import_preserve_costs_and_ignore_cached_balances(client, headers, version):
    content = _excel_import_file() if version == 91 else current_import_file()
    result = client.post('/api/orders/import-excel?filename=synthetic.xlsx', content=content, headers=headers)
    assert result.status_code == 200, result.text
    with db() as conn:
        rows = conn.execute(text('SELECT * FROM v_order_line_finance')).mappings().all()
    assert len(rows) == 1
    row = rows[0]
    assert row['labor_cost'] == Decimal('12.34')
    assert row['other_cost'] == Decimal('5.67')
    assert row['delivery_accounts_receivable'] == Decimal('-113')
    assert row['invoice_accounts_receivable'] == 0
    export = client.get('/api/orders/export', headers=headers)
    assert export.status_code == 200, export.text
    wb = load_workbook(BytesIO(export.content), data_only=True)
    assert [wb.active.cell(3,i).value for i in (88,89,91,92)] == [-113,0,12.34,5.67]
    wb.close()


@pytest.mark.parametrize('version', [91,92])
def test_maintenance_normalization_preserves_tail(version):
    content = _excel_import_file() if version == 91 else current_import_file()
    normalized, layout, _, originals = normalize_source(content)
    assert layout == f'{version}列标准模板'
    wb = load_workbook(BytesIO(normalized), data_only=True)
    assert [wb.active.cell(3,i).value for i in (90,91,92)] == ['进行中',12.34,5.67]
    assert len(originals[3]) == version
    wb.close()


@pytest.mark.parametrize('version', [91,92])
def test_preview_commit_preserves_shifted_costs_with_extra_blank_columns(client, headers, version):
    content = _excel_import_file() if version == 91 else current_import_file()
    wb = load_workbook(BytesIO(content))
    # A styled empty trailing column must not change layout recognition.
    wb.active.cell(3, 98).number_format = '0.00'
    content = workbook_bytes(wb)
    preview = client.post('/api/orders/import-preview?filename=preview.xlsx',content=content,headers=headers)
    assert preview.status_code == 200,preview.text
    session = preview.json()['session_id']
    committed = client.post(f'/api/orders/import-preview/{session}/commit',content=content,headers=headers)
    assert committed.status_code == 200,committed.text
    with db() as conn:
        row = conn.execute(text('SELECT * FROM v_order_line_finance')).mappings().one()
    assert row['labor_cost'] == Decimal('12.34') and row['other_cost'] == Decimal('5.67')
    assert row['delivery_accounts_receivable'] == Decimal('-113')
    assert row['invoice_accounts_receivable'] == 0


def test_invalid_tail_and_partial_failure_roll_back(client, headers):
    _create_order(client,headers,'KEEP-AR')
    with db() as conn:
        before = [dict(r) for r in conn.execute(text('SELECT * FROM v_order_line_finance')).mappings()]
    wb = load_workbook(BytesIO(current_import_file()))
    values = [c.value for c in wb.active[3]]
    values[1],values[12],values[91] = 'BAD-TAIL','BAD-TAIL',-1
    wb.active.append(values)
    response=client.post('/api/orders/import-excel?filename=bad-tail.xlsx',content=workbook_bytes(wb),headers=headers)
    assert response.status_code == 422,response.text
    with db() as conn:
        assert [dict(r) for r in conn.execute(text('SELECT * FROM v_order_line_finance')).mappings()] == before
    wb = load_workbook(BytesIO(current_import_file()))
    wb.active['CN2']='错误的成本字段'
    content=workbook_bytes(wb)
    response=client.post('/api/orders/import-excel?filename=bad-header.xlsx',content=content,headers=headers)
    assert response.status_code == 422,response.text
    response=client.post('/api/orders/import-preview?filename=bad-header.xlsx',content=content,headers=headers)
    assert response.status_code == 422,response.text


def test_all_phases_changes_soft_delete_and_api_aggregates(client, headers):
    line, payload = _create_order(client, headers, 'AR-SPLIT')  # delivery 565, order 1130
    for amount in ('200.11','300.22','100.33'):
        r = client.post(f'/api/sales/{line}/invoices', json={'invoice_amount':amount}, headers=headers)
        assert r.status_code == 200, r.text
    for amount in ('250.01','200.02','100.03'):
        r = client.post(f'/api/sales/{line}/receipts', json={'receipt_amount':amount}, headers=headers)
        assert r.status_code == 200, r.text
    expected = (Decimal('14.94'), Decimal('50.60'))
    def check(row, values=expected):
        assert Decimal(str(row['delivery_accounts_receivable'])) == values[0]
        assert Decimal(str(row['invoice_accounts_receivable'])) == values[1]
    check(_finance(line))
    for path in ('/api/orders','/api/sales','/api/ledgers'):
        response = client.get(path,headers=headers)
        assert response.status_code == 200,response.text
        check(response.json()['items'][0])
    check(client.get(f'/api/sales/{line}',headers=headers).json()['summary'])
    dashboard = client.get('/api/dashboard/summary',headers=headers).json()
    assert dashboard['deliveryAccountsReceivable'] == '14.94'
    assert dashboard['invoiceAccountsReceivable'] == '50.60'
    editor = client.post('/api/sales/batch-editor/rows',json={'order_line_ids':[line]},headers=headers).json()
    assert tuple(Decimal(str(v)) for v in editor['rows'][0]['values'][87:89]) == expected
    # Forging either new read-only field cannot change source data.
    for field in ('delivery_accounts_receivable','invoice_accounts_receivable'):
        r = client.put('/api/sales/batch',json={'items':[{'order_line_id':line,field:'0'}]},headers=headers)
        assert r.status_code == 422,r.text
    with db() as conn:
        invoice_id = conn.execute(text('SELECT id FROM sales_invoice WHERE phase_no=3')).scalar_one()
        receipt_id = conn.execute(text('SELECT id FROM sales_receipt WHERE phase_no=3')).scalar_one()
    r=client.delete(f'/api/sales/invoices/{invoice_id}',headers=headers)
    assert r.status_code == 200,r.text
    check(_finance(line),(Decimal('14.94'),Decimal('-49.73')))
    r=client.delete(f'/api/sales/receipts/{receipt_id}',headers=headers)
    assert r.status_code == 200,r.text
    check(_finance(line),(Decimal('114.97'),Decimal('50.30')))
    r=client.put(f'/api/orders/{line}',json={**payload,'delivery_quantity':'0'},headers=headers)
    assert r.status_code == 200,r.text
    check(_finance(line),(Decimal('-450.03'),Decimal('50.30')))


def test_view_upgrade_is_repeatable_and_preserves_business_records(client, headers):
    line,_ = _create_order(client,headers,'VIEW-UPGRADE')
    from app.backup import BACKUP_TABLES
    def snapshot():
        with db() as conn:
            return {t:[dict(r) for r in conn.execute(text(f'SELECT * FROM `{t}` ORDER BY id')).mappings()] for t in BACKUP_TABLES}
    before = snapshot()
    schema = (ROOT_DIR/'docs/erp_ledger_schema.sql').read_text(encoding='utf-8')
    views = [s for s in _split_sql(schema) if s.upper().startswith(('DROP VIEW','CREATE VIEW'))]
    # Simulate a populated pre-upgrade database, then apply the new view definitions.
    with db() as conn:
        for sql in views:
            old_sql = '\n'.join(line for line in sql.splitlines() if not any(
                field in line for field in ('delivery_accounts_receivable','invoice_accounts_receivable')))
            conn.execute(text(old_sql))
    assert 'delivery_accounts_receivable' not in _finance(line)
    for _ in range(2):
        with db() as conn:
            for sql in views:
                conn.execute(text(sql))
        assert snapshot() == before
        row = _finance(line)
        assert row['delivery_accounts_receivable'] == Decimal('565')
        assert row['invoice_accounts_receivable'] == 0
