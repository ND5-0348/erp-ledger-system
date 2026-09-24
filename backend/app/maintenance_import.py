"""Read-only replacement preview, then signed/checked atomic replacement."""
import base64
import hashlib
import hmac
import json
import time
from io import BytesIO
from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook, Workbook
from pydantic import TypeAdapter
from fastapi import HTTPException
from sqlalchemy import text
from .config import DOCS_DIR, settings
from .ledger_excel import TEMPLATE_HEADERS, NUMBER_COLUMNS, DATE_COLUMNS, PERCENT_COLUMNS, is_template_sample_row, standard_template_version
from .importer import _detect_business_header_row, _is_latest_layout, _has_business_payload, _as_date, _as_tax_rate, import_excel
from .legacy_import_service import parse_row, build_row_state, cross_row_issues, build_normalized_workbook, append_phases, register_histories, FINANCE_SOURCES
from .legacy_resolution import phase_business_issues
from .financial_calculations import calculate_line, refresh_balances
from .validation import Money, PreciseNumber, BusinessDate
from .backup import create_backup, BACKUP_TABLES
from .edit_versions import read_context, bump_epoch
from .audit import write_operation_log


def selected_file(name):
    root=DOCS_DIR.resolve(); path=(root/name).resolve()
    if path.parent!=root or path.suffix.lower()!='.xlsx' or not path.is_file():
        raise HTTPException(400,'请选择受控目录中存在的 Excel 文件')
    if path.stat().st_size>20*1024*1024: raise HTTPException(413,'文件不能超过20MB')
    return path


def dataset(conn):
    ids=conn.execute(text('SELECT id FROM project')).scalars().all()
    context=read_context(conn,ids)
    digest=hashlib.sha256(json.dumps(context,sort_keys=True,separators=(',',':')).encode()).hexdigest()
    return digest,context['data_epoch']


def normalize_source(content):
    try: wb=load_workbook(BytesIO(content),read_only=True,data_only=True)
    except Exception as exc: raise HTTPException(422,'Excel 文件损坏或格式不支持') from exc
    candidates=[]
    for ws in wb.worksheets:
        try: header=_detect_business_header_row(ws)
        except ValueError: continue
        names=[str(v or '').strip() for v in next(ws.iter_rows(min_row=header,max_row=header,values_only=True))]
        candidates.append((ws,header,names))
    if len(candidates)!=1:
        wb.close();raise HTTPException(422,'必须能唯一确定一个业务工作表，请检查表头')
    ws,header,names=candidates[0]
    while names and not names[-1]: names.pop()
    version = standard_template_version(names)
    if version == 92:
        mapping={i:i for i in range(1,93)};layout='92列标准模板'
    elif version == 91:
        mapping={i:i for i in range(1,89)}
        mapping.update({90:89,91:90,92:91});layout='91列标准模板'
    elif len(names)==87 and not _is_latest_layout(names):
        mapping={**{i:i for i in range(1,19)},**{i:i-1 for i in range(20,25)},
                 **{i:i-2 for i in range(26,50)},54:52,**{i:i-2 for i in range(55,61)},
                 **{i:i-2 for i in range(68,90)}}
        layout='旧版87列布局'
    elif len(names)==96 and _is_latest_layout(names):
        mapping={**{i:i for i in range(1,50)},54:65,**{i:i+11 for i in range(55,61)},
                 **{i:i+8 for i in range(68,72)},**{i:i+7 for i in range(73,90)}}
        layout='扩展96列布局'
    else:
        wb.close();raise HTTPException(422,'表头不是支持的92列、91列、87列或96列布局，请核对原文件')
    if version is None:
        mapping[90] = mapping.pop(89)
    # Legacy recognition cannot rely merely on the presence of three names.
    if any(names[i-1]!=label for i,label in ((2,'项目编号'),(13,'销售订单号' if version == 92 else '订单号'),(14,'项目名称'))):
        wb.close();raise HTTPException(422,'业务表头列位置错误')
    normalized=Workbook(); target=normalized.active;target.title='Sheet1';target.append(['业务台账']);target.append(TEMPLATE_HEADERS)
    source_rows={};original_rows={}
    for number,cells in enumerate(ws.iter_rows(min_row=header+1),start=header+1):
        if not any(c.value not in (None,'') for c in cells): continue
        values=[cells[mapping[i]-1].value if i in mapping and mapping[i]<=len(cells) else None for i in range(1,93)]
        if is_template_sample_row(tuple(values)): continue
        target.append(values);source_rows[target.max_row]=number
        original_rows[target.max_row]=[c.value for c in cells]
        for col in (19,25):
            if col in mapping and mapping[col]<=len(cells):target.cell(target.max_row,col).number_format=cells[mapping[col]-1].number_format or 'General'
    wb.close();out=BytesIO();normalized.save(out);normalized.close()
    return out.getvalue(),layout,source_rows,original_rows


def analyze(content):
    normalized,layout,source_rows,original_rows=normalize_source(content)
    wb=load_workbook(BytesIO(normalized),read_only=True,data_only=True);ws=wb.active
    states=[];errors=[];skipped=0;frames={};subprojects={};identities=set()
    def error(row,column,message):errors.append({'row':source_rows.get(row,row),'column':column,'reason':str(message)})
    for row_no,cells in enumerate(ws.iter_rows(min_row=3),start=3):
        values=[c.value for c in cells]
        if not values[1] or not values[12]:
            if _has_business_payload(tuple(values)):error(row_no,'项目编号/订单号','业务行缺少编号')
            else:skipped+=1
            continue
        sources={**FINANCE_SOURCES,'sales_invoice':((74,76,75),)}
        parsed=parse_row(tuple(values),excel_row_no=row_no,sheet_name=ws.title,finance_sources=sources)
        original=original_rows[row_no]
        extra_sources={}
        if layout=='旧版87列布局': extra_sources={'finance_payment_entry':((48,50,49),)}
        elif layout=='扩展96列布局':
            extra_sources={'finance_invoice_check':((51,53,52),(54,56,55)),
                           'finance_payment_entry':((57,59,58),(60,62,61))}
        if extra_sources:
            extra=parse_row(tuple(original),excel_row_no=row_no,sheet_name=ws.title,finance_sources=extra_sources)
            parsed['finance'].update(extra['finance'])
            parsed['issues'].extend(i for i in extra['issues'] if i['columns'] and i['columns'][0] not in ('订单号','客户经理'))
        parsed['source']={'excel_row_no':source_rows[row_no],'values':original,'layout':layout}
        state=build_row_state(row_no,parsed,{},values)
        states.append(state)
        if len(states)>20000: raise HTTPException(413,'单次最多20000条业务行')
        for issue in state.issues:
            if issue.blocking:error(row_no,'/'.join(issue.columns),issue.message)
        for col in (10,15):
            if not str(values[col-1] or '').strip():error(row_no,TEMPLATE_HEADERS[col-1],'必填字段为空')
        try:
            # Validate input columns only. Derived totals are recomputed and may
            # legitimately be negative in an exported historical workbook.
            for col in (18,20,21,23,26,27,29,31,36,42,70,91,92):
                v=values[col-1]
                if v in (None,''):continue
                TypeAdapter(PreciseNumber if col in (18,20,21,26,27,31,36) else Money).validate_python(v)
            for col in (6,30,54,68):
                if values[col-1] not in (None,''): TypeAdapter(BusinessDate).validate_python(_as_date(values[col-1]))
            for business,entry in state.finance.items():
                for pos,phase in enumerate(entry.get('phases',[]),1):
                    for issue in phase_business_issues(phase,label=business,position=pos):error(row_no,business,issue.message)
            if layout=='扩展96列布局' and original[49] not in (None,''):
                TypeAdapter(Money).validate_python(original[49])
                if len(state.finance.get('warehouse_entry',{}).get('phases',[]))!=1:
                    raise ValueError('入库不含税金额不能自动分配给多期，请先核对')
            if len(str(values[72] or ''))>128:raise ValueError('销售开票单号不能超过128字')
            rates={col:_as_tax_rate(values[col-1],cells[col-1].number_format or 'General') for col in (19,25)}
            calc=calculate_line({'quantity':values[17],'sales_tax_rate':rates[19],'sales_unit_price_no_tax':values[19],'sales_unit_price':values[20],'order_value':values[22],
                                 'purchase_tax_rate':rates[25],'purchase_unit_price_no_tax':values[25],'purchase_unit_price':values[26],'purchase_amount':values[28],'delivery_quantity':values[30]})
            if calc.get('order_value') is None:raise ValueError('销售金额缺少可计算的数量、单价或原金额')
            limits={2:64,3:128,4:128,9:128,10:255,11:255,12:128,13:64,14:255,15:255,16:255,17:32,24:255,39:128,40:255,41:255,69:128,71:255}
            for col,limit in limits.items():
                if col==13:
                    if any(len(v)>limit for v in state.order_chain):raise ValueError('订单号超过64字')
                elif len(str(values[col-1] or '').strip())>limit:raise ValueError(f'{TEMPLATE_HEADERS[col-1]}超过{limit}字')
            if any(len(v)>64 for v in state.manager_chain):raise ValueError('客户经理超过64字')
            for group,total_key in (('sales_invoice','order_value'),('sales_receipt','order_value'),('purchase_invoice','purchase_amount'),('purchase_payment','purchase_amount'),('finance_invoice_check','purchase_amount'),('finance_payment_entry','purchase_amount')):
                cap=calc.get(total_key)
                total=sum((Decimal(str(p.get('amount') or 0)) for p in state.finance.get(group,{}).get('phases',[])),Decimal(0))
                if cap is not None and total>Decimal(str(cap)):raise ValueError(f'{group}累计金额超过该明细金额')
            frame=(values[2],values[3],tuple(state.manager_chain),values[8])
            key=values[1]
            if key in frames and frames[key]!=frame:raise ValueError('同框架的经理、部门、分公司或团队不一致')
            frames[key]=frame
            sub=(key,state.order_chain[-1] if state.order_chain else '',values[13])
            fields=tuple(values[9:12])
            if sub in subprojects and subprojects[sub]!=fields:raise ValueError('同子项目的客户、用户或平台不一致')
            subprojects[sub]=fields
            identity=(*sub,values[14],values[15],Decimal(str(values[17] or 0)),Decimal(str(calc.get('sales_unit_price') or 0)),values[23])
            if identity in identities:raise ValueError('同一子项目包含重复明细')
            identities.add(identity)
        except Exception as exc:
            error(row_no,'业务字段',getattr(exc,'detail',str(exc)))
    wb.close()
    for issue in cross_row_issues(None,states):error(issue.row or 0,'历史',issue.message)
    if not states:error(0,'文件','没有有效业务行')
    if len(states)>20000:raise HTTPException(413,'单次最多20000条业务行')
    return {'layout':layout,'total_rows':len(states)+skipped,'valid_rows':len(states),'skipped_rows':skipped,'errors':errors},states,normalized


def sign(data):
    raw=base64.urlsafe_b64encode(json.dumps(data,separators=(',',':')).encode()).decode()
    signature=hmac.new(settings.auth_secret.encode(),raw.encode(),hashlib.sha256).hexdigest()
    return raw+'.'+signature


def verify(token):
    try:
        raw,sig=token.rsplit('.',1)
        if not hmac.compare_digest(sig,hmac.new(settings.auth_secret.encode(),raw.encode(),hashlib.sha256).hexdigest()):raise ValueError()
        data=json.loads(base64.urlsafe_b64decode(raw))
        if data['expires']<time.time():raise ValueError()
        return data
    except (ValueError,KeyError,TypeError,AttributeError):raise HTTPException(409,'预检凭据无效或过期，请重新预检')


def preview(conn,user,name):
    content=selected_file(name).read_bytes(); report,_,_=analyze(content)
    digest,epoch=dataset(conn)
    meta={'file':name,'sha256':hashlib.sha256(content).hexdigest(),'dataset':digest,'epoch':epoch,'user':user.id,'expires':int(time.time())+3600}
    return {**report,'file_name':name,'sha256':meta['sha256'],'database':settings.mysql_database,'data_epoch':epoch,
            'current_rows':conn.execute(text('SELECT COUNT(*) FROM order_line WHERE deleted_at IS NULL')).scalar_one(),
            'error_count':len(report['errors']),'token':sign(meta) if not report['errors'] else None}


def replace(conn,user,name,token,confirmation):
    if confirmation!='替换全部业务数据':raise HTTPException(422,'请输入“替换全部业务数据”确认')
    meta=verify(token);content=selected_file(name).read_bytes()
    if meta['file']!=name or meta['user']!=user.id or meta['sha256']!=hashlib.sha256(content).hexdigest() or meta['dataset']!=dataset(conn)[0]:
        raise HTTPException(409,'文件或业务数据在预检后发生变化，请重新预检')
    report,states,source=analyze(content)
    if report['errors']:raise HTTPException(422,{'message':'预检未通过','errors':report['errors']})
    # Complete pure validation above precedes every DELETE. The verified backup
    # remains as a filesystem recovery copy even if the transaction rolls back.
    backup=create_backup(conn,user,'pre_replace')
    for table in reversed(BACKUP_TABLES):conn.execute(text(f'DELETE FROM `{table}`'))
    normalized,row_map=build_normalized_workbook(states,source_content=source)
    line_ids={}
    result=import_excel(conn,reset=False,workbook_bytes=normalized,source_file_name=name,user=user,strict_template=True,line_id_map=line_ids)
    if result['failed_rows'] or result['success_rows']!=len(states):raise HTTPException(422,'替换写入校验失败，事务已回滚')
    phases=append_phases(conn,states,line_ids,row_map);register_histories(conn,states,line_ids,row_map)
    conn.execute(text('UPDATE import_batch SET source_sha256=:sha WHERE id=:id'),{'sha':meta['sha256'],'id':result['batch_id']})
    for state in states:
        lid=line_ids[row_map[state.excel_row_no]]
        # The document identifier and invoice number are separate inputs.
        if state.raw_values[72] not in (None,''):
            conn.execute(text('UPDATE sales_invoice SET invoice_doc_no=:value WHERE order_line_id=:id'),{'id':lid,'value':str(state.raw_values[72])})
        src=state.parsed['source']
        if src['layout']=='扩展96列布局' and src['values'][49] not in (None,''):
            # Legacy has one warehouse group and a separate net amount. Keep it
            # only if there is exactly one phase; splitting needs confirmation.
            entries=state.finance.get('warehouse_entry',{}).get('phases',[])
            if len(entries)!=1:raise HTTPException(422,'入库不含税金额不能自动分配给多期，事务已回滚')
            conn.execute(text('UPDATE warehouse_entry SET warehouse_amount_no_tax=:value WHERE order_line_id=:id'),{'id':lid,'value':src['values'][49]})
    for lid in line_ids.values():refresh_balances(conn,lid)
    for state in states:
        conn.execute(text('INSERT INTO legacy_import_audit_source (session_id,source_sha256,excel_row_no,raw_json) VALUES (:session,:sha,:row,:raw)'),
                     {'session':'maintenance-'+meta['sha256'][:32],'sha':meta['sha256'],'row':state.excel_row_no,'raw':json.dumps({'values':state.raw_values,'parsed':state.parsed,'source_file':name},default=str,ensure_ascii=False)})
    bump_epoch(conn)
    write_operation_log(conn,user,'系统维护','replace_ledger',f'预检确认后替换全部业务数据，共 {len(states)} 行',after={'sha256':meta['sha256'],'backup_id':backup['id'],'rows':len(states)})
    return {**result,'source_sha256':meta['sha256'],'phases':phases,'backup_id':backup['id'],'data_epoch':dataset(conn)[1]}
