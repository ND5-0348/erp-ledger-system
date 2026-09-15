"""旧台账预检会话服务（H3）。

流程：

1. `create_session` 上传并逐行解析，结果落库（原始值与解析结果分开存），
   **不写任何业务表**；有阻断问题时照样建会话，让用户看到问题清单。
2. `list_session_rows` 分页读取，避免一次返回两万条。
3. `apply_resolutions` 接收结构化的逐项修正（不是让用户重新拼斜杠字符串），
   重新执行全部校验。
4. `commit_session` 在写锁内复核权限、文件摘要与阻断状态，再一次事务写入。

会话默认 24 小时有效、随机不可猜 ID；原始文件不落盘，只保存摘要与解析结果。
"""
from __future__ import annotations

import hashlib
import json
import secrets
from dataclasses import dataclass
from datetime import datetime, timedelta
from decimal import Decimal
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import text
from sqlalchemy.engine import Connection

from .audit import write_operation_log
from .ledger_history import (
    conflicts_in_project,
    find_order_by_number,
    merge_manager_chain,
    merge_order_chain,
    register_manager_history,
    register_order_numbers,
)
from .legacy_ledger_parser import (
    CURRENT_MANAGER_CONFLICT,
    DUPLICATE_PHASE_SUSPECTED,
    ORDER_ALIAS_CONFLICT,
    ORDER_HISTORY_CONFLICT,
    UNPARSABLE_VALUE,
    Issue,
    build_phases,
    merge_chain_sequences,
    merge_manager_history,
    merge_order_number_history,
    parse_amount_sequence,
    parse_date_sequence,
    parse_name_sequence,
)
from .legacy_resolution import (
    RevisionError,
    RowResolution,
    chain_revision_values,
    merge_finance,
    parse_row_resolution,
    phase_business_issues,
    validate_finance_revision,
)

PARSER_VERSION = "legacy-parser-2"
SESSION_TTL_HOURS = 24
MAX_PREVIEW_ROWS = 20000

# 财务列组：业务组 → 该组在 91 列模板里占用的列组列表，每组为
# (日期列, 金额列, 票据列)。列号是模板绝对列号。
#
# 一个业务组可能有**多个模板列组**：标准模板里付款与回款各有两组列
# （付款 55–57 与 58–60、回款 79–82 与 83–86）。期次按“组内列组顺序、
# 列组内位置顺序”展开，两组的来源都保留，不去重也不双计。
FINANCE_SOURCES: dict[str, tuple[tuple[int, int, int | None], ...]] = {
    "sales_invoice": ((74, 76, 73),),
    "sales_receipt": ((79, 81, 80), (83, 85, 84)),
    "purchase_payment": ((55, 57, 56), (58, 60, 59)),
    "purchase_invoice": ((44, 46, 45),),
    "finance_invoice_check": ((50, 52, 51),),
    "warehouse_entry": ((47, 49, 48),),
}

FINANCE_LABELS = {
    "sales_invoice": "销售开票",
    "sales_receipt": "销售回款",
    "purchase_payment": "采购付款",
    "purchase_invoice": "采购收票",
    "finance_invoice_check": "采购入账",
    "warehouse_entry": "采购入库",
}

LABEL_TO_GROUP = {label: name for name, label in FINANCE_LABELS.items()}

ORDER_NO_COLUMN = 13
MANAGER_COLUMN = 5
PROJECT_CODE_COLUMN = 2
PROJECT_NAME_COLUMN = 14
GOODS_COLUMN = 15


def _cell(row: tuple[Any, ...], column: int) -> Any:
    index = column - 1
    return row[index] if 0 <= index < len(row) else None


def _text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip() or None


def _issue_payload(issue: Issue) -> dict[str, Any]:
    return {
        "code": issue.code,
        "message": issue.message,
        "blocking": issue.blocking,
        "sheet": issue.sheet,
        "row": issue.row,
        "columns": list(issue.columns),
    }


def _jsonable(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime,)):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _restore_value(value: Any) -> Any:
    """把会话里存的 ISO 字符串还原成日期对象，再写回规范化工件。

    原始行中的日期在 JSON 里被序列化成字符串；如果直接写回 Excel，
    importer 的日期解析会拒绝带 T 的 ISO 形式。
    """
    if isinstance(value, str) and len(value) >= 10:
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return value
    return value


def _duplicate_phase_issues(
    entry: dict[str, Any], *, label: str, sheet: str, row: int
) -> list[Issue]:
    """同一笔出现在模板的多个列组里时给出提示（不阻断、也不自动去重）。

    标准模板的付款与回款各有两组列，历史文件还可能在单元格内写多期；
    同一组内不同列组出现"同日期同金额"时无法判断是两笔真实业务还是重复填写，
    所以既不去重也不双计，只标记出来交人工确认。
    """
    issues: list[Issue] = []
    seen: dict[tuple[Any, Any], Any] = {}
    for phase in entry.get("phases") or []:
        key = (phase.get("date"), phase.get("amount"))
        group = phase.get("source_group")
        if key in seen and seen[key] != group:
            issues.append(
                Issue(
                    code=DUPLICATE_PHASE_SUSPECTED,
                    message=(
                        f"{label}在模板的不同列组里都填了同一笔"
                        f"（{phase.get('date')} / {phase.get('amount')}）；"
                        "系统没有自动去重、也没有重复计入，请人工确认是否为两笔业务"
                    ),
                    blocking=False,
                    sheet=sheet,
                    row=row,
                    columns=(label,),
                )
            )
        else:
            seen.setdefault(key, group)
    return issues


def parse_row(
    row: tuple[Any, ...],
    *,
    excel_row_no: int,
    sheet_name: str,
) -> dict[str, Any]:
    """解析一行：订单号别名链、负责人交接链、各组财务期次，以及全部问题。"""
    issues: list[Issue] = []

    order_cell = _cell(row, ORDER_NO_COLUMN)
    order_parse = merge_order_number_history([order_cell], order_key=_text(order_cell) or "")
    issues.extend(issue.with_location(sheet=sheet_name, row=excel_row_no) for issue in order_parse.issues)

    manager_cell = _cell(row, MANAGER_COLUMN)
    manager_parse = merge_manager_history([manager_cell], project_code=_text(_cell(row, PROJECT_CODE_COLUMN)) or "")
    issues.extend(issue.with_location(sheet=sheet_name, row=excel_row_no) for issue in manager_parse.issues)

    finance: dict[str, Any] = {}
    for name, sources in FINANCE_SOURCES.items():
        label = FINANCE_LABELS[name]
        entry: dict[str, Any] = {
            "label": label,
            "date_raw": None,
            "amount_raw": None,
            "document_raw": None,
            # 每个模板列组一条来源记录：两组付款/回款列并存时都保留来源，
            # 便于前端逐组展示，也便于人工确认时按组修正。
            "raw_sources": [],
            "phases": [],
        }
        for date_column, amount_column, document_column in sources:
            date_cell = _cell(row, date_column)
            amount_cell = _cell(row, amount_column)
            document_cell = _cell(row, document_column) if document_column else None
            if date_cell in (None, "") and amount_cell in (None, "") and document_cell in (None, ""):
                continue
            entry["raw_sources"].append(
                {
                    "date_column": date_column,
                    "amount_column": amount_column,
                    "document_column": document_column,
                    "date_raw": _text(date_cell),
                    "amount_raw": _text(amount_cell),
                    "document_raw": _text(document_cell),
                }
            )
            phases, group_issues = build_phases(
                date_cell=date_cell,
                amount_cell=amount_cell,
                document_cell=document_cell,
                columns=(label,),
                sheet=sheet_name,
                row=excel_row_no,
            )
            issues.extend(group_issues)
            entry["phases"].extend(
                {
                    "position": len(entry["phases"]) + index + 1,
                    "date": phase.date.isoformat() if phase.date else None,
                    "amount": format(phase.amount, "f") if phase.amount is not None else None,
                    "document_no": phase.document_no,
                    "source_column": date_column,
                    "source_group": len(entry["raw_sources"]),
                }
                for index, phase in enumerate(phases)
            )
        if not entry["raw_sources"]:
            continue
        first = entry["raw_sources"][0]
        entry["date_raw"] = first["date_raw"]
        entry["amount_raw"] = first["amount_raw"]
        entry["document_raw"] = first["document_raw"]
        issues.extend(_duplicate_phase_issues(entry, label=label, sheet=sheet_name, row=excel_row_no))
        finance[name] = entry

    order_tokens = [str(item) for item in order_parse.values]
    manager_chain = [str(item) for item in manager_parse.history]
    return {
        "excel_row_no": excel_row_no,
        "project_code": _text(_cell(row, PROJECT_CODE_COLUMN)),
        "order_no": {
            "raw": _text(order_cell),
            "current": order_tokens[-1] if order_tokens else None,
            "history": [_jsonable(item) for item in order_tokens],
        },
        "manager": {
            "raw": _text(manager_cell),
            "current": manager_chain[-1] if manager_chain else None,
            "history": manager_chain,
        },
        "project_name": _text(_cell(row, PROJECT_NAME_COLUMN)),
        "goods_name": _text(_cell(row, GOODS_COLUMN)),
        "finance": finance,
        "issues": [_issue_payload(issue) for issue in issues],
    }


CHAIN_LABEL_ORDER = "订单号"
CHAIN_LABEL_MANAGER = "客户经理"
CHAIN_CODES = frozenset(
    {
        CURRENT_MANAGER_CONFLICT,
        ORDER_ALIAS_CONFLICT,
        ORDER_HISTORY_CONFLICT,
        UNPARSABLE_VALUE,
    }
)


@dataclass
class RowState:
    """一行在"解析 + 人工确认"之后的完整状态。"""

    excel_row_no: int
    parsed: dict[str, Any]
    revision: RowResolution | None
    stored_resolution: dict[str, Any]
    finance: dict[str, Any]
    order_chain: list[str]
    manager_chain: list[str]
    raw_values: list[Any]
    issues: list[Issue]

    @property
    def blocking(self) -> bool:
        return any(issue.blocking for issue in self.issues)


def _issue_from_payload(payload: dict[str, Any]) -> Issue:
    return Issue(
        code=str(payload.get("code") or ""),
        message=str(payload.get("message") or ""),
        blocking=bool(payload.get("blocking", True)),
        sheet=payload.get("sheet"),
        row=payload.get("row"),
        columns=tuple(str(item) for item in (payload.get("columns") or ())),
    )


def _issue_group(issue: Issue) -> str | None:
    """财务问题的归属组（按标签反查）；链问题返回 None。"""
    for column in issue.columns:
        group = LABEL_TO_GROUP.get(column)
        if group:
            return group
    return None


def _parse_stored_resolution(stored: Any) -> RowResolution | None:
    """读取库里保存的人工确认；`audit` 明细不参与结构校验。"""
    if not isinstance(stored, dict):
        return None
    payload = {key: stored[key] for key in ("order_no", "manager", "finance") if key in stored}
    if not payload:
        return None
    try:
        return RowResolution.model_validate(payload)
    except ValidationError as exc:
        raise HTTPException(
            status_code=409,
            detail="该预检会话里保存的人工确认格式已失效，请重新上传文件建立新会话",
        ) from exc


def build_row_state(
    excel_row_no: int,
    parsed: dict[str, Any],
    stored_resolution: Any,
    raw_values: list[Any] | None = None,
) -> RowState:
    """由"解析结果 + 人工确认"得出该行的最终值与全部问题。

    人工确认过的财务组会用确认值**重新执行**期次数、合计、日期、金额校验；
    未确认的组继续沿用解析阶段的问题——不存在"确认了错误码就放行"的路径。
    """
    revision = _parse_stored_resolution(stored_resolution)
    revised_groups = set(revision.finance) if revision else set()
    revised_order = bool(revision and revision.order_no)
    revised_manager = bool(revision and revision.manager)

    issues: list[Issue] = []
    for payload in parsed.get("issues") or []:
        issue = _issue_from_payload(payload)
        group = _issue_group(issue)
        if group is not None:
            if group in revised_groups:
                continue
        elif issue.code in CHAIN_CODES:
            is_order = CHAIN_LABEL_ORDER in issue.columns
            is_manager = CHAIN_LABEL_MANAGER in issue.columns
            if is_order and revised_order:
                continue
            if is_manager and revised_manager:
                continue
        issues.append(issue)

    parsed_finance = parsed.get("finance") or {}
    if revision:
        for name, entry_revision in revision.finance.items():
            label = FINANCE_LABELS.get(name, name)
            _, group_issues, _ = validate_finance_revision(
                parsed_finance.get(name) or {}, entry_revision, label=label
            )
            issues.extend(
                issue.with_location(
                    sheet=parsed.get("sheet"), row=excel_row_no, columns=(label,)
                )
                for issue in group_issues
            )

    order_chain = (chain_revision_values(revision.order_no) if revision else None) or [
        str(item) for item in (parsed.get("order_no", {}).get("history") or [])
    ]
    manager_chain = (chain_revision_values(revision.manager) if revision else None) or [
        str(item) for item in (parsed.get("manager", {}).get("history") or [])
    ]

    return RowState(
        excel_row_no=excel_row_no,
        parsed=parsed,
        revision=revision,
        stored_resolution=dict(stored_resolution) if isinstance(stored_resolution, dict) else {},
        finance=merge_finance(parsed_finance, revision.finance if revision else None),
        order_chain=[str(item) for item in order_chain],
        manager_chain=[str(item) for item in manager_chain],
        raw_values=list(raw_values or []),
        issues=issues,
    )


def summarize_states(states: list[RowState]) -> dict[str, Any]:
    blocking = sum(1 for state in states if state.blocking)
    warning = sum(
        1 for state in states if any(not issue.blocking for issue in state.issues)
    )
    multi_value = sum(
        1 for state in states if len(state.order_chain) > 1 or len(state.manager_chain) > 1
    )
    return {
        "total_rows": len(states),
        "blocking_rows": blocking,
        "warning_rows": warning,
        "multi_value_rows": multi_value,
        "comparable": blocking == 0,
    }


def create_session(
    conn: Connection,
    *,
    user_id: int,
    file_name: str,
    content: bytes,
    sheet_index: int = 0,
    header_row: int = 2,
    data_start_row: int = 3,
) -> dict[str, Any]:
    """解析整份文件并落库为预检会话；**不写任何业务表**。"""
    digest = hashlib.sha256(content).hexdigest()
    workbook = load_workbook(_BytesReader(content), read_only=True, data_only=True)
    if len(workbook.worksheets) <= sheet_index:
        raise HTTPException(status_code=422, detail="Excel 中缺少业务工作表")
    worksheet = workbook.worksheets[sheet_index]
    sheet_name = worksheet.title
    parsed_rows: list[tuple[tuple[Any, ...], dict[str, Any]]] = []
    for index, row in enumerate(
        worksheet.iter_rows(min_row=data_start_row, values_only=True), start=data_start_row
    ):
        values = tuple(row)
        if not any(value not in (None, "") for value in values):
            continue  # 空行不进会话
        parsed_rows.append((values, parse_row(values, excel_row_no=index, sheet_name=sheet_name)))
    workbook.close()
    rows = [parsed for _, parsed in parsed_rows]
    if len(rows) > MAX_PREVIEW_ROWS:
        raise HTTPException(status_code=413, detail=f"预检单次最多 {MAX_PREVIEW_ROWS} 行，请拆分文件")

    states = [
        build_row_state(parsed["excel_row_no"], parsed, {}, list(values))
        for values, parsed in parsed_rows
    ]
    summary = _summary_with_cross(states, cross_row_issues(conn, states))
    session_id = secrets.token_urlsafe(24)
    expires_at = datetime.now() + timedelta(hours=SESSION_TTL_HOURS)
    conn.execute(
        text(
            """
            INSERT INTO legacy_import_session
              (id, created_by, expires_at, mode, source_file_name, source_sha256,
               parser_version, status, summary_json)
            VALUES
              (:id, :created_by, :expires_at, 'legacy_multi_value', :file_name, :sha256,
               :parser_version, 'pending', CAST(:summary AS JSON))
            """
        ),
        {
            "id": session_id,
            "created_by": user_id,
            "expires_at": expires_at,
            "file_name": file_name,
            "sha256": digest,
            "parser_version": PARSER_VERSION,
            "summary": json.dumps(summary, ensure_ascii=False),
        },
    )
    for values, row in parsed_rows:
        conn.execute(
            text(
                """
                INSERT INTO legacy_import_source
                  (session_id, sheet_name, excel_row_no, raw_json, parsed_json)
                VALUES
                  (:session_id, :sheet_name, :excel_row_no, CAST(:raw AS JSON), CAST(:parsed AS JSON))
                """
            ),
            {
                "session_id": session_id,
                "sheet_name": sheet_name,
                "excel_row_no": row["excel_row_no"],
                # 原始行用于提交时重建工作簿；只存库，不进日志、不进 Git。
                "raw": json.dumps(
                    {"values": [_jsonable(value) for value in values]}, ensure_ascii=False
                ),
                "parsed": json.dumps(row, ensure_ascii=False, default=_jsonable),
            },
        )
    return {
        "session_id": session_id,
        "source_sha256": digest,
        "expires_at": expires_at.isoformat(),
        "parser_version": PARSER_VERSION,
        "summary": summary,
    }


def load_session(conn: Connection, session_id: str, *, user_id: int, is_admin: bool = False) -> dict[str, Any]:
    row = conn.execute(
        text(
            """
            SELECT id, created_by, created_at, expires_at, status, source_file_name,
                   source_sha256, parser_version, summary_json, result_json
            FROM legacy_import_session WHERE id = :id
            """
        ),
        {"id": session_id},
    ).mappings().first()
    if row is None:
        raise HTTPException(status_code=404, detail="预检会话不存在")
    if int(row["created_by"]) != user_id and not is_admin:
        raise HTTPException(status_code=403, detail="无权访问他人的预检会话")
    if row["status"] == "pending" and row["expires_at"] and row["expires_at"] < datetime.now():
        conn.execute(
            text("UPDATE legacy_import_session SET status = 'expired' WHERE id = :id"),
            {"id": session_id},
        )
        row = {**dict(row), "status": "expired"}
    return dict(row)


def list_rows(conn: Connection, session_id: str, *, limit: int = 50, offset: int = 0) -> dict[str, Any]:
    total = conn.execute(
        text("SELECT COUNT(*) FROM legacy_import_source WHERE session_id = :id"),
        {"id": session_id},
    ).scalar()
    rows = conn.execute(
        text(
            """
            SELECT excel_row_no, parsed_json, resolution_json, resolved_at
            FROM legacy_import_source
            WHERE session_id = :id
            ORDER BY excel_row_no
            LIMIT :limit OFFSET :offset
            """
        ),
        {"id": session_id, "limit": limit, "offset": offset},
    ).mappings().all()
    items = []
    for row in rows:
        parsed = row["parsed_json"]
        parsed = json.loads(parsed) if isinstance(parsed, str) else parsed
        resolution = row["resolution_json"]
        resolution = json.loads(resolution) if isinstance(resolution, str) else resolution
        items.append(
            {
                "excel_row_no": int(row["excel_row_no"]),
                "parsed": parsed,
                "resolution": resolution,
                "resolved_at": row["resolved_at"].isoformat() if row["resolved_at"] else None,
            }
        )
    return {"total": int(total or 0), "items": items}


def apply_resolutions(
    conn: Connection,
    session_id: str,
    *,
    user: Any,
    resolutions: list[dict[str, Any]],
) -> dict[str, Any]:
    """写入逐项人工确认，并对相关行重新执行全部解析与业务校验。

    人工确认是**结构化数据**（每期日期/金额/票据各自是独立字段），不是
    "确认某个错误码"。确认值本身不合法（期次数、合计、日期、金额）时整次请求
    被拒绝，不会写入半成品；其他尚未处理的组保持阻断状态。
    """
    user_id = user.id
    session = load_session(conn, session_id, user_id=user_id)
    if session["status"] != "pending":
        raise HTTPException(status_code=409, detail=f"会话状态为 {session['status']}，不能修改")

    updated = 0
    collected_audits: list[dict[str, Any]] = []
    for item in resolutions:
        excel_row_no = int(item.get("excel_row_no") or 0)
        if not excel_row_no:
            continue
        stored = conn.execute(
            text(
                "SELECT parsed_json, resolution_json FROM legacy_import_source "
                "WHERE session_id = :id AND excel_row_no = :row"
            ),
            {"id": session_id, "row": excel_row_no},
        ).mappings().first()
        if stored is None:
            raise HTTPException(status_code=404, detail=f"会话中没有第 {excel_row_no} 行")

        try:
            revision = parse_row_resolution(
                item.get("resolution"), known_finance_groups=FINANCE_SOURCES
            )
        except RevisionError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        if revision is None:
            continue

        parsed = _json_load(stored["parsed_json"]) or {}
        existing = _json_load(stored["resolution_json"]) or {}
        audits: list[dict[str, Any]] = []
        for name, entry_revision in revision.finance.items():
            label = FINANCE_LABELS.get(name, name)
            _, group_issues, audit = validate_finance_revision(
                (parsed.get("finance") or {}).get(name) or {}, entry_revision, label=label
            )
            blocking = [issue for issue in group_issues if issue.blocking]
            if blocking:
                raise HTTPException(
                    status_code=422,
                    detail="；".join(issue.message for issue in blocking),
                )
            if audit:
                audits.append(audit)
                collected_audits.append({**audit, "excel_row_no": excel_row_no})

        payload = _merge_resolution_payload(existing, revision)
        payload["audit"] = {
            "finance": audits,
            "recorded_by": user_id,
            "recorded_at": datetime.now().isoformat(timespec="seconds"),
        }
        conn.execute(
            text(
                """
                UPDATE legacy_import_source
                SET resolution_json = CAST(:resolution AS JSON),
                    resolved_by = :user_id,
                    resolved_at = CURRENT_TIMESTAMP
                WHERE session_id = :id AND excel_row_no = :row
                """
            ),
            {
                "resolution": json.dumps(payload, ensure_ascii=False, default=_jsonable),
                "user_id": user_id,
                "id": session_id,
                "row": excel_row_no,
            },
        )
        updated += 1

    summary = recompute_summary(conn, session_id)
    conn.execute(
        text(
            "UPDATE legacy_import_session SET summary_json = CAST(:summary AS JSON) WHERE id = :id"
        ),
        {"summary": json.dumps(summary, ensure_ascii=False), "id": session_id},
    )
    if updated:
        # 审计只留元数据：会话、行号、涉及哪些财务组与合计修正。不含原始业务内容。
        write_operation_log(
            conn,
            user,
            "订单管理",
            "resolve_legacy_import",
            f"人工确认旧台账预检会话 {session_id} 的 {updated} 行；仍有 {summary['blocking_rows']} 行阻断",
            after={
                "session_id": session_id,
                "rows": [int(item.get("excel_row_no") or 0) for item in resolutions][:200],
                "audit": collected_audits[:200],
            },
        )
    return {"updated": updated, "summary": summary}


def _merge_resolution_payload(existing: dict[str, Any], revision: RowResolution) -> dict[str, Any]:
    """本次修正与前几次已确认的内容合并：没提到的字段保持原确认。"""
    payload: dict[str, Any] = {
        key: existing[key] for key in ("order_no", "manager", "finance") if existing.get(key)
    }
    if revision.order_no is not None:
        payload["order_no"] = {"history": list(revision.order_no.history)}
    if revision.manager is not None:
        payload["manager"] = {"history": list(revision.manager.history)}
    finance = dict(payload.get("finance") or {})
    for name, entry in revision.finance.items():
        finance[name] = entry.model_dump(mode="json")
    if finance:
        payload["finance"] = finance
    return payload


def _json_load(value: Any) -> Any:
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return None
    return value


def evaluate_session(conn: Connection, session_id: str) -> tuple[list[RowState], list[Issue]]:
    """读取会话的全部行，按"解析 + 人工确认"重算状态与跨行问题。"""
    rows = conn.execute(
        text(
            "SELECT excel_row_no, raw_json, parsed_json, resolution_json "
            "FROM legacy_import_source WHERE session_id = :id ORDER BY excel_row_no"
        ),
        {"id": session_id},
    ).mappings().all()
    states = [
        build_row_state(
            int(row["excel_row_no"]),
            _json_load(row["parsed_json"]) or {},
            _json_load(row["resolution_json"]),
            list((_json_load(row["raw_json"]) or {}).get("values") or []),
        )
        for row in rows
    ]
    return states, cross_row_issues(conn, states)


def load_states(conn: Connection, session_id: str) -> list[RowState]:
    states, _ = evaluate_session(conn, session_id)
    return states


def _project_id_by_code(conn: Connection, project_code: str) -> int | None:
    value = conn.execute(
        text("SELECT id FROM project WHERE project_code = :code AND deleted_at IS NULL"),
        {"code": project_code},
    ).scalar()
    return int(value) if value else None


def _order_components(states: list[RowState]) -> list[list[RowState]]:
    """按"共享任一订单号"把行聚成同一订单的连通分量。

    同一订单的多行必须给出可归一的链；只有一个号的行各自独立（新订单）。
    """
    rows = [state for state in states if state.order_chain]
    if not rows:
        return []
    parent = list(range(len(rows)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def union(left: int, right: int) -> None:
        root_left, root_right = find(left), find(right)
        if root_left != root_right:
            parent[root_right] = root_left

    by_number: dict[str, int] = {}
    for position, state in enumerate(rows):
        for number in state.order_chain:
            if number in by_number:
                union(by_number[number], position)
            else:
                by_number[number] = position

    groups: dict[int, list[RowState]] = {}
    for position, state in enumerate(rows):
        groups.setdefault(find(position), []).append(state)
    return list(groups.values())


def cross_row_issues(conn: Connection, states: list[RowState]) -> list[Issue]:
    """跨行与库内现状的一致性校验（预检与提交前各执行一次）。

    - 同一框架的负责人链必须能归一，且与库内现任/历史一致或不冲突；
    - 同一订单的别名链必须能归一，且不与同框架另一个订单的当前号或历史号冲突。
    """
    issues: list[Issue] = []
    issues.extend(_manager_chain_issues(conn, states))
    issues.extend(_order_chain_issues(conn, states))
    return issues


def _manager_chain_issues(conn: Connection, states: list[RowState]) -> list[Issue]:
    issues: list[Issue] = []
    groups: dict[str, list[RowState]] = {}
    for state in states:
        code = str(state.parsed.get("project_code") or "").strip()
        if code and state.manager_chain:
            groups.setdefault(code, []).append(state)

    for code, group in sorted(groups.items()):
        first = group[0]
        outcome = merge_manager_history(
            ["/".join(state.manager_chain) for state in group], project_code=code
        )
        if outcome.blocking:
            issues.extend(
                issue.with_location(row=first.excel_row_no) for issue in outcome.issues
            )
            continue
        project_id = _project_id_by_code(conn, code)
        if project_id is None:
            continue
        merged = merge_manager_chain(conn, project_id, outcome.history)
        if merged.blocked:
            issues.append(
                Issue(
                    code=CURRENT_MANAGER_CONFLICT,
                    message=f"项目 {code}：{merged.reason}",
                    row=first.excel_row_no,
                )
            )
    return issues


def _order_chain_issues(conn: Connection, states: list[RowState]) -> list[Issue]:
    issues: list[Issue] = []
    for group in _order_components(states):
        first = group[0]
        chain, merge_issues = merge_chain_sequences(
            [state.order_chain for state in group],
            what="订单号",
            entity="订单",
            current_conflict_code=ORDER_ALIAS_CONFLICT,
        )
        if chain is None:
            issues.extend(
                issue.with_location(row=first.excel_row_no) for issue in merge_issues
            )
            continue

        code = str(first.parsed.get("project_code") or "").strip()
        project_id = _project_id_by_code(conn, code) if code else None
        if project_id is None:
            continue

        owner = find_order_by_number(conn, project_id, chain[-1])
        conflicts = conflicts_in_project(
            conn, project_id, chain, exclude_sales_order_id=owner
        )
        if conflicts:
            numbers = "、".join(sorted(conflicts))
            issues.append(
                Issue(
                    code=ORDER_ALIAS_CONFLICT,
                    message=(
                        f"订单号 {numbers} 在同一框架内已属于另一个订单（当前号或历史别名），"
                        "整批不能提交；系统不会自动合并两个订单的财务记录，"
                        "若确属改号请走专门的改号流程"
                    ),
                    row=first.excel_row_no,
                )
            )
            continue
        if owner is None:
            continue
        merged = merge_order_chain(conn, owner, chain)
        if merged.blocked:
            issues.append(
                Issue(
                    code=ORDER_HISTORY_CONFLICT,
                    message=merged.reason,
                    row=first.excel_row_no,
                )
            )
    return issues


def _summary_with_cross(states: list[RowState], cross_issues: list[Issue]) -> dict[str, Any]:
    summary = summarize_states(states)
    blocking_rows = {state.excel_row_no for state in states if state.blocking}
    blocking_rows |= {
        int(issue.row) for issue in cross_issues if issue.blocking and issue.row
    }
    summary["blocking_rows"] = len(blocking_rows)
    summary["cross_row_issues"] = [_issue_payload(issue) for issue in cross_issues]
    summary["comparable"] = len(blocking_rows) == 0
    return summary


def recompute_summary(conn: Connection, session_id: str) -> dict[str, Any]:
    """按"解析结果 + 人工确认"重算阻断情况；确认值本身要重新通过全部校验。"""
    states, cross_issues = evaluate_session(conn, session_id)
    return _summary_with_cross(states, cross_issues)


def finish_session(conn: Connection, session_id: str, *, result: dict[str, Any]) -> None:
    conn.execute(
        text(
            """
            UPDATE legacy_import_session
            SET status = 'committed', committed_at = CURRENT_TIMESTAMP, result_json = CAST(:result AS JSON)
            WHERE id = :id
            """
        ),
        {"result": json.dumps(result, ensure_ascii=False, default=_jsonable), "id": session_id},
    )


def check_commit_ready(conn: Connection, session: dict[str, Any], *, content: bytes) -> dict[str, Any]:
    """提交前复核：状态、文件摘要、阻断问题。

    阻断判定与预检阶段**同一套逻辑**：人工确认过的内容要重新通过期次数、合计、
    日期、金额校验，还要通过与库内现状的别名/负责人一致性检查。
    """
    if session["status"] == "committed":
        stored = session.get("result_json")
        return {"already_committed": True, "result": json.loads(stored) if isinstance(stored, str) else stored}
    if session["status"] != "pending":
        raise HTTPException(status_code=409, detail=f"会话状态为 {session['status']}，不能提交")

    digest = hashlib.sha256(content).hexdigest()
    if digest != session["source_sha256"]:
        raise HTTPException(
            status_code=409,
            detail="上传的文件与预检时的文件不一致（内容已改变），请重新预检",
        )

    summary = recompute_summary(conn, session["id"])
    if not summary["comparable"]:
        raise HTTPException(
            status_code=422,
            detail=f"仍有 {summary['blocking_rows']} 行存在阻断问题，请先处理后再提交",
        )
    return {"already_committed": False, "summary": summary}


def commit_session(
    conn: Connection,
    *,
    session_id: str,
    user: Any,
    content: bytes,
    is_admin: bool = False,
) -> dict[str, Any]:
    """提交已通过预检的会话：写锁内复核 → 规范化 → 复用导入服务写入 → 追加期次与历史。

    整个过程在调用方的写锁事务内完成：任一期财务、任何历史修改或会话状态出错，
    业务写入与历史记录一起回滚。重复提交同一会话只返回原结果，不再新增业务记录。
    """
    from .importer import import_excel  # 延迟导入，避免模块级循环依赖

    session = load_session(conn, session_id, user_id=user.id, is_admin=is_admin)
    ready = check_commit_ready(conn, session, content=content)
    if ready.get("already_committed"):
        return {"already_committed": True, **(ready.get("result") or {})}

    states = load_states(conn, session_id)
    workbook_bytes, row_map = build_normalized_workbook(states)
    line_ids: dict[int, int] = {}
    result = import_excel(
        conn,
        reset=False,
        workbook_bytes=workbook_bytes,
        source_file_name=session["source_file_name"],
        user=user,
        strict_template=True,
        line_id_map=line_ids,
    )
    if result["failed_rows"]:
        raise HTTPException(
            status_code=422,
            detail=f"提交失败，已整批回滚：{'；'.join(result.get('errors') or [])}",
        )

    phases = append_phases(conn, states, line_ids, row_map)
    history_records = register_histories(conn, states, line_ids, row_map)

    payload = {
        "session_id": session_id,
        "success_rows": result["success_rows"],
        "skipped_rows": result["skipped_rows"],
        "phases": phases,
        "history_records": history_records,
        "file_sha256": session["source_sha256"],
    }
    finish_session(conn, session_id, result=payload)
    # 审计只记录可核对的元数据：会话、文件摘要、条数、期次数。不含原始业务内容。
    write_operation_log(
        conn,
        user,
        "订单管理",
        "commit_legacy_import",
        (
            f"提交旧台账预检会话：成功 {result['success_rows']} 行，"
            f"写入期次 {sum(phases.values())} 条，历史记录 {history_records} 条"
        ),
        after={
            "session_id": session_id,
            "file_sha256": session["source_sha256"],
            "success_rows": result["success_rows"],
            "skipped_rows": result["skipped_rows"],
            "phases": phases,
        },
    )
    return {"already_committed": False, **payload}


# 业务 → (日期列, 票据列, 金额列)。写期次时按这张表拼列名。
PHASE_TABLE_COLUMNS: dict[str, tuple[str, str, str]] = {
    "sales_invoice": ("invoice_date", "invoice_no", "invoice_amount"),
    "sales_receipt": ("receipt_date", "payment_notice_no", "receipt_amount"),
    "purchase_payment": ("payment_date", "payment_voucher_no", "payment_amount"),
    "purchase_invoice": ("received_invoice_date", "invoice_no", "invoice_amount"),
    "finance_invoice_check": ("received_invoice_date", "voucher_code", "received_invoice_amount"),
    "warehouse_entry": ("warehouse_date", "voucher_no", "warehouse_amount"),
}

# 每组财务在模板里占用的**全部**列。规范化工件要把这些列整体清空：
# 只要漏掉一列（例如开票组的发票号列、或付款组的到期付款日），importer 的固定列
# 逻辑就会先写一期（哪怕只有到期付款日），追加逻辑再写 N 期，期次号与条数都错。
FINANCE_GROUP_COLUMNS: dict[str, tuple[int, ...]] = {
    "sales_invoice": (73, 74, 75, 76, 77),
    "sales_receipt": (79, 80, 81, 82, 83, 84, 85, 86),
    "purchase_payment": (54, 55, 56, 57, 58, 59, 60, 61),
    "purchase_invoice": (44, 45, 46),
    "finance_invoice_check": (50, 51, 52, 53),
    "warehouse_entry": (47, 48, 49),
}

# 到期付款日不参与多值解析，作为采购付款各期共用的行级值保留（原始行里读）。
DUE_PAYMENT_DATE_COLUMN = 54

PHASE_SOURCE_COLUMNS: tuple[int, ...] = tuple(
    column for columns in FINANCE_GROUP_COLUMNS.values() for column in columns
)


def build_normalized_workbook(states: list[RowState]) -> tuple[bytes, dict[int, int]]:
    """按"原始行 + 人工确认"重建一份规范化工件。

    订单号写当前号（链尾），客户经理写现任；财务源列全部清空，改由
    append_phases 按确认后的期次写入，因此支持任意期数而不是模板的固定两期。

    返回（工作簿字节、原始 Excel 行号 → 规范化行号）。原始文件中间可能有空行，
    规范化后是紧凑排列的；调用方必须用这张映射把明细 id 对回原始行，
    否则空行之后的财务与历史会整体错位。
    """
    from .ledger_excel import TEMPLATE_HEADERS, template_bytes

    workbook = load_workbook(_BytesReader(template_bytes()))
    worksheet = workbook.worksheets[0]
    total_columns = len(TEMPLATE_HEADERS)
    for excel_row in range(3, 3 + len(states) + 1):
        for column in range(1, total_columns + 1):
            worksheet.cell(excel_row, column).value = None

    row_map: dict[int, int] = {}
    for offset, state in enumerate(states):
        normalized_row = 3 + offset
        row_map[state.excel_row_no] = normalized_row
        values = list(state.raw_values)
        values += [None] * (total_columns - len(values))
        for column in PHASE_SOURCE_COLUMNS:
            if 1 <= column <= total_columns:
                values[column - 1] = None
        if state.order_chain:
            values[ORDER_NO_COLUMN - 1] = state.order_chain[-1]
        if state.manager_chain:
            values[MANAGER_COLUMN - 1] = state.manager_chain[-1]
        for column, value in enumerate(values[:total_columns], start=1):
            worksheet.cell(normalized_row, column, _restore_value(value))

    buffer = _BytesWriter()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue(), row_map


def append_phases(
    conn: Connection,
    states: list[RowState],
    line_ids: dict[int, int],
    row_map: dict[int, int],
) -> dict[str, int]:
    """把确认后的每一期财务按顺序写入对应期次表，期号连续分配、不设上限。

    写入前对每一期重新执行与单条录入一致的业务校验（金额非负且最多两位小数、
    日期真实存在且年份不超过 2099、票据号长度），任一期不合法就整批回滚。
    """
    written: dict[str, int] = {}
    for state in states:
        order_line_id = line_ids.get(row_map.get(state.excel_row_no, -1))
        if not order_line_id:
            continue
        line_groups: set[str] = set()
        for business, entry in state.finance.items():
            columns = PHASE_TABLE_COLUMNS.get(business)
            if not columns:
                continue
            label = FINANCE_LABELS.get(business, business)
            date_column, document_column, amount_column = columns
            phases = entry.get("phases") or []
            phase_issues: list[Issue] = []
            for position, phase in enumerate(phases, start=1):
                phase_issues.extend(
                    phase_business_issues(phase, label=label, position=position)
                )
            if phase_issues:
                raise HTTPException(
                    status_code=422,
                    detail="；".join(issue.message for issue in phase_issues),
                )
            for phase_no, phase in enumerate(phases, start=1):
                conn.execute(
                    text(
                        f"""
                        INSERT INTO {business}
                          (order_line_id, phase_no, {date_column}, {date_column}_text,
                           {document_column}, {amount_column})
                        VALUES
                          (:order_line_id, :phase_no, :date_value, :date_text,
                           :document_value, :amount_value)
                        """
                    ),
                    {
                        "order_line_id": order_line_id,
                        "phase_no": phase_no,
                        "date_value": phase.get("date"),
                        "date_text": phase.get("date"),
                        "document_value": phase.get("document_no"),
                        "amount_value": phase.get("amount"),
                    },
                )
                if business == "purchase_payment":
                    _set_due_payment_date(conn, order_line_id, phase_no, state)
                written[business] = written.get(business, 0) + 1
            if phases:
                line_groups.add(business)

        # 到期付款日是行级单值。整行只填了它、没有付款日期与金额时，也要保留下来，
        # 否则规范化清空源列之后这个字段会被静默丢掉。
        due = _due_payment_date(state)
        if due not in (None, "") and "purchase_payment" not in line_groups:
            conn.execute(
                text(
                    "INSERT INTO purchase_payment (order_line_id, phase_no, due_payment_date) "
                    "VALUES (:order_line_id, 1, :due)"
                ),
                {"order_line_id": order_line_id, "due": due},
            )
            written["purchase_payment"] = written.get("purchase_payment", 0) + 1
    return written


def _due_payment_date(state: RowState) -> Any:
    index = DUE_PAYMENT_DATE_COLUMN - 1
    if index >= len(state.raw_values):
        return None
    return _restore_value(state.raw_values[index])


def _set_due_payment_date(
    conn: Connection, order_line_id: int, phase_no: int, state: RowState
) -> None:
    """到期付款日在模板里是行级单值，写入该期时一并带上。"""
    due = _due_payment_date(state)
    if due in (None, ""):
        return
    conn.execute(
        text(
            "UPDATE purchase_payment SET due_payment_date = :due "
            "WHERE order_line_id = :order_line_id AND phase_no = :phase_no"
        ),
        {"due": due, "order_line_id": order_line_id, "phase_no": phase_no},
    )


def _remember_chain(
    chains: dict[int, list[str]], entity_id: int, chain: list[str], *, what: str
) -> None:
    """同一实体多行的链必须一致；不一致说明预检校验被绕过，整批回滚。"""
    existing = chains.get(entity_id)
    if existing is None:
        chains[entity_id] = list(chain)
        return
    if existing != chain:
        raise HTTPException(
            status_code=422,
            detail=f"{what}在同一实体下出现多条不同的历史链，已整批回滚，请重新预检",
        )


def register_histories(
    conn: Connection,
    states: list[RowState],
    line_ids: dict[int, int],
    row_map: dict[int, int],
) -> int:
    """登记订单号别名链与负责人交接链（链尾为当前值），并同步主记录的当前值。

    写法是"与库内已确认的链合并"：文件只填当前值时不截短已有历史，文件补充了
    缺失历史时按顺序补上；现任不同或链条冲突则整批回滚（改号/交接走另行授权的流程）。
    """
    order_chains: dict[int, list[str]] = {}
    project_chains: dict[int, list[str]] = {}
    for state in states:
        order_line_id = line_ids.get(row_map.get(state.excel_row_no, -1))
        if not order_line_id:
            continue
        ids = conn.execute(
            text(
                """
                SELECT so.id AS sales_order_id, p.id AS project_id
                FROM order_line ol
                JOIN sales_order so ON so.id = ol.sales_order_id
                JOIN project p ON p.id = so.project_id
                WHERE ol.id = :order_line_id
                """
            ),
            {"order_line_id": order_line_id},
        ).mappings().first()
        if ids is None:
            continue
        sales_order_id = int(ids["sales_order_id"])
        project_id = int(ids["project_id"])

        if state.order_chain:
            _remember_chain(order_chains, sales_order_id, state.order_chain, what="订单号")
        if state.manager_chain:
            _remember_chain(project_chains, project_id, state.manager_chain, what="客户经理")

    written = 0
    for sales_order_id, chain in order_chains.items():
        # 只补不删，且现任必须与库内一致；改号或历史冲突在这里整批回滚。
        merged = merge_order_chain(conn, sales_order_id, chain)
        if merged.blocked:
            raise HTTPException(status_code=422, detail=merged.reason)
        register_order_numbers(conn, sales_order_id, merged.chain)
        conn.execute(
            text("UPDATE sales_order SET order_no = :order_no WHERE id = :id"),
            {"order_no": merged.chain[-1], "id": sales_order_id},
        )
        written += 1

    for project_id, chain in project_chains.items():
        merged = merge_manager_chain(conn, project_id, chain)
        if merged.blocked:
            raise HTTPException(status_code=422, detail=merged.reason)
        register_manager_history(conn, project_id, merged.chain)
        conn.execute(
            text("UPDATE project SET account_manager = :name WHERE id = :project_id"),
            {"name": merged.chain[-1], "project_id": project_id},
        )
        written += 1
    return written


class _BytesWriter:
    def __init__(self) -> None:
        self._chunks: list[bytes] = []

    def write(self, data: bytes) -> int:
        self._chunks.append(bytes(data))
        return len(data)

    def getvalue(self) -> bytes:
        return b"".join(self._chunks)

    def flush(self) -> None:
        return None

    def seekable(self) -> bool:
        return False

    def tell(self) -> int:
        return sum(len(chunk) for chunk in self._chunks)


class _BytesReader:
    """openpyxl 只接受文件路径或类文件对象。"""

    def __init__(self, content: bytes) -> None:
        self._buffer = memoryview(content)
        self._position = 0

    def read(self, size: int = -1) -> bytes:
        if size is None or size < 0:
            chunk = bytes(self._buffer[self._position :])
            self._position = len(self._buffer)
            return chunk
        chunk = bytes(self._buffer[self._position : self._position + size])
        self._position += len(chunk)
        return chunk

    def seek(self, offset: int, whence: int = 0) -> int:
        if whence == 0:
            self._position = offset
        elif whence == 1:
            self._position += offset
        else:
            self._position = len(self._buffer) + offset
        return self._position

    def tell(self) -> int:
        return self._position

    def seekable(self) -> bool:
        return True
