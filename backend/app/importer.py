from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.engine import Connection

from .audit import write_operation_log
from .auth import CurrentUser, can_access_department
from .config import DOCS_DIR
from .line_identity import find_duplicate_line
from .ledger_excel import SAMPLE_ORDER_NO, SAMPLE_PROJECT_CODE, TEMPLATE_HEADERS, is_template_sample_row
from .validation import validate_business_date


BUSINESS_SHEET_INDEX = 2
HEADER_ROW = 3
DATA_START_ROW = 5
LATEST_LAYOUT_HEADERS = frozenset({"销售税率", "采购税率", "物资/服务名称"})
REQUIRED_BUSINESS_HEADERS = frozenset({"项目编号", "订单号", "项目名称"})


def _json_default(value: Any) -> str | float | int | None:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value) if value is not None else None


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return validate_business_date(value.date())
    if isinstance(value, date):
        return validate_business_date(value)
    if isinstance(value, str) and value.strip():
        normalized = value.strip().replace("/", "-").replace(".", "-")
        for date_format in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return validate_business_date(datetime.strptime(normalized, date_format).date())
            except ValueError:
                continue
        raise ValueError(f"日期格式错误：{value}，请使用 YYYY-MM-DD")
    return None


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    text_value = str(value).strip()
    return text_value or None


def _as_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except (TypeError, ValueError, ArithmeticError):
        return None


def _as_tax_rate(value: Any) -> Decimal | None:
    rate = _as_decimal(value)
    if rate is None:
        return None
    if Decimal("0") < rate <= Decimal("1"):
        rate *= Decimal("100")
    if rate < 0 or rate > 100:
        raise ValueError("税率必须在0到100之间")
    return rate


def _calculated_prices(
    quantity: Decimal | None,
    tax_rate: Decimal | None,
    unit_price_no_tax: Decimal | None,
    unit_price: Decimal | None,
) -> tuple[Decimal | None, Decimal | None, Decimal | None]:
    if tax_rate is None:
        return unit_price, None, None
    if unit_price_no_tax is not None:
        unit_price = (unit_price_no_tax * (Decimal("1") + tax_rate / Decimal("100"))).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
    amount_no_tax = None if quantity is None or unit_price_no_tax is None else (quantity * unit_price_no_tax).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    amount = None if quantity is None or unit_price is None else (quantity * unit_price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return unit_price, amount_no_tax, amount


def _row_value(row: tuple[Any, ...], position: int) -> Any:
    index = position - 1
    return row[index] if 0 <= index < len(row) else None


# 业务载荷列（模板布局下的绝对列号）：部门、订单日期、项目名称、货物名称、
# 规格型号、数量、单价、订单价值、采购厂商。只填了项目编号或订单号、
# 这些列全空的行不是业务数据。
BUSINESS_PAYLOAD_POSITIONS = (3, 6, 14, 15, 16, 18, 21, 23, 24)


def _has_business_payload(row: tuple[Any, ...]) -> bool:
    return any(_as_text(_row_value(row, position)) for position in BUSINESS_PAYLOAD_POSITIONS)


# 框架项目级共享字段：同一框架下全部订单、子项目和明细必须一致。
# 客户经理属于框架项目（旧单元格里的“甲/乙/丙”是历次交接，最后值为现任）。
SHARED_PROJECT_FIELDS: tuple[tuple[str, int, int, str], ...] = (
    ("department", 3, 3, "部门"),
    ("branch_company", 4, 4, "分公司"),
    ("account_manager", 5, 5, "客户经理"),
    ("team_level3_name", 9, 9, "三级团队"),
)

# 子项目级共享字段：同一订单的同一子项目内必须一致；不同子项目可以不同，
# 包括同一订单下的不同子项目。这些字段不能校验在框架项目上，否则同框架下
# 不同子项目的合法差异会被整批拦下。
SUB_PROJECT_FIELDS: tuple[tuple[str, int, int, str], ...] = (
    ("customer_unit_name", 10, 10, "客户单位"),
    ("end_user_name", 11, 11, "最终用户"),
    ("regional_platform", 12, 12, "区域平台"),
)


def _is_latest_layout(headers: list[Any] | tuple[Any, ...]) -> bool:
    header_names = {str(header or "").strip() for header in headers}
    return bool(LATEST_LAYOUT_HEADERS & header_names)


def _detect_business_header_row(worksheet, *, max_scan_rows: int = 10) -> int:
    for row_no, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
        start=1,
    ):
        names = {str(value or "").strip() for value in row}
        if REQUIRED_BUSINESS_HEADERS.issubset(names):
            return row_no
    raise ValueError("Excel 前10行中未找到项目编号、订单号、项目名称等业务表头")


def _column_position(latest_layout: bool, latest: int, legacy: int | None = None) -> int:
    return latest if latest_layout else int(legacy or 0)


def _execute_scalar(conn: Connection, sql: str, params: dict[str, Any]) -> int:
    result = conn.execute(text(sql), params)
    if result.lastrowid:
        return int(result.lastrowid)
    return int(conn.execute(text("SELECT LAST_INSERT_ID()")).scalar() or 0)


def _reset_business_data(conn: Connection) -> None:
    tables = [
        "sales_receipt",
        "sales_invoice",
        "sales_contract",
        "purchase_payment",
        "finance_payment_entry",
        "finance_invoice_check",
        "warehouse_entry",
        "purchase_invoice",
        "purchase_contract",
        "delivery_record",
        "purchase_info",
        "order_line",
        "sub_project",
        "sales_order",
        "project",
        "ledger_raw_row",
        "import_batch",
    ]
    conn.execute(text("SET FOREIGN_KEY_CHECKS=0"))
    for table in tables:
        conn.execute(text(f"TRUNCATE TABLE {table}"))
    conn.execute(text("SET FOREIGN_KEY_CHECKS=1"))


def _diff_row_against_stored(
    row: tuple[Any, ...],
    position,
    stored: Mapping[str, Any],
    fields: tuple[tuple[str, int, int, str], ...],
) -> list[str]:
    """比较 Excel 行与库中记录在这些字段上的取值，返回中文冲突描述。"""
    conflicts: list[str] = []
    for column, latest, legacy, label in fields:
        excel_value = _as_text(_row_value(row, position(latest, legacy)))
        stored_value = _as_text(stored[column])
        # 只在两边都有值且不同时报冲突：Excel 留空表示不改动，台账为空表示可补充。
        if excel_value and stored_value and excel_value != stored_value:
            conflicts.append(f"{label}（台账为“{stored_value}”，Excel 为“{excel_value}”）")
    return conflicts


def _find_shared_field_conflicts(
    conn: Connection,
    project_code: str,
    row: tuple[Any, ...],
    position,
    user: CurrentUser | None,
) -> tuple[int | None, list[str]]:
    """校验导入行与台账中已有框架项目的兼容性（客户经理、分公司、部门、三级团队）。

    返回（已有框架项目 id, 冲突字段描述）。若项目不存在则返回 (None, [])。
    数据库原部门对当前账号不可录入时直接抛 PermissionError——这挡住“拿别的
    部门的项目编号、把部门填成自己能访问的部门”来绕过部门边界。
    """
    existing = conn.execute(
        text(
            """
            SELECT id, department, branch_company, account_manager, team_level3_name
            FROM project
            WHERE project_code = :project_code AND deleted_at IS NULL
            """
        ),
        {"project_code": project_code},
    ).mappings().first()
    if existing is None:
        return None, []

    if user is not None and not can_access_department(user, existing["department"], require_entry=True):
        raise PermissionError(
            f"项目 {project_code} 属于部门“{existing['department'] or '空'}”，"
            "无权向该部门的项目导入数据"
        )

    return int(existing["id"]), _diff_row_against_stored(row, position, existing, SHARED_PROJECT_FIELDS)


def _find_or_create_sub_project(
    conn: Connection,
    sales_order_id: int,
    row: tuple[Any, ...],
    position,
    project_name: str | None,
    excel_row_no: int,
) -> tuple[int, list[str]]:
    """按（订单, 子项目名称）定位子项目，必要时新建。

    返回（子项目 id, 冲突描述）。客户单位/最终用户/区域平台存在子项目上：
    同一订单下不同子项目可以有不同值，这与框架共享字段是两回事。
    """
    name = project_name or ""
    existing = conn.execute(
        text(
            """
            SELECT id, customer_unit_name, end_user_name, regional_platform
            FROM sub_project
            WHERE sales_order_id = :sales_order_id AND name = :name AND deleted_at IS NULL
            """
        ),
        {"sales_order_id": sales_order_id, "name": name},
    ).mappings().first()

    if existing is not None:
        return int(existing["id"]), _diff_row_against_stored(row, position, existing, SUB_PROJECT_FIELDS)

    values = {
        column: _as_text(_row_value(row, position(latest, legacy)))
        for column, latest, legacy, _label in SUB_PROJECT_FIELDS
    }
    sub_project_id = _execute_scalar(
        conn,
        """
        INSERT INTO sub_project
          (sales_order_id, name, customer_unit_name, end_user_name, regional_platform)
        VALUES
          (:sales_order_id, :name, :customer_unit_name, :end_user_name, :regional_platform)
        ON DUPLICATE KEY UPDATE id = LAST_INSERT_ID(id)
        """,
        {"sales_order_id": sales_order_id, "name": name, **values},
    )
    if not sub_project_id:
        raise ValueError(f"第 {excel_row_no} 行无法为子项目“{name or '（空名称）'}”建立档案")
    return sub_project_id, []


def import_excel(
    conn: Connection,
    reset: bool = True,
    *,
    workbook_bytes: bytes | None = None,
    source_file_name: str | None = None,
    user: CurrentUser | None = None,
    strict_template: bool = False,
) -> dict[str, Any]:
    workbook_path: Path | None = None
    if workbook_bytes is None:
        workbooks = sorted(DOCS_DIR.glob("2026*.xlsx"))
        if not workbooks:
            raise FileNotFoundError("docs 目录中未找到 2026*.xlsx 业务台账文件")
        workbook_path = workbooks[0]
        source_file_name = workbook_path.name
        workbook_source: Path | BytesIO = workbook_path
        sheet_index = BUSINESS_SHEET_INDEX
        header_row = HEADER_ROW
        data_start_row = DATA_START_ROW
    else:
        workbook_source = BytesIO(workbook_bytes)
        source_file_name = source_file_name or "市场部业务台账.xlsx"
        sheet_index = 0
        header_row = 2 if strict_template else 0
        data_start_row = 3 if strict_template else 0
    if reset:
        _reset_business_data(conn)

    workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    if len(workbook.worksheets) <= sheet_index:
        raise ValueError("Excel 中缺少业务台账工作表")
    worksheet = workbook.worksheets[sheet_index]
    if not strict_template and header_row == 0:
        header_row = _detect_business_header_row(worksheet)
        data_start_row = header_row + 1
    headers = list(next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
    normalized_headers = [str(header or "").strip() for header in headers]
    if strict_template and normalized_headers != TEMPLATE_HEADERS:
        raise ValueError("Excel 表头与“市场部业务台账模板”不一致，请重新下载模板后填写")
    template_layout = strict_template and normalized_headers == TEMPLATE_HEADERS
    latest_layout = _is_latest_layout(headers)

    def position(latest: int, legacy: int | None = None) -> int:
        return _column_position(latest_layout, latest, legacy)

    batch_id = _execute_scalar(
        conn,
        """
        INSERT INTO import_batch
          (source_file_name, source_sheet_name, header_row_no, data_start_row_no, status, uploaded_by)
        VALUES
          (:file_name, :sheet_name, :header_row_no, :data_start_row_no, 'parsing', :uploaded_by)
        """,
        {
            "file_name": source_file_name,
            "sheet_name": worksheet.title,
            "header_row_no": header_row,
            "data_start_row_no": data_start_row,
            "uploaded_by": user.id if user else None,
        },
    )

    success_rows = 0
    failed_rows = 0
    skipped_rows = 0
    error_messages: list[str] = []

    for excel_row_no, row in enumerate(
        worksheet.iter_rows(min_row=data_start_row, values_only=True),
        start=data_start_row,
    ):
        if strict_template and is_template_sample_row(row):
            continue
        project_code = _as_text(_row_value(row, position(2, 2)))
        order_no = _as_text(_row_value(row, position(13, 13)))
        if not any(value not in (None, "") for value in row):
            continue
        if strict_template and (project_code == SAMPLE_PROJECT_CODE or order_no == SAMPLE_ORDER_NO):
            raise ValueError(f"第 {excel_row_no} 行仍包含示例占位内容，请完整替换项目编号和销售订单号")
        if not project_code or not order_no:
            message = f"第 {excel_row_no} 行缺少项目编号或订单号"
            # 生成式 AI 工具会在表格末尾追加 AIGC 标识行（浅色字体，只有前两列有
            # 内容，没有订单号）。这类行没有任何业务载荷，按非业务行跳过并计数；
            # 有业务载荷却缺编号的属于用户漏填，仍整批报错提示修改。
            if strict_template and _has_business_payload(row):
                raise ValueError(message)
            skipped_rows += 1
            continue

        row_dict = {
            str(headers[i] or f"column_{i + 1}"): _json_default(value)
            for i, value in enumerate(row)
        }
        raw_json = json.dumps(row_dict, ensure_ascii=False, default=_json_default)
        row_hash = hashlib.sha256(raw_json.encode("utf-8")).hexdigest()

        try:
            department = _as_text(_row_value(row, position(3, 3)))
            if user and not can_access_department(user, department, require_entry=True):
                raise PermissionError(f"无权向部门“{department or '空'}”导入数据")

            raw_row_id = _execute_scalar(
                conn,
                """
                INSERT INTO ledger_raw_row
                  (import_batch_id, excel_row_no, row_hash, raw_json, parse_status, project_code, order_no)
                VALUES
                  (:batch_id, :excel_row_no, :row_hash, CAST(:raw_json AS JSON), 'success', :project_code, :order_no)
                ON DUPLICATE KEY UPDATE id = LAST_INSERT_ID(id), parse_status = 'success'
                """,
                {
                    "batch_id": batch_id,
                    "excel_row_no": excel_row_no,
                    "row_hash": row_hash,
                    "raw_json": raw_json,
                    "project_code": project_code,
                    "order_no": order_no,
                },
            )

            existing_project_id, conflicts = _find_shared_field_conflicts(
                conn, project_code, row, position, user
            )
            if conflicts:
                raise ValueError(
                    f"第 {excel_row_no} 行项目 {project_code} 已在台账中，"
                    "以下项目级共享字段与台账不一致，普通导入只能追加明细、不能修改它们："
                    + "、".join(conflicts)
                )

            if existing_project_id is not None:
                project_id = existing_project_id
            else:
                # 框架项目只承载框架级字段：客户经理、分公司、部门、三级团队。
                # 客户单位/最终用户/区域平台属于子项目，写在下面的 sub_project 上。
                project_id = _execute_scalar(
                    conn,
                    """
                    INSERT INTO project
                      (project_code, project_name, department, branch_company,
                       account_manager, team_level3_name)
                    VALUES
                      (:project_code, :project_name, :department, :branch_company,
                       :account_manager, :team_level3_name)
                    ON DUPLICATE KEY UPDATE
                      id = LAST_INSERT_ID(id),
                      project_name = COALESCE(project_name, VALUES(project_name))
                    """,
                    {
                        "project_code": project_code,
                        "project_name": _as_text(_row_value(row, position(14, 14))),
                        "department": department,
                        "branch_company": _as_text(_row_value(row, position(4, 4))),
                        "account_manager": _as_text(_row_value(row, position(5, 5))),
                        "team_level3_name": _as_text(_row_value(row, position(9, 9))),
                    },
                )

            sales_order_id = _execute_scalar(
                conn,
                """
                INSERT INTO sales_order
                  (project_id, import_batch_id, source_excel_row_no, gross_net_type, order_no,
                   order_date, business_type, statistic_category, close_status)
                VALUES
                  (:project_id, :batch_id, :excel_row_no, :gross_net_type, :order_no,
                   :order_date, :business_type, :statistic_category, :close_status)
                ON DUPLICATE KEY UPDATE
                  id = LAST_INSERT_ID(id),
                  order_date = COALESCE(VALUES(order_date), order_date),
                  business_type = COALESCE(VALUES(business_type), business_type),
                  statistic_category = COALESCE(VALUES(statistic_category), statistic_category),
                  close_status = COALESCE(VALUES(close_status), close_status)
                """,
                {
                    "project_id": project_id,
                    "batch_id": batch_id,
                    "excel_row_no": excel_row_no,
                    "gross_net_type": _as_text(_row_value(row, position(1, 1))),
                    "order_no": order_no,
                    "order_date": _as_date(_row_value(row, position(6, 6))),
                    "business_type": _as_text(_row_value(row, position(7, 7))),
                    "statistic_category": _as_text(_row_value(row, position(8, 8))),
                    "close_status": _as_text(_row_value(row, 89 if template_layout else position(96, 87))),
                },
            )

            # 子项目是订单下的独立层级：同一订单可以有多个子项目，
            # 每个子项目各自保存客户单位、最终用户、区域平台。
            project_name = _as_text(_row_value(row, position(14, 14)))
            sub_project_id, sub_project_conflicts = _find_or_create_sub_project(
                conn, sales_order_id, row, position, project_name, excel_row_no
            )
            if sub_project_conflicts:
                raise ValueError(
                    f"第 {excel_row_no} 行订单 {order_no} 的子项目“{project_name or '（空名称）'}”已在台账中，"
                    "以下子项目字段与台账不一致，普通导入只能追加明细、不能修改它们："
                    + "、".join(sub_project_conflicts)
                )

            quantity = _as_decimal(_row_value(row, position(18, 18)))
            sales_tax_rate = _as_tax_rate(_row_value(row, position(19))) if latest_layout else None
            sales_unit_price_no_tax = _as_decimal(_row_value(row, position(20, 19)))
            sales_unit_price = _as_decimal(_row_value(row, position(21, 20)))
            sales_unit_price, revenue_no_tax, order_value = _calculated_prices(
                quantity, sales_tax_rate, sales_unit_price_no_tax, sales_unit_price
            )
            revenue_no_tax = revenue_no_tax or _as_decimal(_row_value(row, position(22, 21)))
            order_value = order_value or _as_decimal(_row_value(row, position(23, 22)))

            # 判重限定在同一子项目内：物资名称、规格、销售单价、数量、采购厂商五项组合唯一。
            # 跨子项目、跨订单、跨框架允许五项完全相同。
            if strict_template and _order_line_exists(
                conn,
                sub_project_id,
                _as_text(_row_value(row, position(15, 15))),
                _as_text(_row_value(row, position(16, 16))),
                quantity,
                sales_unit_price,
                _as_text(_row_value(row, position(24, 23))),
            ):
                raise ValueError(
                    f"订单 {order_no} 的子项目“{project_name or '（空名称）'}”下，"
                    "同名同规格同数量同单价同采购厂商的明细已存在"
                )

            order_line_id = _execute_scalar(
                conn,
                """
                INSERT INTO order_line
                  (sales_order_id, sub_project_id, raw_row_id, source_excel_row_no, project_name,
                   goods_name, specification_model, unit_name, quantity, sales_tax_rate,
                   sales_unit_price_no_tax, sales_unit_price, revenue_no_tax, order_value)
                VALUES
                  (:sales_order_id, :sub_project_id, :raw_row_id, :excel_row_no, :project_name,
                   :goods_name, :specification_model, :unit_name, :quantity, :sales_tax_rate,
                   :sales_unit_price_no_tax, :sales_unit_price, :revenue_no_tax, :order_value)
                """,
                {
                    "sales_order_id": sales_order_id,
                    "sub_project_id": sub_project_id,
                    "raw_row_id": raw_row_id,
                    "excel_row_no": excel_row_no,
                    "project_name": project_name,
                    "goods_name": _as_text(_row_value(row, position(15, 15))),
                    "specification_model": _as_text(_row_value(row, position(16, 16))),
                    "unit_name": _as_text(_row_value(row, position(17, 17))),
                    "quantity": quantity,
                    "sales_tax_rate": sales_tax_rate,
                    "sales_unit_price_no_tax": sales_unit_price_no_tax,
                    "sales_unit_price": sales_unit_price,
                    "revenue_no_tax": revenue_no_tax,
                    "order_value": order_value,
                },
            )

            purchase_tax_rate = _as_tax_rate(_row_value(row, position(25))) if latest_layout else None
            purchase_unit_price_no_tax = _as_decimal(_row_value(row, position(26, 24)))
            purchase_unit_price = _as_decimal(_row_value(row, position(27, 25)))
            purchase_unit_price, cost_no_tax, purchase_amount = _calculated_prices(
                quantity, purchase_tax_rate, purchase_unit_price_no_tax, purchase_unit_price
            )
            cost_no_tax = cost_no_tax or _as_decimal(_row_value(row, position(28, 26)))
            purchase_amount = purchase_amount or _as_decimal(_row_value(row, position(29, 27)))

            conn.execute(
                text(
                    """
                    INSERT INTO purchase_info
                      (order_line_id, supplier_name, purchase_tax_rate, purchase_unit_price_no_tax,
                       purchase_unit_price, cost_no_tax, purchase_amount, labor_cost, other_cost)
                    VALUES
                      (:order_line_id, :supplier_name, :purchase_tax_rate, :purchase_unit_price_no_tax,
                       :purchase_unit_price, :cost_no_tax, :purchase_amount, :labor_cost, :other_cost)
                    """
                ),
                {
                    "order_line_id": order_line_id,
                    "supplier_name": _as_text(_row_value(row, position(24, 23))),
                    "purchase_tax_rate": purchase_tax_rate,
                    "purchase_unit_price_no_tax": purchase_unit_price_no_tax,
                    "purchase_unit_price": purchase_unit_price,
                    "cost_no_tax": cost_no_tax,
                    "purchase_amount": purchase_amount,
                    "labor_cost": _as_decimal(_row_value(row, 90)) if template_layout else None,
                    "other_cost": _as_decimal(_row_value(row, 91)) if template_layout else None,
                },
            )

            conn.execute(
                text(
                    """
                    INSERT INTO delivery_record
                      (order_line_id, delivery_date, delivery_quantity, delivery_revenue_no_tax,
                       delivery_value, delivery_cost_no_tax, delivery_cost, pending_delivery_quantity,
                       pending_delivery_amount_no_tax, pending_delivery_amount)
                    VALUES
                      (:order_line_id, :delivery_date, :delivery_quantity, :delivery_revenue_no_tax,
                       :delivery_value, :delivery_cost_no_tax, :delivery_cost, :pending_delivery_quantity,
                       :pending_delivery_amount_no_tax, :pending_delivery_amount)
                    """
                ),
                {
                    "order_line_id": order_line_id,
                    "delivery_date": _as_date(_row_value(row, position(30, 28))),
                    "delivery_quantity": _as_decimal(_row_value(row, position(31, 29))),
                    "delivery_revenue_no_tax": _as_decimal(_row_value(row, position(32, 30))),
                    "delivery_value": _as_decimal(_row_value(row, position(33, 31))),
                    "delivery_cost_no_tax": _as_decimal(_row_value(row, position(34, 32))),
                    "delivery_cost": _as_decimal(_row_value(row, position(35, 33))),
                    "pending_delivery_quantity": _as_decimal(_row_value(row, position(36, 34))),
                    "pending_delivery_amount_no_tax": _as_decimal(_row_value(row, position(37, 35))),
                    "pending_delivery_amount": _as_decimal(_row_value(row, position(38, 36))),
                },
            )

            conn.execute(
                text(
                    """
                    INSERT INTO purchase_contract
                      (order_line_id, purchase_contract_no, payment_terms, performance_period, signed_amount, unsigned_amount)
                    VALUES
                      (:order_line_id, :purchase_contract_no, :payment_terms, :performance_period, :signed_amount, :unsigned_amount)
                    """
                ),
                {
                    "order_line_id": order_line_id,
                    "purchase_contract_no": _as_text(_row_value(row, position(39, 37))),
                    "payment_terms": _as_text(_row_value(row, position(40, 38))),
                    "performance_period": _as_text(_row_value(row, position(41, 39))),
                    "signed_amount": _as_decimal(_row_value(row, position(42, 40))),
                    "unsigned_amount": _as_decimal(_row_value(row, position(43, 41))),
                },
            )

            _insert_phase(conn, "purchase_invoice", order_line_id, 1, row, position(44, 42), position(45, 43), position(46, 44))
            _insert_warehouse(
                conn,
                order_line_id,
                row,
                position(47, 45),
                position(48, 46),
                position(49, 47),
                0 if template_layout else position(50),
            )
            if template_layout:
                _insert_finance_payment(conn, order_line_id, 1, row, 50, 51, 52)
            elif latest_layout:
                _insert_finance_check(conn, order_line_id, 1, row, 51, 52, 53)
                _insert_finance_check(conn, order_line_id, 2, row, 54, 55, 56)
                _insert_finance_payment(conn, order_line_id, 1, row, 57, 58, 59)
                _insert_finance_payment(conn, order_line_id, 2, row, 60, 61, 62)
            else:
                _insert_finance_payment(conn, order_line_id, 1, row, 48, 49, 50)
            _insert_payment(
                conn,
                order_line_id,
                1,
                _as_date(_row_value(row, 54 if template_layout else position(65, 52))),
                _as_date(_row_value(row, 55 if template_layout else position(66, 53))),
                _as_text(_row_value(row, 56 if template_layout else position(67, 54))),
                _as_decimal(_row_value(row, 57 if template_layout else position(68, 55))),
            )
            _insert_payment(
                conn,
                order_line_id,
                2,
                None,
                _as_date(_row_value(row, 58 if template_layout else position(69, 56))),
                _as_text(_row_value(row, 59 if template_layout else position(70, 57))),
                _as_decimal(_row_value(row, 60 if template_layout else position(71, 58))),
            )

            conn.execute(
                text(
                    """
                    INSERT INTO sales_contract
                      (order_line_id, contract_signed_date, sales_contract_no, contract_value,
                       performance_period, unsigned_contract_amount)
                    VALUES
                      (:order_line_id, :contract_signed_date, :sales_contract_no, :contract_value,
                       :performance_period, :unsigned_contract_amount)
                    """
                ),
                {
                    "order_line_id": order_line_id,
                    "contract_signed_date": _as_date(_row_value(row, 68 if template_layout else position(76, 66))),
                    "sales_contract_no": _as_text(_row_value(row, 69 if template_layout else position(77, 67))),
                    "contract_value": _as_decimal(_row_value(row, 70 if template_layout else position(78, 68))),
                    "performance_period": _as_text(_row_value(row, 71 if template_layout else position(79, 69))),
                    "unsigned_contract_amount": (
                        _as_decimal(_row_value(row, 72))
                        if template_layout
                        else (None if latest_layout else _as_decimal(_row_value(row, 70)))
                    ),
                },
            )

            conn.execute(
                text(
                    """
                    INSERT INTO sales_invoice
                      (order_line_id, phase_no, invoice_doc_no, invoice_date, invoice_date_text,
                       invoice_no, invoice_amount, pending_invoice_amount, delivered_not_invoiced_amount)
                    VALUES
                      (:order_line_id, 1, :invoice_doc_no, :invoice_date, :invoice_date_text,
                       :invoice_no, :invoice_amount, :pending_invoice_amount, :delivered_not_invoiced_amount)
                    """
                ),
                {
                    "order_line_id": order_line_id,
                    "invoice_doc_no": _as_text(_row_value(row, 73 if template_layout else position(80, 71))),
                    "invoice_date": _as_date(_row_value(row, 74 if template_layout else position(81, 72))),
                    "invoice_date_text": _as_text(_row_value(row, 74 if template_layout else position(81, 72))),
                    "invoice_no": _as_text(_row_value(row, 75 if template_layout else position(82, 73))),
                    "invoice_amount": _as_decimal(_row_value(row, 76 if template_layout else position(83, 74))),
                    "pending_invoice_amount": _as_decimal(_row_value(row, 77 if template_layout else position(84, 75))),
                    "delivered_not_invoiced_amount": _as_decimal(_row_value(row, 78 if template_layout else position(85, 76))),
                },
            )

            _insert_receipt(
                conn,
                order_line_id,
                1,
                row,
                79 if template_layout else position(86, 77),
                80 if template_layout else position(87, 78),
                81 if template_layout else position(88, 79),
                82 if template_layout else position(89, 80),
            )
            _insert_receipt(
                conn,
                order_line_id,
                2,
                row,
                83 if template_layout else position(90, 81),
                84 if template_layout else position(91, 82),
                85 if template_layout else position(92, 83),
                86 if template_layout else position(93, 84),
            )

            success_rows += 1
        except PermissionError:
            raise
        except Exception as exc:  # noqa: BLE001 - record row-level import issue and continue
            failed_rows += 1
            error_messages.append(f"第 {excel_row_no} 行：{exc}")
            conn.execute(
                text(
                    """
                    INSERT INTO ledger_raw_row
                      (import_batch_id, excel_row_no, row_hash, raw_json, parse_status, parse_message, project_code, order_no)
                    VALUES
                      (:batch_id, :excel_row_no, :row_hash, CAST(:raw_json AS JSON), 'failed', :message, :project_code, :order_no)
                    ON DUPLICATE KEY UPDATE parse_status = 'failed', parse_message = VALUES(parse_message)
                    """
                ),
                {
                    "batch_id": batch_id,
                    "excel_row_no": excel_row_no,
                    "row_hash": row_hash,
                    "raw_json": raw_json,
                    "message": str(exc)[:1000],
                    "project_code": project_code,
                    "order_no": order_no,
                },
            )

    conn.execute(
        text(
            """
            UPDATE import_batch
            SET total_rows = :total_rows,
                success_rows = :success_rows,
                failed_rows = :failed_rows,
                status = :status
            WHERE id = :batch_id
            """
        ),
        {
            "batch_id": batch_id,
            "total_rows": success_rows + failed_rows,
            "success_rows": success_rows,
            "failed_rows": failed_rows,
            "status": "completed" if failed_rows == 0 else "failed",
        },
    )

    workbook.close()
    if strict_template and success_rows + failed_rows == 0:
        raise ValueError("导入文件中没有可导入的业务数据")

    detail = f"导入 {source_file_name}: 成功 {success_rows} 行，失败 {failed_rows} 行"
    if skipped_rows:
        detail = f"{detail}，跳过非业务行 {skipped_rows} 行"
    if user and failed_rows == 0:
        write_operation_log(
            conn,
            user,
            "订单管理",
            "import_excel",
            detail,
            after={"batch_id": batch_id, "success_rows": success_rows, "source_file": source_file_name},
        )
    elif user is None:
        conn.execute(
            text(
                """
                INSERT INTO operation_log (user_name, module_name, action_name, detail, status)
                VALUES ('system', '数据导入', 'import_excel', :detail, :status)
                """
            ),
            {"detail": detail, "status": "success" if failed_rows == 0 else "failed"},
        )

    return {
        "batch_id": batch_id,
        "source_file": source_file_name,
        "success_rows": success_rows,
        "failed_rows": failed_rows,
        "skipped_rows": skipped_rows,
        "errors": error_messages[:20],
    }


def _order_line_exists(
    conn: Connection,
    sub_project_id: int,
    goods_name: str | None,
    specification_model: str | None,
    quantity: Decimal | None,
    sales_unit_price: Decimal | None,
    supplier_name: str | None,
) -> bool:
    """同一子项目内，物资名称、规格、销售单价、数量、采购厂商五项组合是否已存在。

    sub_project_id 已经隐含了框架项目与订单，所以判重限定在同一子项目内；
    跨子项目允许五项完全相同。
    """
    return (
        find_duplicate_line(
            conn,
            sub_project_id,
            goods_name,
            specification_model,
            quantity,
            sales_unit_price,
            supplier_name,
        )
        is not None
    )


def _insert_phase(conn: Connection, table: str, order_line_id: int, phase_no: int, row: tuple[Any, ...], date_pos: int, code_pos: int, amount_pos: int) -> None:
    value = _as_decimal(_row_value(row, amount_pos))
    if value is None and _row_value(row, date_pos) is None and _row_value(row, code_pos) is None:
        return
    if table == "purchase_invoice":
        conn.execute(
            text(
                """
                INSERT INTO purchase_invoice
                  (order_line_id, phase_no, received_invoice_date, received_invoice_date_text, invoice_no, invoice_amount)
                VALUES
                  (:order_line_id, :phase_no, :date_value, :date_text, :code_value, :amount_value)
                """
            ),
            {
                "order_line_id": order_line_id,
                "phase_no": phase_no,
                "date_value": _as_date(_row_value(row, date_pos)),
                "date_text": _as_text(_row_value(row, date_pos)),
                "code_value": _as_text(_row_value(row, code_pos)),
                "amount_value": value,
            },
        )
    elif table == "warehouse_entry":
        conn.execute(
            text(
                """
                INSERT INTO warehouse_entry
                  (order_line_id, phase_no, warehouse_date, warehouse_date_text, voucher_no, warehouse_amount)
                VALUES
                  (:order_line_id, :phase_no, :date_value, :date_text, :code_value, :amount_value)
                """
            ),
            {
                "order_line_id": order_line_id,
                "phase_no": phase_no,
                "date_value": _as_date(_row_value(row, date_pos)),
                "date_text": _as_text(_row_value(row, date_pos)),
                "code_value": _as_text(_row_value(row, code_pos)),
                "amount_value": value,
            },
        )


def _insert_warehouse(
    conn: Connection,
    order_line_id: int,
    row: tuple[Any, ...],
    date_pos: int,
    voucher_pos: int,
    amount_pos: int,
    amount_no_tax_pos: int,
) -> None:
    date_value = _row_value(row, date_pos)
    voucher = _row_value(row, voucher_pos)
    amount = _as_decimal(_row_value(row, amount_pos))
    amount_no_tax = _as_decimal(_row_value(row, amount_no_tax_pos))
    if date_value is None and voucher is None and amount is None and amount_no_tax is None:
        return
    conn.execute(
        text(
            """
            INSERT INTO warehouse_entry
              (order_line_id, phase_no, warehouse_date, warehouse_date_text, voucher_no,
               warehouse_amount, warehouse_amount_no_tax)
            VALUES
              (:order_line_id, 1, :warehouse_date, :warehouse_date_text, :voucher_no,
               :warehouse_amount, :warehouse_amount_no_tax)
            """
        ),
        {
            "order_line_id": order_line_id,
            "warehouse_date": _as_date(date_value),
            "warehouse_date_text": _as_text(date_value),
            "voucher_no": _as_text(voucher),
            "warehouse_amount": amount,
            "warehouse_amount_no_tax": amount_no_tax,
        },
    )


def _insert_finance_check(
    conn: Connection,
    order_line_id: int,
    phase_no: int,
    row: tuple[Any, ...],
    date_pos: int,
    amount_pos: int,
    voucher_pos: int,
) -> None:
    date_value = _row_value(row, date_pos)
    amount = _as_decimal(_row_value(row, amount_pos))
    voucher = _row_value(row, voucher_pos)
    if date_value is None and amount is None and voucher is None:
        return
    conn.execute(
        text(
            """
            INSERT INTO finance_invoice_check
              (order_line_id, phase_no, received_invoice_date, received_invoice_date_text,
               received_invoice_amount, voucher_code)
            VALUES
              (:order_line_id, :phase_no, :received_invoice_date, :received_invoice_date_text,
               :received_invoice_amount, :voucher_code)
            """
        ),
        {
            "order_line_id": order_line_id,
            "phase_no": phase_no,
            "received_invoice_date": _as_date(date_value),
            "received_invoice_date_text": _as_text(date_value),
            "received_invoice_amount": amount,
            "voucher_code": _as_text(voucher),
        },
    )


def _insert_finance_payment(
    conn: Connection,
    order_line_id: int,
    phase_no: int,
    row: tuple[Any, ...],
    date_pos: int,
    voucher_pos: int,
    amount_pos: int,
) -> None:
    date_value = _row_value(row, date_pos)
    voucher = _row_value(row, voucher_pos)
    amount = _as_decimal(_row_value(row, amount_pos))
    if date_value is None and voucher is None and amount is None:
        return
    conn.execute(
        text(
            """
            INSERT INTO finance_payment_entry
              (order_line_id, phase_no, payment_date, payment_date_text, voucher_code, booked_amount)
            VALUES
              (:order_line_id, :phase_no, :payment_date, :payment_date_text, :voucher_code, :booked_amount)
            """
        ),
        {
            "order_line_id": order_line_id,
            "phase_no": phase_no,
            "payment_date": _as_date(date_value),
            "payment_date_text": _as_text(date_value),
            "voucher_code": _as_text(voucher),
            "booked_amount": amount,
        },
    )


def _insert_payment(conn: Connection, order_line_id: int, phase_no: int, due_date: date | None, payment_date: date | None, voucher_no: str | None, amount: Decimal | None) -> None:
    if due_date is None and payment_date is None and voucher_no is None and amount is None:
        return
    conn.execute(
        text(
            """
            INSERT INTO purchase_payment
              (order_line_id, phase_no, due_payment_date, payment_date, payment_date_text, payment_voucher_no, payment_amount)
            VALUES
              (:order_line_id, :phase_no, :due_date, :payment_date, :payment_date_text, :voucher_no, :amount)
            """
        ),
        {
            "order_line_id": order_line_id,
            "phase_no": phase_no,
            "due_date": due_date,
            "payment_date": payment_date,
            "payment_date_text": payment_date.isoformat() if payment_date else None,
            "voucher_no": voucher_no,
            "amount": amount,
        },
    )


def _insert_receipt(conn: Connection, order_line_id: int, phase_no: int, row: tuple[Any, ...], date_pos: int, notice_pos: int, amount_pos: int, ratio_pos: int) -> None:
    amount = _as_decimal(_row_value(row, amount_pos))
    if amount is None and _row_value(row, date_pos) is None and _row_value(row, notice_pos) is None:
        return
    conn.execute(
        text(
            """
            INSERT INTO sales_receipt
              (order_line_id, phase_no, receipt_date, receipt_date_text, payment_notice_no, receipt_amount, receipt_ratio)
            VALUES
              (:order_line_id, :phase_no, :receipt_date, :receipt_date_text, :payment_notice_no, :receipt_amount, :receipt_ratio)
            """
        ),
        {
            "order_line_id": order_line_id,
            "phase_no": phase_no,
            "receipt_date": _as_date(_row_value(row, date_pos)),
            "receipt_date_text": _as_text(_row_value(row, date_pos)),
            "payment_notice_no": _as_text(_row_value(row, notice_pos)),
            "receipt_amount": amount,
            "receipt_ratio": _as_decimal(_row_value(row, ratio_pos)),
        },
    )
