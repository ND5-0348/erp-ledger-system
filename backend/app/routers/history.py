"""Authorized stable-identity changes and complete history/phase export (H4)."""
from __future__ import annotations

from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import text, bindparam
from ..auth import CurrentUser, get_current_user, require_permission, can_access_department
from ..edit_versions import line_context
from ..db import db
from ..audit import write_operation_log
from ..write_guard import business_write
from ..ledger_history import conflicts_in_project, register_current_number, register_current_manager
from ..ledger_excel import filtered_export_rows, content_disposition
from ..serializers import clean_row, clean_rows
from ..validation import BusinessDate

router = APIRouter(prefix='/api/history', tags=['history'])


def context(conn, line_id, user, write=False):
    row = conn.execute(text('''SELECT ol.id AS order_line_id, so.id AS sales_order_id,
      p.id AS project_id, p.project_code, so.order_no, ol.source_preserved,
      CASE WHEN ol.source_preserved=1 THEN ol.line_account_manager ELSE p.account_manager END AS account_manager,
      CASE WHEN ol.source_preserved=1 THEN ol.line_department ELSE p.department END AS department,
      CASE WHEN ol.source_preserved=1 THEN ol.line_branch_company ELSE p.branch_company END AS branch_company,
      CASE WHEN ol.source_preserved=1 THEN ol.line_team_level3_name ELSE p.team_level3_name END AS team_level3_name
      FROM order_line ol JOIN sales_order so ON so.id=ol.sales_order_id AND so.deleted_at IS NULL
      JOIN project p ON p.id=so.project_id AND p.deleted_at IS NULL
      WHERE ol.id=:id AND ol.deleted_at IS NULL'''), {'id':line_id}).mappings().first()
    if not row:
        raise HTTPException(404, '订单明细不存在')
    if not can_access_department(user, row['department'], require_entry=write):
        raise HTTPException(403, '没有该部门的访问权限')
    if write and row['source_preserved']:
        raise HTTPException(409, '原表导入按明细保存归属，不能用框架整体交接或历史改号覆盖原表多值；请按明细维护')
    return dict(row)


@router.get('/lines/{line_id}')
def read_history(line_id: int, user: CurrentUser = Depends(get_current_user)):
    with db() as conn:
        row = context(conn, line_id, user)
        edit_context = line_context(conn,[line_id])
        numbers = conn.execute(text('SELECT order_no, history_order, source FROM sales_order_number_history WHERE sales_order_id=:id ORDER BY history_order'), {'id':row['sales_order_id']}).mappings().all()
        managers = conn.execute(text('SELECT manager_name, history_order, effective_from, source FROM project_manager_history WHERE project_id=:id ORDER BY history_order'), {'id':row['project_id']}).mappings().all()
        if row['source_preserved']:
            numbers = [{'order_no':row['order_no'], 'history_order':1, 'source':'原表记录（未推断改号）'}]
            managers = [{'manager_name':row['account_manager'], 'history_order':1, 'effective_from':None, 'source':'原表明细归属（未推断交接）'}]
        count = conn.execute(text('SELECT COUNT(*) FROM order_line ol JOIN sales_order so ON so.id=ol.sales_order_id AND so.deleted_at IS NULL WHERE so.project_id=:id AND ol.deleted_at IS NULL'), {'id':row['project_id']}).scalar_one()
    return {'edit_context':edit_context,'current':clean_row(row), 'order_numbers':clean_rows(numbers), 'managers':clean_rows(managers), 'affected_lines':count}


class Rename(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)
    order_line_id: int = Field(gt=0)
    expected_order_no: str = Field(min_length=1, max_length=64)
    order_no: str = Field(min_length=1, max_length=64)
    reason: str = Field(min_length=1, max_length=500)

    @field_validator('order_no')
    @classmethod
    def single_number(cls, value):
        if any(c in value for c in '/／;；\n\r'):
            raise ValueError('一次只能变更为一个订单号，历史由系统保留')
        return value


class Renames(BaseModel):
    items: list[Rename] = Field(min_length=1, max_length=500)


@router.post('/rename-orders')
def rename_orders(payload: Renames, user: CurrentUser = Depends(require_permission('order_edit'))):
    with business_write() as conn:
        seen = set()
        for item in payload.items:
            row = context(conn, item.order_line_id, user, True)
            oid = row['sales_order_id']
            if oid in seen:
                raise HTTPException(422, '同一订单只能提交一次改号')
            seen.add(oid)
            if row['order_no'] != item.expected_order_no:
                raise HTTPException(409, '订单号已变更，请重新读取后确认')
            if row['order_no'] == item.order_no:
                continue
            if conflicts_in_project(conn, row['project_id'], [item.order_no], exclude_sales_order_id=oid):
                raise HTTPException(409, '目标编号已属于同框架的其他订单（含历史号）')
            register_current_number(conn, oid, row['order_no'], source='manual_change')
            conn.execute(text('''INSERT INTO sales_order_number_history (sales_order_id, order_no, history_order, source)
              SELECT :id, :number, COALESCE(MAX(history_order),0)+1, 'manual_change'
              FROM sales_order_number_history WHERE sales_order_id=:id'''), {'id':oid, 'number':item.order_no})
            conn.execute(text('UPDATE sales_order SET order_no=:number WHERE id=:id'), {'id':oid,'number':item.order_no})
            write_operation_log(conn, user, '订单管理', 'rename_order', f'变更订单号：{item.reason}', before={'sales_order_id':oid,'order_no':row['order_no']}, after={'sales_order_id':oid,'order_no':item.order_no})
    return {'updated':len(seen)}


class Transfer(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)
    order_line_id: int = Field(gt=0)
    expected: dict[str, str | None]
    account_manager: str = Field(min_length=1, max_length=64)
    department: str = Field(min_length=1, max_length=64)
    branch_company: str | None = Field(default=None, max_length=64)
    team_level3_name: str | None = Field(default=None, max_length=64)
    effective_from: BusinessDate | None = None
    reason: str = Field(min_length=1, max_length=500)

    @field_validator('account_manager')
    @classmethod
    def single_manager(cls, value):
        if any(c in value for c in '/／;；\n\r'):
            raise ValueError('一次只能交接给一位客户经理')
        return value


@router.post('/transfer-project')
def transfer_project(payload: Transfer, user: CurrentUser = Depends(require_permission('order_edit'))):
    keys = ('account_manager', 'department', 'branch_company', 'team_level3_name')
    with business_write() as conn:
        row = context(conn, payload.order_line_id, user, True)
        if not can_access_department(user, payload.department, require_entry=True):
            raise HTTPException(403, '没有目标部门的维护权限')
        if set(payload.expected) != set(keys) or any(payload.expected[k] != row[k] for k in keys):
            raise HTTPException(409, '框架归属已变更，请重新读取后确认')
        data = {k:getattr(payload,k) for k in keys}
        if all(data[k] == row[k] for k in keys):
            return {'updated':False}
        pid=row['project_id']
        register_current_manager(conn, pid, row['account_manager'], source='manual_change')
        if row['account_manager'] != payload.account_manager:
            conn.execute(text('''INSERT INTO project_manager_history (project_id, manager_name, history_order, effective_from, source)
              SELECT :id, :manager, COALESCE(MAX(history_order),0)+1, :date, 'manual_change'
              FROM project_manager_history WHERE project_id=:id'''), {'id':pid,'manager':payload.account_manager,'date':payload.effective_from})
        conn.execute(text('UPDATE project SET account_manager=:account_manager, department=:department, branch_company=:branch_company, team_level3_name=:team_level3_name WHERE id=:id'), {'id':pid,**data})
        write_operation_log(conn,user,'订单管理','transfer_project',f'框架整体交接：{payload.reason}', before={k:row[k] for k in keys},after={'project_id':pid,**data})
    return {'updated':True}


@router.get('/export')
def export_history(project_id: str | None = None, department: str | None = None,
                   manager: str | None = None, include_history_manager: bool = False,
                   client_unit: str | None = None, order_id: str | None = None,
                   order_status: str | None = None, supplier_name: str | None = None,
                   start_date: BusinessDate | None = None, end_date: BusinessDate | None = None,
                   invoice_start_date: BusinessDate | None = None, invoice_end_date: BusinessDate | None = None,
                   user: CurrentUser = Depends(get_current_user)):
    filters = dict(project_id=project_id, department=department, manager=manager,
                   include_history_manager=include_history_manager, client_unit=client_unit,
                   order_id=order_id, order_status=order_status, supplier_name=supplier_name,
                   start_date=start_date,end_date=end_date,invoice_start_date=invoice_start_date,invoice_end_date=invoice_end_date)
    if (start_date and end_date and start_date>end_date) or (invoice_start_date and invoice_end_date and invoice_start_date>invoice_end_date):
        raise HTTPException(422, '开始日期不能晚于结束日期')
    wb=Workbook(); info=wb.active; info.title='说明'
    info.append(['历史及完整期次明细；仅供查询核对，不是91列回导模板。金额以文本保留精度。'])
    with db() as conn:
        rows=filtered_export_rows(conn,user,filters)
        ids=[r['order_line_id'] for r in rows]
        sheets=[('明细',rows)]
        if ids:
            base='FROM order_line ol JOIN sales_order so ON so.id=ol.sales_order_id JOIN project p ON p.id=so.project_id'
            for title,sql in [
                ('订单号历史',f'SELECT DISTINCT p.project_code, so.id AS sales_order_id, so.order_no AS current_order_no, h.order_no, h.history_order, h.source {base} JOIN sales_order_number_history h ON h.sales_order_id=so.id WHERE ol.id IN :ids ORDER BY p.project_code, so.id, h.history_order'),
                ('负责人历史',f'SELECT DISTINCT p.project_code, p.account_manager AS current_manager, h.manager_name, h.history_order, h.effective_from, h.source {base} JOIN project_manager_history h ON h.project_id=p.id WHERE ol.id IN :ids ORDER BY p.project_code, h.history_order')]:
                sheets.append((title,conn.execute(text(sql).bindparams(bindparam('ids',expanding=True)),{'ids':ids}).mappings().all()))
            for title,table in [('采购收票','purchase_invoice'),('采购入库','warehouse_entry'),('财务收票','finance_invoice_check'),('财务付款','finance_payment_entry'),('采购付款','purchase_payment'),('销售开票','sales_invoice'),('销售回款','sales_receipt')]:
                sql=f'SELECT p.project_code, so.order_no, ol.project_name, ol.goods_name, t.* {base} JOIN {table} t ON t.order_line_id=ol.id AND t.deleted_at IS NULL WHERE ol.id IN :ids ORDER BY ol.id,t.phase_no,t.id'
                sheets.append((title,conn.execute(text(sql).bindparams(bindparam('ids',expanding=True)),{'ids':ids}).mappings().all()))
        for title,data in sheets:
            ws=wb.create_sheet(title)
            if data:
                ws.append(list(data[0].keys()))
                for row in data:
                    ws.append([None if v is None else str(v) for v in row.values()])
                    for cell in ws[ws.max_row]:
                        if cell.value is not None: cell.data_type='s'
                ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        write_operation_log(conn,user,'订单管理','export_history',f'导出历史及期次明细，共 {len(ids)} 条明细')
    out=BytesIO(); wb.save(out); wb.close()
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':content_disposition('历史及期次明细.xlsx')})
